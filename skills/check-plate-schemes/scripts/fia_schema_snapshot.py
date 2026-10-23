"""Read-only FIA schema snapshot; see references/compatibility.json and LICENSE-FIA.txt.

Extracted unchanged validation functions from the recorded upstream commit.
This module performs no image analysis and no workbook writes.
"""
import hashlib
import io
import json
import re
import unicodedata
from pathlib import Path
from openpyxl import load_workbook

METHODS = ('Nuclei', 'Foci', 'Nuclear intensity', 'Field signal', 'Fiber')
PREFIXES = dict(zip(METHODS, ('Nuclei', 'Foci', 'NuclearIntensity', 'FieldSignal', 'Fiber')))

def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     separators=(',', ':')).encode('utf-8')).hexdigest()


def marker_id(method, name):
    """Portable, collision-resistant identity, independent of source channel."""
    label = unicodedata.normalize('NFC', name.strip())
    slug = re.sub(r'[^A-Za-z0-9]+', '_', label).strip('_')[:40] or 'marker'
    return f'{PREFIXES[method]}_{slug}_{digest([method, label])[:12]}'


def _literal(cell):
    if cell.data_type in ('f', 'e'):
        raise ValueError(f'{cell.parent.title}!{cell.coordinate}: use literal values, not formulas/errors')
    if cell.value is None:
        return ''
    if not isinstance(cell.value, str):
        # This is a workbook-content validation error, not a Python API type error.
        raise ValueError(f'{cell.parent.title}!{cell.coordinate}: expected text')  # noqa: TRY004
    return unicodedata.normalize('NFC', cell.value.strip())


def _color(cell):
    color = cell.fill.fgColor
    if cell.fill.patternType != 'solid' or color.type != 'rgb':
        raise ValueError(f'{cell.parent.title}!{cell.coordinate}: use a direct solid RGB block fill')
    return str(color.rgb)[-6:].upper()


class PlateDesign:
    def __init__(self, path, content, wells, panels):
        self.path = Path(path)
        self.content = content
        self.wells = wells
        self.panels = panels
        self.workbook_sha256 = hashlib.sha256(content).hexdigest()

    def channel(self, well, method, marker=None):
        values = [row['channel'] for row in self.wells.get(well, {}).get('channels', [])
                  if row and row['method'] == method and (marker is None or row['id'] == marker)]
        if len(values) != 1:
            raise ValueError(f'{well}: expected one {method} channel for {marker or method}; found {len(values)}')
        return values[0]

    def assigned(self, well, method=None, marker=None):
        return well in self.wells and (method is None or any(
            row and row['method'] == method and (marker is None or row['id'] == marker)
            for row in self.wells[well]['channels']))

    def validate_channels(self, well, count):
        expected = sum(row is not None for row in self.wells[well]['channels'])
        if int(count) != expected:
            raise ValueError(f'{well}: Excel declares {expected} existing channels, original has {count}. '
                             'Use Ignore for an existing unused channel; blank only for absent channels.')

    def markers(self, method):
        result = {}
        for well, item in self.wells.items():
            for row in item['channels']:
                if row and row['method'] == method:
                    marker = result.setdefault(row['id'], {k: row[k] for k in ('id', 'name', 'method')})
                    marker.setdefault('channels', {})[well] = row['channel']
        return list(result.values())

    def scope_hash(self, method=None, marker=None):
        rows = {}
        for well, item in self.wells.items():
            if not self.assigned(well, method, marker):
                continue
            if method == 'Nuclei':
                assignments = [self.channel(well, 'Nuclei')]
            elif method:
                assignments = [row for row in item['channels'] if row and (
                    row['method'] == 'Nuclei' or (row['method'] == method and
                                                (marker is None or row['id'] == marker)))]
            else:
                assignments = item['channels']
            rows[well] = {'channels': assignments,
                          'count': sum(row is not None for row in item['channels'])}
        return digest(rows)

    def report_rows(self):
        order = {group: index for index, group in enumerate(
            (g for panel in self.panels for g in panel['groups']), 1)}
        positions = {group: panel for panel in self.panels for group in panel['groups']}
        return [{'Well': well, 'Group': item['condition'], 'Color': item['block'],
                 'Is_control': item['control'], 'Excel_cell': item['cell'],
                 'Group_order': order[item['condition']], 'Order_source': 'plot_layout',
                 'Order_cell': positions[item['condition']]['cells'][item['condition']],
                 'Order_group_cell': positions[item['condition']]['cells'][item['condition']]}
                for well, item in self.wells.items()]


def read_design(path, allow_empty=False):
    path = Path(path)
    content = path.read_bytes()
    book = load_workbook(io.BytesIO(content), data_only=False)
    try:
        if not {'Plate Map', 'Channels'} <= set(book.sheetnames):
            raise ValueError(f'{path.name}: required sheets are Plate Map and Channels; use the current FIA template')
        plate, channels = book['Plate Map'], book['Channels']
        if ([plate.cell(1, c).value for c in range(2, 14)] != list(range(1, 13)) or
                [plate.cell(r, 1).value for r in range(2, 10)] != list('ABCDEFGH')):
            raise ValueError('Plate Map: expected the 96-well grid at A1:M9')
        if ([channels.cell(5, c).value for c in range(3, 15)] != list(range(1, 13)) or
                [channels.cell(43, c).value for c in range(3, 15)] != list(range(1, 13))):
            raise ValueError('Channels: expected channel grid at C6:N37 and intersections at C44:N51')
        for sheet, regions in ((plate, [(2, 9, 2, 13), (15, 21, 2, 13)]),
                               (channels, [(6, 37, 3, 14), (44, 51, 3, 14)])):
            for area in sheet.merged_cells.ranges:
                if any(area.min_row <= r2 and area.max_row >= r1 and area.min_col <= c2 and area.max_col >= c1
                       for r1, r2, c1, c2 in regions):
                    raise ValueError(f'{sheet.title}: merged input cells are not supported ({area})')
            for rule in sheet.conditional_formatting:
                if any(a.min_row <= r2 and a.max_row >= r1 and a.min_col <= c2 and a.max_col >= c1
                       for a in rule.sqref.ranges for r1, r2, c1, c2 in regions):
                    raise ValueError(f'{sheet.title}: use direct formatting in input cells')
        wells, conditions = {}, {}
        for row_index, letter in enumerate('ABCDEFGH'):
            if channels.cell(6 + 4 * row_index, 1).value != letter:
                raise ValueError(f'Channels: missing row label {letter}')
            if [channels.cell(6 + 4 * row_index + k, 2).value for k in range(4)] != ['C1', 'C2', 'C3', 'C4']:
                raise ValueError(f'Channels: {letter} must have C1, C2, C3, C4 in that order')
            for column in range(1, 13):
                well = f'{letter}{column:02d}'
                cell = plate.cell(row_index + 2, column + 1)
                condition = _literal(cell)
                cells = [channels.cell(6 + 4 * row_index + k, column + 2) for k in range(4)]
                entries = [_literal(c) for c in cells]
                intersection_cell = channels.cell(44 + row_index, column + 2)
                requested = _literal(intersection_cell)
                if not condition:
                    if any(entries) or requested:
                        raise ValueError(f'{well}: channels/intersections specified for an unannotated well')
                    continue
                block, control = _color(cell), bool(cell.font.bold)
                if condition in conditions and conditions[condition] != (block, control):
                    raise ValueError(f'{cell.coordinate}: condition {condition!r} has inconsistent block/control formatting')
                conditions[condition] = block, control
                assignments, seen, absent = [], set(), False
                for index, (entry, source_cell) in enumerate(zip(entries, cells), 1):
                    if not entry:
                        absent = True
                        assignments.append(None)
                        continue
                    if absent:
                        raise ValueError(f'{well}: missing channel before C{index}; use Ignore for existing unused channels')
                    if entry == 'Ignore':
                        assignments.append({'channel': index, 'method': 'Ignore', 'name': '', 'id': ''})
                        continue
                    method, sep, label = entry.partition(':')
                    method, label = method.strip(), label.strip()
                    if not sep or method not in METHODS or not label:
                        raise ValueError(f'Channels!{source_cell.coordinate}: expected Method: marker name, or Ignore')
                    identity = marker_id(method, label)
                    if identity in seen:
                        raise ValueError(f'{well}: duplicate {method}: {label}; use distinct marker names')
                    seen.add(identity)
                    assignments.append({'channel': index, 'method': method, 'name': label, 'id': identity})
                if sum(bool(a and a['method'] == 'Nuclei') for a in assignments) != 1:
                    raise ValueError(f'{well}: exactly one Nuclei channel is required')
                intersections = []
                for value in requested.split(';') if requested else []:
                    value = value.strip()
                    if not re.fullmatch(r'C[1-4](\s*\+\s*C[1-4]){1,2}', value):
                        raise ValueError(f'{well}: invalid intersection {value!r}; use C2+C3; C2+C3+C4')
                    numbers = sorted(int(v.strip()[1:]) for v in value.split('+'))
                    if len(set(numbers)) != len(numbers):
                        raise ValueError(f'{well}: repeated channel in {value}')
                    if any(not assignments[n-1] or assignments[n-1]['method'] != 'Foci' for n in numbers):
                        raise ValueError(f'{well}: all channels in {value} must be assigned Foci')
                    if numbers not in intersections:
                        intersections.append(numbers)
                wells[well] = {'condition': condition, 'block': block, 'control': control,
                               'cell': cell.coordinate, 'channels': assignments, 'intersections': intersections}
        if not wells and not allow_empty:
            raise ValueError('Plate Map has no annotated wells; fill the template before analysis')
        panels, located, blocks = [], set(), set()
        for row in range(15, 22):
            column = 2
            while column <= 13:
                if not _literal(plate.cell(row, column)):
                    column += 1
                    continue
                start, groups, cells = column, [], {}
                block = _color(plate.cell(row, column))
                if block in blocks:
                    raise ValueError('PLOT LAYOUT: each comparison block must occupy one horizontal strip')
                while column <= 13 and _literal(plate.cell(row, column)):
                    cell = plate.cell(row, column)
                    group = _literal(cell)
                    if group not in conditions or conditions[group] != (_color(cell), bool(cell.font.bold)):
                        raise ValueError(f'PLOT LAYOUT!{cell.coordinate}: label/fill/bold must match Plate Map')
                    if _color(cell) != block:
                        raise ValueError(f'PLOT LAYOUT!{cell.coordinate}: leave a blank column between blocks')
                    if group in located:
                        raise ValueError(f'PLOT LAYOUT: repeated condition {group!r}')
                    located.add(group)
                    groups.append(group)
                    cells[group] = cell.coordinate
                    column += 1
                expected = {g for g, (color, _) in conditions.items() if color == block}
                if set(groups) != expected:
                    raise ValueError(f'PLOT LAYOUT: strip at row {row} must contain every condition of its color block')
                blocks.add(block)
                panels.append({'row': row, 'column': start, 'end_column': column - 1,
                               'block': block, 'groups': groups, 'cells': cells})
        if located != set(conditions):
            raise ValueError('PLOT LAYOUT: missing conditions: ' + ', '.join(sorted(set(conditions) - located)))
        return PlateDesign(path, content, wells, panels)
    finally:
        book.close()
