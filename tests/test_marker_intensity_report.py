"""Report contracts: provenance, plate design, observation units and literal exports."""

import csv
import json
import os
import subprocess
import sys
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

import collect_marker_intensity_results as collector
import marker_intensity_report as report
import marker_report_data as data_tools
import numpy as np
import pytest
from marker_report_plots import plot_rows
from marker_report_statistics import calculate_statistics, holm_adjust, observations
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from scipy.stats import ttest_ind
from test_collect_marker_intensity_results import (
    collect,
    create_intensity,
    create_morphology,
    sheet_rows,
)

MARKER = 'Foci_1_Channel_2'
RAW = MARKER + '_Marker_RawIntDen'


def plate(path, groups=None):
    groups = groups or {'A02': ('Control', True), 'A03': ('Treatment', False)}
    book = Workbook()
    sheet = book.active
    sheet.title = 'Plate Map'
    for number in range(1, 13):
        sheet.cell(1, number + 1, number)
    for number, letter in enumerate('ABCDEFGH', 2):
        sheet.cell(number, 1, letter)
    for well, (group, control) in groups.items():
        cell = sheet.cell(ord(well[0]) - ord('A') + 2, int(well[1:]) + 1, group)
        cell.fill = PatternFill('solid', fgColor='ABCDEF')
        cell.font = Font(bold=control)
    book.save(path)
    return path


def collection(tmp_path, empty=False, legacy=False, missing=False):
    morph = create_morphology(tmp_path / 'experiment', empty=empty)
    create_intensity(morph)
    markers = [MARKER]
    if legacy:
        create_intensity(morph, stamp='20260924_110000', legacy=True)
        markers.append('Channel_2')
    if missing:
        markers.append('Foci_2_Channel_3')
    ok, path = collect(morph, markers)
    assert ok
    plate(path / 'arbitrary layout.xlsx')
    return path


def annotated():
    """Unequal image populations deliberately distinguish pooled and equal-image means."""
    images, nuclei = [], []
    design = [{'Well': well, 'Group': group} for well, group in
              [('A02', 'Control'), ('A03', 'Control'), ('A04', 'Treatment'), ('A05', 'Treatment')]]
    populations = [('A02', [1] * 100), ('A02', [11]), ('A03', [5, 7]), ('A04', [20, 40]), ('A05', [60, 90])]
    specs = data_tools.metric_specs([MARKER])
    for index, (well, values) in enumerate(populations):
        identity = {'Well': well, 'Group': 'Control' if well in ('A02', 'A03') else 'Treatment',
                    'Mask_name': f'mask_{index}', 'Image_name': f'field{index}_Well{well}.nd2'}
        image = dict(identity, Non_border_nuclei_count=len(values))
        for _, _, field, _ in specs[1:]:
            for stat, value in zip(collector.STATS, (np.mean(values), np.median(values),
                                                    np.percentile(values, 75) - np.percentile(values, 25))):
                image[field + '_' + stat] = float(value)
        images.append(image)
        for nucleus_id, value in enumerate(values, 1):
            nuclei.append(dict(identity, Nucleus_ID=nucleus_id, **{field: value for _, _, field, _ in specs[1:]}))
    return data_tools.aggregate({'images': images, 'nuclei': nuclei, 'design': design,
                                 'groups': ['Control', 'Treatment'], 'markers': [MARKER],
                                 'blocks': {'one color': {'Control': True, 'Treatment': False}},
                                 'stats_unit': 'nucleus', 'run_id': 'Final_Nuclei_Mask_20260924_090000'})


def test_archive_validation_and_legacy_selection(tmp_path):
    path = collection(tmp_path, legacy=True)
    data = data_tools.load_collection(path)
    markers, notes = data_tools.marker_catalog(data)
    assert markers == [MARKER] and 'Channel_2 excluded' in notes[0]
    data_tools.annotate(data, data_tools.find_template(path), markers, 'nucleus')
    data_tools.aggregate(data)
    assert len(data['nuclei']) == 2 and data['images'][0][data_tools.COUNT] == 2
    assert all(row['Touches_border'] is False for row in data['nuclei'])
    assert data['nuclei'][0]['Well'] == 'A02'
    assert not any(key.startswith('Channel_2_') for key in data['nuclei'][0])
    assert len(plot_rows(data)) == 3
    # Archived sources suffice even after the original raw image is removed.
    (path.parent / 'field, WellA2.nd2').unlink()
    assert len(data_tools.load_collection(path)['nuclei']) == 2


def test_archived_source_change_rejected(tmp_path):
    path = collection(tmp_path)
    with (path / 'Nuclei_Morphology.csv').open('a') as handle:
        handle.write('\n')
    with pytest.raises(data_tools.ValidationError, match='changed since collection'):
        data_tools.load_collection(path)


def test_combined_csv_workbook_disagreement_rejected(tmp_path):
    path = collection(tmp_path)
    csv_path = path / 'FIA_Marker_Intensity_Nuclei.csv'
    header, rows = collector.csv_snapshot(csv_path, {}, [])
    rows[0]['Area_px2'] = '999'
    collector.write_csv(csv_path, header, rows)
    with pytest.raises(data_tools.ValidationError, match='differs from CSV'):
        data_tools.load_collection(path)


def test_plate_discovery_ignores_hidden_lock_and_analytical_workbooks(tmp_path):
    path = collection(tmp_path)
    for name in ('._hidden.xlsx', '~$locked.xlsx'):
        plate(path / name)
    (path / 'bad.xlsx').write_text('not a workbook')
    assert data_tools.find_template(path).name == 'arbitrary layout.xlsx'
    plate(path / 'second layout.xlsx')
    with pytest.raises(data_tools.ValidationError, match='found 2'):
        data_tools.find_template(path)


@pytest.mark.parametrize('groups,reason', [
    ({'A03': ('Treatment', False)}, 'Unannotated imaged well'),
    ({'A02': ('Control', False), 'A03': ('Treatment', False)}, 'one control'),
    ({'A02': ('Control', True), 'A03': ('Control', False)}, 'Inconsistent'),
])
def test_invalid_plate_design_rejected(tmp_path, groups, reason):
    path = collection(tmp_path)
    mapping = plate(tmp_path / 'invalid.xlsx', groups)
    with pytest.raises(data_tools.ValidationError, match=reason):
        data_tools.annotate(data_tools.load_collection(path), mapping, [MARKER], 'nucleus')


def test_descriptive_template_needs_no_color_or_control(tmp_path):
    path = collection(tmp_path)
    template = path / 'arbitrary layout.xlsx'
    book = load_workbook(template)
    for cell in (book.active['C2'], book.active['D2']):
        cell.fill = PatternFill()
        cell.font = Font()
    book.save(template)
    data = data_tools.annotate(data_tools.load_collection(path), template, [MARKER], None)
    calculate_statistics(data)
    assert data['statistics'] == []
    with pytest.raises(data_tools.ValidationError, match='solid fill'):
        data_tools.annotate(data_tools.load_collection(path), template, [MARKER], 'well')


def test_well_filename_consistency_and_marker_identity(tmp_path):
    with pytest.raises(data_tools.ValidationError, match='Well/Point'):
        data_tools.image_well({'Well': 'A02', 'Image_name': 'WellA2_PointA3.nd2'})
    path = collection(tmp_path, legacy=True)
    data = data_tools.load_collection(path)
    with pytest.raises(data_tools.ValidationError, match='catalog'):
        data_tools.annotate(data, data_tools.find_template(path), ['Channel_2'], None)
    data['markers'].append('Foci_2_Channel_2')
    with pytest.raises(data_tools.ValidationError, match='same channel'):
        data_tools.annotate(data, data_tools.find_template(path), [MARKER, 'Foci_2_Channel_2'], None)


def test_hierarchical_well_means_and_nucleus_plot_population():
    data = annotated()
    expected_points = plot_rows(data)
    control_nuclei, unit = observations(data, RAW, 'Control')
    assert unit == 'nucleus' and len(control_nuclei) == 103
    data['stats_unit'] = 'well'
    wells, unit = observations(data, RAW, 'Control')
    assert unit == 'well' and wells == [6.0, 6.0]
    assert np.mean(control_nuclei) != np.mean(wells)
    assert plot_rows(data) == expected_points
    assert sum(row['Observation'] == 'nucleus' for row in expected_points) == 107
    assert sum(row['Observation'] == 'image' for row in expected_points) == 5
    counts, _ = observations(data, data_tools.COUNT, 'Control')
    assert counts == [50.5, 2.0]


@pytest.mark.parametrize('unit', ['nucleus', 'well'])
def test_welch_direction_interval_family_and_count_exception(unit):
    data = annotated()
    data['stats_unit'] = unit
    calculate_statistics(data)
    assert len(data['statistics']) == 19
    row = next(row for row in data['statistics'] if row['Metric'] == RAW)
    control, _ = observations(data, RAW, 'Control')
    treatment, _ = observations(data, RAW, 'Treatment')
    expected = ttest_ind(treatment, control, equal_var=False)
    interval = expected.confidence_interval()
    assert row['P_raw'] == pytest.approx(expected.pvalue)
    assert row['CI95_low'] == pytest.approx(interval.low)
    assert row['CI95_high'] == pytest.approx(interval.high)
    assert row['Difference_treatment_minus_control'] > 0
    assert row['Family_size_planned'] == 19
    assert row['N_control'] == (103 if unit == 'nucleus' else 2)
    count = data['statistics'][0]
    assert count['Effective_unit'] == ('image' if unit == 'nucleus' else 'well')
    assert count['N_control'] == (3 if unit == 'nucleus' else 2)


def test_holm_uses_planned_family_including_unavailable_tests():
    assert holm_adjust([0.04, 0.01, 0.03], 5) == pytest.approx([0.12, 0.05, 0.12])
    data = annotated()
    data['groups'].append('No observations')
    data['blocks']['one color']['No observations'] = False
    calculate_statistics(data)
    assert len(data['statistics']) == 38
    missing = [r for r in data['statistics'] if r['Treatment'] == 'No observations']
    assert all(r['N_treatment'] == 0 and r['P_raw'] is None and r['P_Holm'] is None for r in missing)
    assert all(r['Family_size_planned'] == 38 for r in data['statistics'])


def test_one_well_no_test_and_missing_marker_never_zero(tmp_path):
    path = collection(tmp_path, missing=True)
    data = data_tools.annotate(data_tools.load_collection(path), data_tools.find_template(path),
                              [MARKER, 'Foci_2_Channel_3'], 'well')
    data_tools.aggregate(data)
    calculate_statistics(data)
    assert len(data['statistics']) == 25
    assert all(r['Status'] == 'NOT_TESTED' and r['P_raw'] is None for r in data['statistics'])
    missing = [r for r in data['well_values'] if r['Marker'] == 'Foci_2_Channel_3']
    assert missing and all(r['Value'] is None for r in missing)
    assert len(plot_rows(data)) == 3


def test_zero_variance_not_tested_and_mean_reconciliation():
    data = annotated()
    for row in data['nuclei']:
        row[RAW] = 10
    calculate_statistics(data)
    statistic = next(r for r in data['statistics'] if r['Metric'] == RAW)
    assert statistic['Status'] == 'NOT_TESTED' and 'zero variance' in statistic['Reason']
    with pytest.raises(data_tools.ValidationError, match='summary disagrees'):
        data_tools.aggregate(data)


def test_complete_report_keeps_inputs_and_has_embedded_plots(tmp_path):
    path = collection(tmp_path, legacy=True)
    before = {p: p.read_bytes() for p in path.iterdir() if p.is_file()}
    ok, output = report.create_report(path, path.parent, markers='all', stats_unit='nucleus')
    assert ok
    status = json.loads((output / 'report_status.json').read_text())
    assert status['Status'] == 'SUCCESS' and status['Non_border_nuclei'] == 2
    assert status['Planned_comparisons'] == 19
    assert len(sheet_rows(output / report.WORKBOOK, 'Nuclei')) == 2
    assert len(sheet_rows(output / report.WORKBOOK, 'Statistics')) == 19
    by_condition = sheet_rows(output / report.MORPHOLOGY_WORKBOOK, report.MORPHOLOGY_SHEETS[0])
    comparisons = sheet_rows(output / report.MORPHOLOGY_WORKBOOK, report.MORPHOLOGY_SHEETS[1])
    assert len(by_condition) == 12 and len(comparisons) == 12
    assert {r['Metric'] for r in by_condition} == set(collector.MORPH_METRICS)
    assert by_condition == sheet_rows(output / report.WORKBOOK, report.MORPHOLOGY_SHEETS[0])
    assert comparisons == sheet_rows(output / report.WORKBOOK, report.MORPHOLOGY_SHEETS[1])
    for title in report.MORPHOLOGY_SHEETS:
        assert (output / (title + '.csv')).is_file()
    with ZipFile(output / report.MORPHOLOGY_WORKBOOK) as archive:
        assert not any(name.startswith('xl/media/') for name in archive.namelist())
    assert len(list((output / 'Plots').glob('*.png'))) == 2
    assert [r['Points'] for r in sheet_rows(output / report.WORKBOOK, 'Plot_Info')] == [1, 2]
    with ZipFile(output / report.WORKBOOK) as archive:
        assert len([name for name in archive.namelist() if name.startswith('xl/media/')]) == 2
    assert all(p.read_bytes() == content for p, content in before.items())
    assert report.create_output(path.parent) != output
    assert (output / 'Inputs' / 'Collection' / 'Nuclei_Morphology.csv').read_bytes() == before[path / 'Nuclei_Morphology.csv']


def test_empty_nuclei_are_zero_counts_not_zero_intensities(tmp_path):
    path = collection(tmp_path, empty=True)
    ok, output = report.create_report(path, path.parent, markers='all', stats_unit='well')
    assert ok
    assert sheet_rows(output / report.WORKBOOK, 'Nuclei') == []
    summary = sheet_rows(output / report.MORPHOLOGY_WORKBOOK, report.MORPHOLOGY_SHEETS[0])
    assert len(summary) == 12
    assert all(r['Control: N'] == 0 and r['Control: Mean'] is None and r['Control: SD'] is None for r in summary)
    comparisons = sheet_rows(output / report.MORPHOLOGY_WORKBOOK, report.MORPHOLOGY_SHEETS[1])
    assert all(r['Status'] == 'NOT_TESTED' and r['P_Holm'] is None for r in comparisons)
    points = sheet_rows(output / report.WORKBOOK, 'Plot_Data')
    assert len(points) == 1 and points[0]['Value'] == 0 and points[0]['Observation'] == 'image'
    with (output / 'Nuclei.csv').open(encoding='utf-8-sig') as handle:
        assert 'Area_px2' in next(csv.reader(handle))


def test_changed_input_cannot_publish_success(tmp_path, monkeypatch):
    path = collection(tmp_path)
    original = report.write_workbook

    def mutate(*args):
        original(*args)
        with (path / 'Nuclei_Images.csv').open('a') as handle:
            handle.write('\n')

    monkeypatch.setattr(report, 'write_workbook', mutate)
    ok, output = report.create_report(path, path.parent, markers='all')
    assert not ok
    status = json.loads((output / 'report_status.json').read_text())
    assert status['Status'] == 'FAILED' and 'Input changed' in status['Error']
    assert (output / 'Report_Diagnostics.xlsx').is_file()


def test_formula_like_group_is_literal_in_workbook(tmp_path):
    path = collection(tmp_path)
    template = path / 'arbitrary layout.xlsx'
    book = load_workbook(template)
    book.active['C2'].value = '=Control'
    book.active['C2'].data_type = 's'
    book.save(template)
    ok, output = report.create_report(path, path.parent, markers='none')
    assert ok
    book = load_workbook(output / report.WORKBOOK, read_only=True, data_only=False)
    try:
        assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
        assert '=Control' in [r['Group'] for r in sheet_rows(output / report.WORKBOOK, 'Nuclei')]
    finally:
        book.close()
    assert sheet_rows(output / report.MORPHOLOGY_WORKBOOK, report.MORPHOLOGY_SHEETS[1]) == []
    summary = sheet_rows(output / report.MORPHOLOGY_WORKBOOK, report.MORPHOLOGY_SHEETS[0])
    assert summary[0]['=Control: N'] == 2


def test_cli_noninteractive_and_no_imaging_imports(tmp_path):
    path = collection(tmp_path)
    manifest = tmp_path / 'input_paths.json'
    manifest.write_text(json.dumps({'paths_to_files': [str(path.parent)]}))
    script = Path(report.__file__).resolve()
    result = subprocess.run([sys.executable, str(script), '-i', str(manifest), '--all-experiments',
                             '--collections', 'latest', '--markers', MARKER, '--stats-unit', 'well'],
                            text=True, capture_output=True, timeout=45, check=False)
    assert result.returncode == 0, result.stdout + result.stderr
    assert 'SUCCESS:' in result.stdout
    check = subprocess.run([sys.executable, '-c',
                            ('import sys, marker_intensity_report; '
                             'assert not {"imagej", "scyjava", "stardist", "tensorflow"} & set(sys.modules)')],
                           env={**os.environ, 'PYTHONPATH': str(script.parent)}, text=True,
                           capture_output=True, timeout=30, check=False)
    assert check.returncode == 0, check.stderr


def test_discovery_ignores_unfinished_and_hidden_collections(tmp_path):
    path = collection(tmp_path)
    (path.parent / 'FIA_Marker_Intensity_Combined_Results_20990101_010101').mkdir()
    (path.parent / ('._' + path.name)).mkdir()
    selected, missing = report.choose_collections([path.parent], 'latest')
    assert selected == [path] and missing == []


@pytest.mark.parametrize('unit', ['nucleus', 'well', None])
def test_morphology_summary_matches_analysis_unit_and_preserves_statistics(unit):
    data = annotated()
    data['stats_unit'] = unit
    calculate_statistics(data)
    original = deepcopy(data)
    plots = plot_rows(data)
    tables = report.morphology_tables(data)
    columns, summary = tables[report.MORPHOLOGY_SHEETS[0]]
    assert len(summary) == 12 and len(columns) == 13
    assert {r['Metric'] for r in summary} == set(collector.MORPH_METRICS)
    circularity = next(r for r in summary if r['Metric'] == 'Circularity')
    expected = [6.0, 6.0] if unit == 'well' else [1] * 100 + [11, 5, 7]
    assert circularity['Observation_unit'] == (unit or 'nucleus')
    assert circularity['Control: N'] == len(expected)
    assert circularity['Control: Mean'] == pytest.approx(np.mean(expected))
    assert circularity['Control: SD'] == pytest.approx(np.std(expected, ddof=1))
    assert circularity['Control: Median'] == pytest.approx(np.median(expected))
    assert circularity['Control: IQR'] == pytest.approx(np.percentile(expected, 75) - np.percentile(expected, 25))
    comparison_columns, comparisons = tables[report.MORPHOLOGY_SHEETS[1]]
    expected_tests = [r for r in data['statistics'] if r['Category'] == 'Morphology']
    assert comparisons == [{key: row[key] for key in comparison_columns} for row in expected_tests]
    if unit is not None:
        assert all(row['Family_size_planned'] == 19 for row in comparisons)
    assert data == original and plot_rows(data) == plots


def test_morphology_summary_keeps_unobserved_groups_and_metric_specific_missing_values():
    data = annotated()
    data['groups'].append('No observations')
    # One valid circularity value remains in Control; other morphology fields stay present.
    control = [r for r in data['nuclei'] if r['Group'] == 'Control']
    for nucleus in control[1:]:
        nucleus['Circularity'] = None
    data['stats_unit'] = None
    calculate_statistics(data)
    _, rows = report.morphology_tables(data)[report.MORPHOLOGY_SHEETS[0]]
    circularity = next(r for r in rows if r['Metric'] == 'Circularity')
    area = next(r for r in rows if r['Metric'] == 'Area_px2')
    assert circularity['Control: N'] == 1 and circularity['Control: Mean'] == 1
    assert circularity['Control: SD'] is None
    assert area['Control: N'] == 103
    assert all(row['No observations: N'] == 0 and row['No observations: Mean'] is None for row in rows)


def test_missing_morphology_workbook_prevents_success(tmp_path, monkeypatch):
    path = collection(tmp_path)
    original = report.write_workbook

    def skip_morphology(output, tables, plots, filename=report.WORKBOOK):
        if filename != report.MORPHOLOGY_WORKBOOK:
            original(output, tables, plots, filename)

    monkeypatch.setattr(report, 'write_workbook', skip_morphology)
    ok, output = report.create_report(path, path.parent, markers='all')
    assert not ok
    status = json.loads((output / 'report_status.json').read_text())
    assert status['Status'] == 'FAILED' and status['Stage'] == 'final verification'
