"""Selection/export tests and optional real ImageJ + Bio-Formats integration.

Set FIA_IMAGEJ_JAR and FIA_BIOFORMATS_JAR to installed JAR paths for native tests.
"""

import csv
import json
import os
from pathlib import Path
from unittest.mock import MagicMock

import nuclear_intensity as app
import nuclear_intensity_imagej as backend
import numpy as np
import pytest
import tifffile
from openpyxl import load_workbook
from PIL import Image


@pytest.fixture
def engine(monkeypatch):
    jars = [os.environ.get('FIA_IMAGEJ_JAR'), os.environ.get('FIA_BIOFORMATS_JAR')]
    if not all(jars) or not all(Path(path).is_file() for path in jars):
        pytest.skip('Set FIA_IMAGEJ_JAR and FIA_BIOFORMATS_JAR for native tests')
    jpype = pytest.importorskip('jpype')
    if not jpype.isJVMStarted():
        jpype.startJVM('-Djava.awt.headless=true', classpath=jars)
    else:
        for jar in jars:
            jpype.addClassPath(jar)
    monkeypatch.setattr(backend, 'jimport', jpype.JClass)
    jpype.JClass('loci.common.DebugTools').setRootLevel('ERROR')
    return backend.ImageJEngine(initialize=False)


def make_experiment(tmp_path, shape=(37, 1031)):
    root = tmp_path / 'experiment, A'
    root.mkdir(exist_ok=True)
    raw = root / 'field, WellA2.ome.tif'
    y, x = shape
    pixels = np.zeros((3, 2, y, x), dtype=np.uint16)
    pixels[:, 0] = 65535  # The unselected channel must never enter measurements.
    pixels[0, 1] = 3
    pixels[1, 1] = 8
    pixels[2, 1] = 5
    pixels[:, 1, 10:13, 12:16] = 0  # Include true zero pixels in the ROI.
    pixels[2, 1, 10, 12] = 40
    tifffile.imwrite(raw, pixels, ome=True, metadata={'axes': 'ZCYX'}, photometric='minisblack')
    run = root / 'foci_assay' / 'Final_Nuclei_Mask_20260923_120000'
    qc = run / 'Morphology_QC'
    qc.mkdir(parents=True)
    stem = raw.stem + app.MASK_SUFFIX
    labels = np.zeros(shape, dtype=np.uint16)
    labels[0:3, 2:5] = 7
    labels[10:13, 12:16] = 19
    labels[20:24, 24:29] = 31
    tifffile.imwrite(run / (stem + '.tif'), np.where(labels, 0, 255).astype(np.uint8))
    tifffile.imwrite(qc / (stem + '_ids.tif'), labels)
    metadata = {'schema_version': 1, 'status': 'complete', 'morphology_status': 'complete',
                'processed_files': [raw.stem + '_nuclei_projection_StarDist_processed.tif'],
                'skipped_files': [], 'particle_size_pixels_squared': 2,
                'stardist_folder': str(run.parent / 'StarDist_original')}
    (run / 'nuclei_run.json').write_text(json.dumps(metadata))
    (root / ('._' + raw.name)).write_bytes(b'AppleDouble')
    (run / ('._' + stem + '.tif')).write_bytes(b'AppleDouble')
    manifest = tmp_path / 'input_paths.json'
    manifest.write_text(json.dumps({'paths_to_files': [str(root)]}))
    return manifest, raw, run, pixels, labels


def test_selection_all_several_deduplicate_and_cancel():
    assert app.parse_selection('all', 3) == [0, 1, 2]
    assert app.parse_selection('3,1,3', 3) == [0, 2]
    for bad in ('', '0', '4', '1,x'):
        with pytest.raises(ValueError):
            app.parse_selection(bad, 3)
    with pytest.raises(app.Cancelled):
        app.parse_selection('q', 3)


def test_discovery_hidden_files_missing_paths_and_duplicate_folders(tmp_path):
    root = tmp_path / 'experiment'
    root.mkdir()
    for name in ('image.tif', 'image.nd2', '._image.tif', '.hidden.tif', 'notes.csv'):
        (root / name).touch()
    (root / 'directory.tif').mkdir()
    manifest = tmp_path / 'input.json'
    manifest.write_text(json.dumps({'paths_to_files': [str(root), str(root), str(root / 'missing')]}))
    results = app.discover(manifest)
    assert len(results) == 1
    assert len(results[0]['raw']) == 2
    assert results[0]['runs'] == []


def test_select_latest_all_manual_excludes_incompatible(monkeypatch):
    records = [{'dataset': Path(root), 'path': Path(name), 'eligible': valid, 'metadata': {}}
               for root, name, valid in [('A', '20260101', True), ('A', '20260103', False),
                                         ('A', '20260102', True), ('B', '20260101', True)]]
    monkeypatch.setattr('builtins.input', lambda _: '1')
    assert app.select_runs(records) == records[2:]
    monkeypatch.setattr('builtins.input', lambda _: '2')
    assert app.select_runs(records) == [records[i] for i in (0, 2, 3)]
    answers = iter(['3', '1,3'])
    monkeypatch.setattr('builtins.input', lambda _, answers=answers: next(answers))
    assert app.select_runs(records) == [records[0], records[3]]


def test_native_max_intensity_ids_zeros_and_border_exclusion(tmp_path, engine):
    manifest, raw, run, pixels, labels = make_experiment(tmp_path)
    experiment = app.discover(manifest)[0]
    record = app.inspect_run(experiment, run, 'tiff-stack', 2, engine, {})
    assert record['eligible'], record['errors']
    pair = record['pairs'][0]
    marker = engine.marker(raw, 'tiff-stack', 2, pair['info'])
    output = tmp_path / 'measurements'
    try:
        rows, counts = engine.measure(marker, labels, output)
    finally:
        marker.close()
    assert counts == {'Nuclei_count_total': 3, 'Border_nuclei_count': 1,
                          'Non_border_nuclei_count': 2, 'Nuclei_measured_count': 2, 'Failed_nuclei_count': 0}
    assert [row['Nucleus_ID'] for row in rows] == [19, 31]
    maximum = pixels[:, 1].max(axis=0)
    for row in rows:
        expected = maximum[labels == row['Nucleus_ID']]
        assert row['Area_px2'] == expected.size
        assert row['Marker_Mean'] == pytest.approx(expected.mean())
        assert row['Marker_Median'] == pytest.approx(np.median(expected))
        assert row['Marker_StdDev'] == pytest.approx(expected.std(ddof=1))
        assert row['Marker_Min'] == expected.min()
        assert row['Marker_Max'] == expected.max()
        assert row['Marker_RawIntDen'] == float(expected.sum())
        saved = tifffile.imread(output / row['Nucleus_mask'])
        assert saved.shape == labels.shape
        assert np.array_equal(saved > 0, labels == row['Nucleus_ID'])
        assert (output / row['ROI']).read_bytes().startswith(b'Iout')
    assert not (output / 'nucleus_7.roi').exists()
    saved_marker = tifffile.imread(output / 'marker.tif')
    assert saved_marker.dtype == np.float32
    assert np.array_equal(saved_marker, maximum)
    assert Image.open(output / 'numbered_nuclei.png').size == (labels.shape[1], labels.shape[0])
    assert np.array_equal(tifffile.imread(raw), pixels)


def test_native_2d_float_channel_preserves_values(tmp_path, engine):
    data = np.arange(2 * 13 * 21, dtype=np.float32).reshape(2, 13, 21) / 4 - 10
    path = tmp_path / 'float.ome.tif'
    tifffile.imwrite(path, data, ome=True, metadata={'axes': 'CYX'}, photometric='minisblack')
    info = engine.inspect(path, 'tiff-2d', 2)
    marker = engine.marker(path, 'tiff-2d', 2, info)
    try:
        measured = np.asarray(marker.getProcessor().getPixels()).reshape(13, 21)
        assert np.array_equal(measured, data[1])
    finally:
        marker.close()
    with pytest.raises(ValueError, match='Channel'):
        engine.inspect(path, 'tiff-2d', 3)


def test_native_incompatible_dimensions_incomplete_and_ambiguous(tmp_path, engine):
    manifest, raw, run, _pixels, labels = make_experiment(tmp_path)
    experiment = app.discover(manifest)[0]
    assert not app.inspect_run(experiment, run, 'tiff-2d', 2, engine, {})['eligible']
    id_map = next((run / 'Morphology_QC').glob('*.tif'))
    tifffile.imwrite(id_map, labels[:, :-1])
    record = app.inspect_run(experiment, run, 'tiff-stack', 2, engine, {})
    assert not record['eligible'] and 'dimensions differ' in str(record['errors'])
    tifffile.imwrite(id_map, labels)
    copy = raw.with_suffix('.tiff')
    copy.write_bytes(raw.read_bytes())
    experiment = app.discover(manifest)[0]
    record = app.inspect_run(experiment, run, 'tiff-stack', 2, engine, {})
    assert not record['eligible'] and 'found 2' in str(record['errors'])
    metadata = json.loads((run / 'nuclei_run.json').read_text())
    metadata['status'] = 'running'
    (run / 'nuclei_run.json').write_text(json.dumps(metadata))
    incomplete = app.inspect_run(experiment, run, 'tiff-stack', 2, engine, {})
    assert not incomplete['eligible']
    assert incomplete['mask_count'] == 1


def test_native_reject_multi_series_timepoints_rgb_and_nonfinite(tmp_path, engine):
    info = {'Series_count': 2, 'Timepoints': 1, 'RGB': False, 'Channels': 2, 'Z_planes': 3, 'Pixel_type': 'uint16'}
    with pytest.raises(ValueError, match='Multiple series'):
        engine.validate_info(info, 'tiff-stack', 1)
    info.update(Series_count=1, Timepoints=2)
    with pytest.raises(ValueError, match='timepoints'):
        engine.validate_info(info, 'tiff-stack', 1)
    info.update(Timepoints=1, RGB=True)
    with pytest.raises(ValueError, match='RGB'):
        engine.validate_info(info, 'tiff-stack', 1)
    path = tmp_path / 'nan.tif'
    data = np.ones((12, 17), dtype=np.float32)
    data[2, 3] = np.nan
    tifffile.imwrite(path, data)
    info = engine.inspect(path, 'tiff-2d', 1)
    with pytest.raises(ValueError, match='NaN'):
        engine.marker(path, 'tiff-2d', 1, info)


def test_native_end_to_end_exports_quoted_csv_and_rerun(tmp_path, engine, monkeypatch):
    manifest, raw, run, _pixels, _labels = make_experiment(tmp_path)
    before = {p: p.read_bytes() for p in run.rglob('*') if p.is_file()}
    monkeypatch.setattr(app, 'ImageJEngine', lambda: engine)
    for unused in range(2):
        answers = iter(['all', '1'])
        monkeypatch.setattr('builtins.input', lambda _, answers=answers: next(answers))
        assert app.main(manifest, 2, 'tiff-stack') == 0
    outputs = list(run.parent.glob('Nuclear_Intensity_*'))
    assert len(outputs) == 2
    for output in outputs:
        with (output / 'Nuclear_Intensity.csv').open(encoding='utf-8-sig', newline='') as handle:
            rows = list(csv.DictReader(handle))
        assert len(rows) == 2 and rows[0]['Image_name'] == raw.name
        assert rows[0]['Dataset'] == 'experiment, A'
        book = load_workbook(output / 'Nuclear_Intensity.xlsx')
        assert book.sheetnames == ['Nuclei', 'Images', 'Run_Info']
        headers, row = list(book['Images'].values)
        values = dict(zip(headers, row))
        assert values['Non_border_nuclei_count'] == 2
        assert values['Area_px2_Mean'] == 16
        assert 'Condition' not in headers and 'Biological_replicate' not in headers
        assert values['Status'] == 'complete'
        book.close()
        assert json.loads((output / 'intensity_run.json').read_text())['Status'] == 'complete'
    assert all(p.read_bytes() == content for p, content in before.items())
    assert len(list(tmp_path.glob('Nuclear_Intensity_Batch_*'))) == 2


def test_empty_population_has_blank_statistics_and_failure_is_not_zero(tmp_path):
    assert all(value is None for value in app.summarize([]).values())
    app.save_tables(tmp_path, [], [{'Status': 'failed', 'Error': 'unreadable'}], {'Status': 'incomplete'})
    with (tmp_path / 'Nuclear_Intensity_Images.csv').open(encoding='utf-8-sig') as handle:
        rows = list(csv.DictReader(handle))
    assert rows[0]['Nuclei_count_total'] == ''
    assert rows[0]['Marker_Mean_Mean'] == ''
    assert rows[0]['Error'] == 'unreadable'


def test_native_changed_input_after_selection_is_not_measured(tmp_path, engine):
    manifest, raw, run, _, _ = make_experiment(tmp_path)
    record = app.inspect_run(app.discover(manifest)[0], run, 'tiff-stack', 2, engine, {})
    raw.write_bytes(raw.read_bytes() + b'changed')
    output = app.new_output(run.parent, 'Nuclear_Intensity_')
    assert app.analyze_run(record, output, 'tiff-stack', 2, engine, tmp_path / 'batch.json') == 'incomplete'
    with (output / 'Nuclear_Intensity_Images.csv').open(encoding='utf-8-sig') as handle:
        row = next(csv.DictReader(handle))
    assert row['Nuclei_count_total'] == ''
    assert row['Status'] == 'failed' and 'changed' in row['Error']
    assert not list(output.glob('image_*'))


def test_cancel_creates_no_results(tmp_path, monkeypatch):
    manifest, _, run, _, _ = make_experiment(tmp_path)
    monkeypatch.setattr('builtins.input', lambda _: 'q')
    init = MagicMock()
    monkeypatch.setattr(app, 'ImageJEngine', init)
    assert app.main(manifest) == 130
    init.assert_not_called()
    assert not list(run.parent.glob('Nuclear_Intensity_*'))


def test_native_holes_fractional_values_and_roi_roundtrip(tmp_path, engine):
    import jpype
    labels = np.zeros((19, 29), dtype=np.uint16)
    labels[3:10, 4:11] = 42
    labels[5:8, 6:9] = 0
    pixels = (np.arange(labels.size, dtype=np.float32).reshape(labels.shape) % 37) / 7 - 1
    path = tmp_path / 'fractional.tif'
    tifffile.imwrite(path, pixels)
    info = engine.inspect(path, 'tiff-2d', 1)
    marker = engine.marker(path, 'tiff-2d', 1, info)
    try:
        rows, _ = engine.measure(marker, labels, tmp_path / 'results')
    finally:
        marker.close()
    values = pixels[labels == 42]
    assert rows[0]['Marker_Median'] == pytest.approx(np.median(values), abs=1e-6)
    assert rows[0]['Marker_RawIntDen'] == pytest.approx(values.astype(float).sum())
    roi = jpype.JClass('ij.io.RoiDecoder')(str(tmp_path / 'results' / 'nucleus_42.roi')).getRoi()
    restored = np.array([[bool(roi.contains(x, y)) for x in range(29)] for y in range(19)])
    assert np.array_equal(restored, labels == 42)


def test_native_no_nuclei_is_successful_zero_count(tmp_path, engine):
    path = tmp_path / 'empty.tif'
    data = np.ones((14, 17), dtype=np.uint16)
    tifffile.imwrite(path, data)
    info = engine.inspect(path, 'tiff-2d', 1)
    marker = engine.marker(path, 'tiff-2d', 1, info)
    try:
        rows, counts = engine.measure(marker, np.zeros(data.shape, dtype=np.uint16), tmp_path / 'empty')
    finally:
        marker.close()
    assert rows == [] and all(value == 0 for value in counts.values())
    assert (tmp_path / 'empty' / 'numbered_nuclei.png').exists()


def test_stage4_numbered_and_empty_previews_keep_native_dimensions(tmp_path):
    import sys
    module = sys.modules['foci_quantification']
    data = np.zeros((79, 1401), dtype=np.uint16)
    data[10:20, 20:30] = 9
    for index, labels in enumerate(({9: (15, 25)}, {})):
        path = tmp_path / f'preview{index}.png'
        module.save_labeled_image(data, str(path), 'Preview title', labels)
        assert Image.open(path).size == (1401, 79)


@pytest.mark.parametrize('mode', ['2', '3'])
def test_stage1_native_channel_exports(tmp_path, engine, monkeypatch, mode):
    """Use real ImageJ pixels/projections with a headless command bridge."""
    import re
    import sys

    import jpype
    module = sys.modules['select_channels']
    IJ = jpype.JClass('ij.IJ')
    root = tmp_path / 'originals'
    root.mkdir()
    z = 3 if mode == '2' else 1
    data = np.arange(z * 2 * 73 * 1301, dtype=np.uint16).reshape(z, 2, 73, 1301)
    tifffile.imwrite(root / 'input.tif', data, imagej=True,
                     metadata={'axes': 'ZCYX'}, photometric='minisblack')

    class HeadlessCommands:
        current = None

        @staticmethod
        def openImage(path):
            return IJ.openImage(path)

        @staticmethod
        def saveAs(imp, file_format, path):
            IJ.saveAs(imp, file_format, path)

        @staticmethod
        def run(*args):
            if args == ('Close All',):
                return
            imp, command, options = args
            if command == 'Duplicate...':
                channel = int(re.search(r'channels=(\d+)', options)[1])
                HeadlessCommands.current = jpype.JClass('ij.plugin.Duplicator')().run(
                    imp, channel, channel, 1, imp.getNSlices(), 1, 1)
            elif command == '8-bit':
                jpype.JClass('ij.process.ImageConverter')(imp).convertToGray8()
            else:
                raise AssertionError(f'Unexpected command: {command}')

        @staticmethod
        def getImage():
            return HeadlessCommands.current

    monkeypatch.setattr(module, 'initialize_imagej', lambda: None)
    class HeadlessChannels:
        @staticmethod
        def split(imp):
            # ChannelSplitter opens AWT windows in bare ImageJ without Fiji's headless patch.
            return [jpype.JClass('ij.plugin.Duplicator')().run(
                imp, c, c, 1, imp.getNSlices(), 1, 1)
                for c in range(1, imp.getNChannels() + 1)]

    def java_class(cls):
        if cls == 'ij.IJ':
            return HeadlessCommands
        if cls == 'ij.plugin.ChannelSplitter':
            return HeadlessChannels
        return jpype.JClass(cls)

    monkeypatch.setattr(module, 'jimport', java_class)
    answers = iter([mode, '1', '1', '2'])
    monkeypatch.setattr('builtins.input', lambda _: next(answers))
    module.process_image([str(root)])
    outputs = list((root / 'foci_assay').rglob('*.tif'))
    assert len(outputs) == 2
    for path in outputs:
        prepared = tifffile.imread(path)
        assert prepared.shape == (73, 1301)
        assert prepared.dtype == np.uint8
    metadata = (root / 'foci_assay' / 'image_metadata.txt').read_text()
    assert 'Width: 1301' in metadata and 'Height: 73' in metadata
