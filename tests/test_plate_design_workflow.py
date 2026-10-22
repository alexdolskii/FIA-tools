"""Mixed-staining plate integration using real table exporters and report readers."""

import copy
import json
import logging
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest
import tifffile
from openpyxl import load_workbook

import collect_marker_intensity_results as collect
import marker_intensity_report as report
import marker_report_data as data
import nuclear_intensity as intensity
import nuclei_morphology as morphology
import plate_design as plate
import spatial_calibration as spatial
from marker_report_plots import panel_grid, plot_panels, plot_specs
from marker_report_statistics import calculate_statistics
from test_select_channels_logging import runtime

GOLD = Path(__file__).parents[1] / 'templates/FIA_96_well_Golden_Standard.xlsx'
BLANK = GOLD.with_name('FIA_96_well_Template.xlsx')


@pytest.fixture(autouse=True)
def small_excel_previews(monkeypatch):
    # Vector PDF output is unchanged; integration tests need only small embedded previews.
    import marker_report_plots
    import pfak_plots
    monkeypatch.setattr(marker_report_plots, 'PNG_DPI', 80)
    monkeypatch.setattr(pfak_plots, 'PNG_DPI', 80)


def scheme(root):
    root.mkdir(parents=True, exist_ok=True)
    target = root / 'FIA_scheme.xlsx'
    target.write_bytes(GOLD.read_bytes())
    return plate.load_design(root)


def dataset(root, wells=('A01', 'C01', 'C02', 'H04', 'D01', 'H07')):
    design = scheme(root)
    run = root / 'fia_assay/Final_Nuclei_Mask_20261022_100000'
    (run / 'Morphology_QC').mkdir(parents=True)
    export = morphology.NucleiMorphologyExport(run, run.parent / 'StarDist_source', 2, 'test')
    ids = np.zeros((20, 30), dtype=np.uint16)
    ids[0, 20:23] = 7
    ids[2:5, 2:6] = 2
    ids[8:12, 12:17] = 5
    processed = []
    for well in wells:
        raw = root / f'field_Well{well}_Point{well}_0001.nd2'
        raw.write_bytes(b'synthetic metadata-only source ' + well.encode())
        source = raw.stem + '_nuclei_projection_StarDist_processed.tif'
        mask = Path(source).stem + '_processed.tif'
        processed.append(source)
        tifffile.imwrite(run / mask, np.where(ids, 0, 255).astype(np.uint8))
        tifffile.imwrite(run / 'Morphology_QC' / (Path(mask).stem + '_ids.tif'), ids)
        identity = export.identifiers(source, mask)
        rows = []
        for label, area in [(2, 12), (5, 20)]:
            values = {key: .5 for key in morphology.METRICS}
            rows.append({**identity, **values, 'Nucleus_ID': label, 'Area_px2': area,
                         'Touches_border': False, 'Orientation_deg': 0,
                         'Centroid_X_px': 10, 'Centroid_Y_px': 10,
                         'Label_map': 'Morphology_QC/' + Path(mask).stem + '_ids.tif',
                         'Numbered_image': ''})
        export.nuclei.extend(rows)
        image = {**identity, 'Status': 'complete', 'Error': '', 'Nuclei_count_total': 3,
                 'Non_border_nuclei_count': 2, 'Border_nuclei_count': 1}
        for metric in morphology.METRICS:
            values = [r[metric] for r in rows]
            for name, value in zip(collect.STATS, [np.mean(values), np.median(values), np.percentile(values, 75)-np.percentile(values, 25)]):
                image[metric + '_' + name] = float(value)
        export.images.append(image)
    assert export.save() == 'complete'
    (run / 'nuclei_run.json').write_text(json.dumps({
        'schema_version': 1, 'status': 'complete', 'morphology_status': 'complete',
        'processed_files': processed, 'skipped_files': [], 'particle_size_pixels_squared': 2}))
    plate.save_snapshot(run, design)
    return design, run, ids


class IntensityEngine:
    version = bioformats_version = 'test'

    def __init__(self, design, ids):
        self.design, self.ids, self.calls = design, ids, []

    def inspect(self, source, mode, channel):
        well = plate.well_from_name(source)
        self.calls.append((well, channel))
        return {**spatial.uncalibrated(), 'Channels': sum(x is not None for x in self.design.wells[well]['channels']),
                'Width_px': 30, 'Height_px': 20, 'Z_planes': 3, 'Z_step_um': None, 'Pixel_type': 'uint16'}

    def labels(self, *args):
        return self.ids

    def populations(self, labels):
        return [2, 5, 7], [7], [2, 5]

    def marker(self, source, mode, channel, info):
        self.calls.append(('MEASURE_' + plate.well_from_name(source), channel))
        return SimpleNamespace(channel=channel, close=lambda: None)

    def measure(self, marker, labels, folder):
        rows = []
        for index, (label, area) in enumerate([(2, 12), (5, 20)]):
            rows.append({**{key: float(index * marker.channel) for key in collect.INTENSITY_METRICS},
                         'Nucleus_ID': label, 'Area_px2': area, 'Nucleus_mask': 'mask.tif', 'ROI': 'roi.zip'})
        return rows, {'Nuclei_count_total': 3, 'Border_nuclei_count': 1, 'Non_border_nuclei_count': 2,
                      'Nuclei_measured_count': 2, 'Failed_nuclei_count': 0}


def measured_collection(root):
    design, run, ids = dataset(root)
    raw, _ = plate.sources(root, design)
    experiment = {'root': root, 'raw': raw, 'runs': [run], 'design': design}
    engine = IntensityEngine(design, ids)
    for index, marker in enumerate(design.markers('Nuclear intensity'), 1):
        choice = {'name': marker['id'], 'label': marker['name'], 'channel': None,
                  'folder': root / 'fia_assay/markers' / marker['id']}
        record = intensity.inspect_run(experiment, run, 'nd2', choice, engine, {})
        assert record['eligible'], record['errors']
        output = root / 'fia_assay' / f"Nuclear_Intensity_{marker['id']}_20261022_11000{index}"
        output.mkdir()
        assert intensity._analyze_run(record, output, 'nd2', engine, output/'batch.json', None,
                                      logging.getLogger('test')) == 'complete'
    context = collect.scan_source(root)
    assert len(context['intensity']) == 2
    assert len(context['morphology']) == 1
    success, output = collect.collect_one(context['morphology'][0], context,
                                          sorted(context['markers']), 'test_manifest.json')
    assert success
    return design, run, output, engine


def test_golden_routing_and_layout():
    design = plate.read_design(GOLD)
    assert len(design.wells) == 96
    expected = {'Foci': [42, 39, 27], 'Nuclear intensity': [18, 3], 'Field signal': [18, 3], 'Fiber': [48]}
    for method, counts in expected.items():
        assert [len(m['channels']) for m in design.markers(method)] == counts
    assert design.panels[1]['groups'] == ['Group 2 Treatment 2', 'Group 2 Control', 'Group 2 Treatment 1']
    assert {design.channel(f'F{c:02d}', 'Nuclei') for c in (1, 4, 7, 10)} == {1, 2, 3, 4}
    assert len(design.wells['F01']['intersections']) == 4
    assert len(design.wells['G01']['intersections']) == 1
    assert len(design.wells['G01']['intersections'][0]) == 3
    assert not plate.read_design(BLANK, allow_empty=True).wells
    with pytest.raises(ValueError, match='no annotated wells'):
        plate.read_design(BLANK)


@pytest.mark.parametrize('sheet,cell,value,reason', [
    ('Channels', 'C6', 'Ignore', 'exactly one Nuclei'),
    ('Channels', 'C7', None, 'missing channel'),
    ('Channels', 'C8', 'Foci: X', 'must be assigned Foci'),
    ('Channels', 'C48', 'C1+C2', 'must be assigned Foci'),
    ('Channels', 'C48', 'C2+C2', 'repeated channel'),
    ('Plate Map', 'B15', 'Unknown', 'label/fill/bold'),
])
def test_invalid_assignments_are_actionable(tmp_path, sheet, cell, value, reason):
    path = tmp_path/'FIA_test.xlsx'
    book = load_workbook(GOLD)
    if cell == 'C8':
        book['Channels']['C44'] = 'C2+C3'
    book[sheet][cell] = value
    book.save(path)
    with pytest.raises(ValueError, match=reason):
        plate.read_design(path)


def test_same_marker_changes_channel_and_nothing_else_is_measured(tmp_path):
    design, run, output, engine = measured_collection(tmp_path/'experiment')
    measured = [x for x in engine.calls if x[0].startswith('MEASURE_')]
    assert measured == [('MEASURE_C01', 3), ('MEASURE_C02', 3), ('MEASURE_H04', 2), ('MEASURE_H04', 3)]
    combined = data.load_collection(output)
    assert len(combined['images']) == 6 and len(combined['nuclei']) == 12
    first, second = [m['id'] for m in design.markers('Nuclear intensity')]
    by_well = {r['Well']: r for r in combined['images']}
    assert by_well['A01'][first+'_Availability'] == 'NOT_APPLICABLE'
    assert by_well['C01'][second+'_Availability'] == 'NOT_APPLICABLE'
    assert by_well['H04'][first+'_Source_channel'] == 2
    assert by_well['H04'][second+'_Source_channel'] == 3
    assert [r[first+'_Marker_RawIntDen'] for r in combined['nuclei'] if r['Well'] == 'C01'] == [0, 3]
    assert all(r[first+'_Marker_RawIntDen'] is None for r in combined['nuclei'] if r['Well'] == 'A01')


def test_marker_report_and_nuclei_report_have_separate_populations(tmp_path):
    import nuclei_report
    design, run, collection, _ = measured_collection(tmp_path/'experiment')
    marker = design.markers('Nuclear intensity')[0]['id']
    ok, output = report.create_report(collection, run.parent.parent, markers=[marker], stats_unit='nucleus')
    assert ok
    plots = list((output/'Plots').glob('*.pdf'))
    assert len(plots) == 1 and 'Integrated_density' in plots[0].name
    assert not list(output.glob('Nuclei_Morphology_Summary.xlsx'))
    images = collect.csv_snapshot(output/'Images.csv', {}, [])[1]
    assert {r['Well'] for r in images} == {'C01', 'C02', 'H04'}
    stats = collect.csv_snapshot(output/'Statistics.csv', {}, [])[1]
    assert {r['Category'] for r in stats} == {'Intensity'}
    assert {r['Marker'] for r in stats} == {marker}
    style = json.loads((output/'plot_style.json').read_text())
    assert style['panels'][1]['groups'] == design.panels[1]['groups']
    assert panel_grid(style['panels'])[0] == 2
    ok, nuclear_output = nuclei_report.create_report(run, run.parent.parent)
    assert ok
    assert len(list((nuclear_output/'Plots').glob('*.pdf'))) == 5
    assert len(collect.csv_snapshot(nuclear_output/'Images.csv', {}, [])[1]) == 6


def test_scope_provenance_allows_presentation_changes_only(tmp_path):
    design = scheme(tmp_path/'experiment')
    destination = tmp_path/'snapshot'
    plate.save_snapshot(destination, design)
    book = load_workbook(design.path)
    # Swap two whole condition cells, preserving their formatting.
    sheet = book['Plate Map']
    a, b = sheet['C15'], sheet['D15']
    a.value, b.value = b.value, a.value
    a._style, b._style = copy.copy(b._style), copy.copy(a._style)
    book.save(design.path)
    changed = plate.load_design(design.path.parent)
    plate.validate_snapshot(destination, changed)
    assert design.workbook_sha256 != changed.workbook_sha256
    book['Channels']['C6'], book['Channels']['C7'] = 'Ignore', 'Nuclei: Nuclear stain'
    book.save(design.path)
    with pytest.raises(ValueError, match='assignments changed'):
        plate.validate_snapshot(destination, plate.load_design(design.path.parent))


def test_preflight_checks_all_schemes_before_processing(tmp_path):
    first = tmp_path/'first'
    scheme(first)
    second = tmp_path/'second'
    second.mkdir()
    manifest = tmp_path/'input.json'
    manifest.write_text(json.dumps({'paths_to_files': [str(first), str(second)]}))
    with pytest.raises(ValueError, match='second'):
        plate.preflight(manifest)
    assert not (first/'fia_assay').exists()


def test_ignore_count_and_source_identity(tmp_path):
    design = scheme(tmp_path/'experiment')
    design.validate_channels('A01', 4)
    design.validate_channels('B01', 3)
    with pytest.raises(ValueError, match='declares 3 existing channels'):
        design.validate_channels('B01', 4)
    assert plate.well_from_name('x_WellB1_PointB01_001.nd2') == 'B01'
    with pytest.raises(ValueError, match='disagree'):
        plate.well_from_name('x_WellB1_PointB02_001.nd2')


def test_preparation_uses_per_well_channel_without_assignment_prompts(tmp_path, monkeypatch, runtime):
    import sys
    app = sys.modules['select_channels']
    design = scheme(tmp_path/'experiment')
    raw = design.path.parent/'field_WellF04.nd2'
    raw.write_bytes(b'test source')
    runtime.openImage.return_value.getDimensions.return_value = (17, 10, 4, 4, 1)
    monkeypatch.setattr('builtins.input', lambda _: pytest.fail('No channel prompt is permitted'))
    assert app.process_image([str(raw.parent)], nuclei_only=True) == ['SUCCESS']
    duplicates = [str(call) for call in runtime.run.call_args_list if 'Duplicate...' in str(call)]
    assert len(duplicates) == 1 and 'channels=2' in duplicates[0]
    folder = raw.parent/'fia_assay/Nuclei'
    assert plate.preparation_files(folder) == [raw.stem+'_nuclei_projection.tif']
    (folder/'obsolete_WellA01_nuclei_projection.tif').write_bytes(b'not in current manifest')
    assert len(plate.preparation_files(folder)) == 1
    raw.write_bytes(b'changed original')
    with pytest.raises(ValueError, match='original changed'):
        plate.preparation_files(folder)


def test_explicit_triple_does_not_generate_pairs(tmp_path):
    import sys
    app = sys.modules['foci_quantification']
    labels = np.zeros((15, 16), dtype=np.uint16)
    labels[2:12, 2:12] = 1
    nucleus = tmp_path/'field_WellF01_nuclei_projection_StarDist_processed_processed.tif'
    tifffile.imwrite(nucleus, labels)
    masks = {}
    for index in range(1, 4):
        path = tmp_path/f'marker{index}.tif'
        mask = np.zeros_like(labels)
        mask[4:7, 4:7] = index
        tifffile.imwrite(path, mask)
        masks[f'M{index}'] = str(path)
    rows = app.process_nuclei_image(str(nucleus), masks, {}, str(tmp_path), True,
                                    requested_intersections=[('M1', 'M2', 'M3')])
    assert rows
    assert 'Foci Count (M1+M2+M3)' in rows[0]
    assert not any('(M1+M2)' in key or '(M2+M3)' in key for key in rows[0])
    assert len(list(tmp_path.glob('intersection_*.tif'))) == 1
    assert rows[0]['Foci Count (M1+M2+M3)'] == 1


def test_field_signal_routes_only_its_wells_and_report_keeps_marker_name(tmp_path):
    import pfak_analysis as analysis
    import pfak_io
    import pfak_report
    from pfak_measurements import threshold_settings
    from table_workflow import TableJournal, TableProgress
    design, run, ids = dataset(tmp_path/'experiment')
    bundle = collect.load_morphology(run)
    bundle.update(root=run.parent.parent, design=design)
    marker = design.markers('Field signal')[0]
    tasks = analysis.image_tasks(bundle['root'], bundle, marker)
    assert {plate.well_from_name(t['Image_name']) for t in tasks} == {'D01', 'H07'}
    engine = IntensityEngine(design, ids)
    engine.projection = lambda source, channel, info: np.full(ids.shape, channel * 10, dtype=np.float32)
    audit = TableJournal('test_field_signal.log')
    audit.register(bundle['root'])
    try:
        with TableProgress() as progress:
            ok, output = analysis.analyze_run(bundle, marker, threshold_settings([5, 35]), False,
                                              engine, audit, progress)
        assert ok
    finally:
        audit.finish()
    results = pfak_io.load_run(output)
    assert [(r['Well'], r['Channel']) for r in results['images']] == [('D01', 3), ('H07', 2)]
    assert results['metadata']['Marker_name'] == 'Marker 5'
    assert all(r['FieldSignal_Positive_pixels'] == ids.size for r in results['images'])
    ok, report_path = pfak_report.create_report(output, bundle['root'], stats_unit='image')
    assert ok
    assert not list((report_path/'Plots').glob('Nuclei_count*'))
    plots = collect.csv_snapshot(report_path/'Plot_Files.csv', {}, [])[1]
    assert plots
    comparisons = collect.csv_snapshot(report_path/'Statistics.csv', {}, [])[1]
    assert {row['Marker'] for row in comparisons} == {marker['id']}


def test_holm_is_separate_by_marker_and_keeps_filtered_planned_tests():
    design = plate.read_design(GOLD)
    first, second = design.markers('Nuclear intensity')
    specs = [('Intensity', m['id'], m['id']+'_x', 'raw') for m in (first, second)]
    d = {'stats_unit': 'nucleus', 'metric_specs': specs, 'images': [], 'nuclei': [],
         'image_values': [], 'well_values': [], 'design': design.report_rows(),
         'eligible_wells': {m['id']: set(m['channels']) for m in (first, second)},
         'blocks': {'test': {'Group 1 Control': True, 'Group 1 Treatment 1': False}},
         'marker_labels': {m['id']: m['name'] for m in (first, second)}}
    calculate_statistics(d)
    by_marker = {row['Marker']: row for row in d['statistics']}
    assert by_marker[first['id']]['Status'] == 'NOT_TESTED'
    assert by_marker[first['id']]['Family_size_planned'] == 1
    assert by_marker[second['id']]['Status'] == 'NOT_APPLICABLE'
    assert by_marker[second['id']]['Family_size_planned'] == 0
