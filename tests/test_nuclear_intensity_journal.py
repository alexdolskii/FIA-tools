"""Exercise real selection, exports and journal I/O with only the ImageJ engine mocked."""

import csv
import json
from pathlib import Path
from unittest.mock import MagicMock

import nuclear_intensity as app
import numpy as np
import plate_design
import pytest
from openpyxl import load_workbook
from plate_helpers import design
from test_nuclear_intensity import save_morphology_tables
from test_nuclear_intensity import two_marker_experiment as make_experiment


@pytest.fixture
def batch_setup(tmp_path, monkeypatch):
    manifest, _, run, _, labels = make_experiment(tmp_path)
    save_morphology_tables(run, np.zeros_like(labels))
    second = run.parent / 'markers' / 'Foci_2_Channel_1'
    second.mkdir()
    (second / 'prepared.tif').touch()
    home = tmp_path / 'home'
    home.mkdir()
    monkeypatch.setattr(Path, 'home', classmethod(lambda cls: home))
    engine = MagicMock(version='test', bioformats_version='test')
    engine.inspect.return_value = {'Width_px': labels.shape[1], 'Height_px': labels.shape[0],
                                   'Z_planes': 3, 'Pixel_type': 'uint16', 'Channels': 3}
    engine.labels.return_value = np.zeros_like(labels)
    engine.populations.return_value = (set(), set(), set())
    engine.measure.return_value = ([], {key: 0 for key in app.COUNTS})
    monkeypatch.setattr(app, 'ImageJEngine', lambda: engine)
    answers = iter(['all'])
    monkeypatch.setattr('builtins.input', lambda _: next(answers))
    return manifest, run, engine, home


def result_journals(run):
    return sorted(run.parent.glob('Nuclear_Intensity_*/batch.json'))


@pytest.mark.parametrize('failure, exit_code, collection_status', [
    (OSError('collection disk unavailable'), 1, 'failed'),
    (KeyboardInterrupt(), 130, 'canceled'),
])
def test_collection_failure_preserves_complete_measurements(batch_setup, monkeypatch, failure, exit_code, collection_status):
    manifest, run, _, _ = batch_setup
    def fail(*_args, **_kwargs):
        raise failure
    monkeypatch.setattr(app.collection, 'collect_one', fail)
    assert app.main(manifest, mode='tiff-stack') == exit_code
    paths = result_journals(run)
    assert len(paths) == 2 and paths[0].read_bytes() == paths[1].read_bytes()
    for path in paths:
        journal = json.loads(path.read_text())
        assert journal['measurement_status'] == 'complete'
        assert journal['collection_status'] == ('incomplete' if collection_status == 'failed' else 'canceled')
        assert journal['collections'][0]['status'] == collection_status
        assert all(r['status'] == 'complete' for r in journal['runs'])
        assert json.loads((path.parent / 'intensity_run.json').read_text())['Status'] == 'complete'
    text = (run.parent / '4_collect_marker_intensity.log').read_text()
    assert 'status=COMPLETE' not in text.split('RUN_FINISHED |')[1]
    assert 'status=COMPLETE' not in (run.parent / '3_nuclei_intensity.log').read_text().split('RUN_FINISHED |')[1]


def test_collection_log_failure_sets_nonzero_exit_and_keeps_results(batch_setup, monkeypatch):
    manifest, run, _, _ = batch_setup
    original = app.collection.TableJournal.finish
    def fail_log(audit):
        original(audit)
        return False
    monkeypatch.setattr(app.collection.TableJournal, 'finish', fail_log)
    assert app.main(manifest, mode='tiff-stack') == 1
    journal = json.loads(result_journals(run)[0].read_text())
    assert journal['measurement_status'] == 'complete' and journal['collection_status'] == 'incomplete'
    assert journal['collections'][0]['status'] == 'complete'
    assert Path(journal['collections'][0]['output']).is_dir()


def test_success_keeps_identical_journals_and_valid_metadata_in_each_result(batch_setup):
    manifest, run, _, home = batch_setup
    assert app.main(manifest, mode='tiff-stack') == 0
    paths = result_journals(run)
    assert len(paths) == 2 and paths[0].read_bytes() == paths[1].read_bytes()
    journal = json.loads(paths[0].read_text())
    assert journal['status'] == 'complete'
    assert len(journal['runs']) == 2 and all(r['status'] == 'complete' for r in journal['runs'])
    for path in paths:
        output = path.parent
        info = json.loads((output / 'intensity_run.json').read_text())
        assert info['Batch_journal'] == str(path)
        with (output / 'Nuclear_Intensity_Run_Info.csv').open(encoding='utf-8-sig') as handle:
            parameters = {row['Parameter']: row['Value'] for row in csv.DictReader(handle)}
        assert parameters['Batch_journal'] == str(path)
        book = load_workbook(output / 'Nuclear_Intensity.xlsx', read_only=True)
        try:
            parameters = dict(book['Run_Info'].iter_rows(min_row=2, max_col=2, values_only=True))
            assert parameters['Batch_journal'] == str(path)
        finally:
            book.close()
    assert not list(manifest.parent.glob('Nuclear_Intensity_Batch_*'))
    assert not list(home.rglob('batch.json'))


def test_failed_image_updates_all_copies_and_retains_diagnostics(batch_setup):
    manifest, run, engine, home = batch_setup
    engine.measure.side_effect = [engine.measure.return_value, OSError('simulated image read failure')]
    assert app.main(manifest, mode='tiff-stack') == 1
    paths = result_journals(run)
    assert len(paths) == 2 and paths[0].read_bytes() == paths[1].read_bytes()
    journal = json.loads(paths[0].read_text())
    assert journal['status'] == 'incomplete'
    assert [item['status'] for item in journal['runs']] == ['complete', 'incomplete']
    failed = Path(journal['runs'][1]['output'])
    assert json.loads((failed / 'intensity_run.json').read_text())['Status'] == 'incomplete'
    assert 'simulated image read failure' in (run.parent / '3_nuclei_intensity.log').read_text()
    assert not list(home.rglob('batch.json'))


def test_interruption_preserves_completed_interrupted_and_pending_runs(batch_setup):
    manifest, run, engine, _ = batch_setup
    scheme = design(run.parent.parent, channels=['Nuclei: DAPI', 'Nuclear intensity: Marker 1', 'Nuclear intensity: Marker 2', 'Nuclear intensity: Marker 3'])
    plate_design.save_snapshot(run, scheme)
    engine.inspect.return_value['Channels'] = 4
    third = run.parent / 'markers' / 'unused'
    third.mkdir()
    (third / 'prepared.tif').touch()
    engine.measure.side_effect = [engine.measure.return_value, KeyboardInterrupt()]
    assert app.main(manifest, mode='tiff-stack') == 130
    paths = result_journals(run)
    assert len(paths) == 2 and paths[0].read_bytes() == paths[1].read_bytes()
    journal = json.loads(paths[0].read_text())
    assert journal['status'] == 'canceled'
    assert [item['status'] for item in journal['runs']] == ['complete', 'interrupted', 'pending']
    assert journal['runs'][2]['output'] is None
    interrupted = Path(journal['runs'][1]['output'])
    assert json.loads((interrupted / 'intensity_run.json').read_text())['Status'] == 'running'


def test_unavailable_result_volume_saves_a_fallback(batch_setup, monkeypatch, capsys):
    manifest, run, _, home = batch_setup
    original = app.new_output

    def fail_results(parent, prefix):
        if parent == run.parent:
            raise OSError('result volume unavailable')
        return original(parent, prefix)

    monkeypatch.setattr(app, 'new_output', fail_results)
    assert app.main(manifest, mode='tiff-stack') == 1
    paths = list((home / '.fia-tools' / 'logs').glob('Nuclear_Intensity_Batch_*/batch.json'))
    assert len(paths) == 1
    journal = json.loads(paths[0].read_text())
    assert journal['status'] == 'incomplete'
    assert all(item['status'] == 'failed' and item['output'] is None for item in journal['runs'])
    assert all('result volume unavailable' in item['error'] for item in journal['runs'])
    assert str(paths[0]) in capsys.readouterr().out
    assert not list(manifest.parent.glob('Nuclear_Intensity_Batch_*'))


def test_journal_write_failure_keeps_fallback_links_and_signals_failure(batch_setup, monkeypatch):
    manifest, run, _, home = batch_setup
    original = app.write_json

    def fail_result_journals(path, content):
        if path.name == 'batch.json' and path.parent.parent == run.parent:
            raise PermissionError('journal copy unavailable')
        return original(path, content)

    monkeypatch.setattr(app, 'write_json', fail_result_journals)
    assert app.main(manifest, mode='tiff-stack') == 1
    fallback = next(home.rglob('batch.json'))
    journal = json.loads(fallback.read_text())
    assert journal['status'] == 'complete'
    assert len(journal['journal_copy_errors']) == 2
    assert all('journal copy unavailable' in row['error'] for row in journal['journal_copy_errors'])
    for item in journal['runs']:
        info = json.loads((Path(item['output']) / 'intensity_run.json').read_text())
        assert info['Status'] == 'complete' and info['Batch_journal'] == str(fallback)


def test_cancel_before_measurement_creates_text_log_but_no_result_journal(batch_setup, monkeypatch):
    manifest, run, _, home = batch_setup
    import shutil
    shutil.copytree(run, run.with_name('Final_Nuclei_Mask_20260923_130000'))
    monkeypatch.setattr('builtins.input', lambda _: 'q')
    assert app.main(manifest, mode='tiff-stack') == 130
    assert not result_journals(run) and not list(home.rglob('batch.json'))
    assert not list(manifest.parent.glob('Nuclear_Intensity_Batch_*'))
