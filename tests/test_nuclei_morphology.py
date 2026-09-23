"""Morphology exports and optional integration checks against real ImageJ.

Set FIA_IMAGEJ_JAR to an existing ij.jar to run the native measurement tests.
The tests do not download dependencies or run StarDist/Fiji initialization.
"""

import csv
import json
import math
import os
import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

import numpy as np
import pytest
from openpyxl import load_workbook

import nuclei_morphology as morphology


@pytest.fixture
def exporter(tmp_path):
    output = tmp_path / "condition" / "foci_assay" / "Final_Nuclei_Mask_test"
    output.mkdir(parents=True)
    return morphology.NucleiMorphologyExport(
        output, output.parent / "Nuclei_StarDist_mask_processed_test", 2000, "test")


def read_sheet(workbook, name):
    rows = list(workbook[name].values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def test_empty_and_failed_images_have_distinct_counts(exporter):
    with patch.object(morphology, "measure_final_mask", return_value=([], MagicMock())), \
         patch.object(morphology, "jimport"), \
         patch.object(morphology, "save_numbered_image"):
        exporter.add_image(MagicMock(), "WellA02_empty_StarDist_processed.tif", "empty.tif")
    exporter.record_failure("bad_StarDist_processed.tif", "Cannot read mask")
    assert exporter.save() == "incomplete"
    with (exporter.output / "Nuclei_Images.csv").open(encoding="utf-8-sig") as handle:
        csv_rows = list(csv.DictReader(handle))
    assert csv_rows[0]["Nuclei_count"] == "0"
    assert csv_rows[1]["Nuclei_count"] == ""
    book = load_workbook(exporter.output / "Nuclei_Morphology.xlsx")
    try:
        assert book.sheetnames == ["Nuclei", "Images", "Run_Info"]
        assert book["Nuclei"].max_row == 1
        rows = read_sheet(book, "Images")
        assert rows[0]["Nuclei_count"] == 0
        assert rows[0]["Area_px2_Median"] is None
        assert rows[0]["Well"] == "A2"
        assert rows[0]["Condition"] is None
        assert rows[1]["Nuclei_count"] is None
        assert rows[1]["Status"] == "failed"
        assert rows[1]["Error"] == "Cannot read mask"
        assert "Orientation_deg_Median" not in rows[0]
    finally:
        book.close()


def test_measurement_failure_is_reported_without_losing_final_mask(tmp_path):
    mod = sys.modules["nuclei_mask_generation"]
    source = tmp_path / "foci_assay" / "Nuclei_StarDist_mask_processed_test"
    source.mkdir(parents=True)
    (source / "cell_StarDist_processed.tif").touch()
    (source / "._cell_StarDist_processed.tif").write_bytes(b"AppleDouble")
    ij, windows = MagicMock(), MagicMock()
    ij.saveAs.side_effect = lambda imp, fmt, path: Path(path).write_bytes(b"final mask")
    with patch.object(mod, "initialize_imagej"), \
         patch.object(mod, "jimport", side_effect=lambda cls: ij if cls == "ij.IJ" else windows), \
         patch.object(morphology, "measure_final_mask", side_effect=ValueError("bad geometry")):
        mod.process_nuclei([str(source)], 2000)
    output = next(source.parent.glob("Final_Nuclei_Mask_*"))
    assert ij.openImage.call_count == 1
    assert (output / "cell_StarDist_processed_processed.tif").read_bytes() == b"final mask"
    metadata = json.loads((output / "nuclei_run.json").read_text())
    assert metadata["status"] == "complete"
    assert metadata["morphology_status"] == "incomplete"
    book = load_workbook(output / "Nuclei_Morphology.xlsx")
    try:
        row = read_sheet(book, "Images")[0]
        assert row["Status"] == "failed"
        assert row["Nuclei_count"] is None
        assert row["Error"] == "bad geometry"
        assert row["Particle_size_px2"] == 2000
    finally:
        book.close()


def test_spreadsheet_write_error_is_not_silenced(exporter):
    exporter.record_failure("cell.tif", "bad mask")
    with patch.object(morphology.Workbook, "save", side_effect=OSError("disk full")):
        with pytest.raises(OSError, match="disk full"):
            exporter.save()


@pytest.fixture
def imagej_classes(monkeypatch):
    jar = os.environ.get("FIA_IMAGEJ_JAR")
    if not jar:
        pytest.skip("Set FIA_IMAGEJ_JAR to run real ImageJ measurements")
    assert Path(jar).is_file(), "FIA_IMAGEJ_JAR must point to an existing ij.jar"
    jpype = pytest.importorskip("jpype")
    if not jpype.isJVMStarted():
        jpype.startJVM("-Djava.awt.headless=true", classpath=[jar])
    monkeypatch.setattr(morphology, "jimport", jpype.JClass)
    return jpype.JClass


def image_from_pixels(classes, pixels, inverted=False):
    import jpype

    height, width = pixels.shape
    array = jpype.JArray(jpype.JByte)(pixels.astype(np.int8).ravel())
    processor = classes("ij.process.ByteProcessor")(width, height, array, None)
    if inverted:
        processor.invertLut()
    return classes("ij.ImagePlus")("Synthetic final mask", processor)


def synthetic_objects():
    pixels = np.zeros((128, 128), dtype=np.uint8)
    pixels[:12, 4:24] = 255
    pixels[30:50, 30:70] = 255
    row, col = np.ogrid[:128, :128]
    pixels[(row - 90) ** 2 + (col - 90) ** 2 <= 15 ** 2] = 255
    return pixels


@pytest.mark.parametrize("inverted", [False, True])
def test_native_geometry_pixel_units_and_mask_preservation(imagej_classes, inverted):
    classes = imagej_classes
    foreground = synthetic_objects()
    mask = image_from_pixels(classes, foreground if inverted else 255 - foreground, inverted)
    calibration = classes("ij.measure.Calibration")()
    calibration.pixelWidth, calibration.pixelHeight = 0.3, 0.8
    calibration.xOrigin, calibration.yOrigin = 4.0, 7.0
    calibration.setUnit("micron")
    mask.setCalibration(calibration)
    mask.setGlobalCalibration(calibration)
    labels = None
    try:
        before = morphology.processor_pixels(mask)
        records, labels = morphology.measure_final_mask(mask)
        assert [r["Area_px2"] for r in records] == [240, 800, 709]
        assert [r["Touches_border"] for r in records] == [True, False, False]
        np.testing.assert_array_equal(morphology.processor_pixels(labels) > 0, foreground > 0)
        rectangle = records[1]
        assert rectangle["Aspect_ratio"] == pytest.approx(2.0)
        assert rectangle["Solidity"] == pytest.approx(1.0)
        assert rectangle["Feret_max_px"] == pytest.approx(math.hypot(40, 20))
        assert rectangle["Feret_min_px"] == pytest.approx(20.0)
        assert rectangle["Equivalent_diameter_px"] == pytest.approx(math.sqrt(3200 / math.pi))
        assert rectangle["Eccentricity"] == pytest.approx(math.sqrt(0.75))
        assert rectangle["Centroid_X_px"] == pytest.approx(50.0)
        assert rectangle["Centroid_Y_px"] == pytest.approx(40.0)
        assert rectangle["Circularity"] == pytest.approx(
            4 * math.pi * 800 / rectangle["Perimeter_px"] ** 2)
        assert records[2]["Aspect_ratio"] == pytest.approx(1.0)
        np.testing.assert_array_equal(before, morphology.processor_pixels(mask))
        assert mask.getCalibration().pixelWidth == 0.3
        assert bool(mask.isInvertedLut()) == inverted
    finally:
        mask.setGlobalCalibration(None)
        if labels is not None:
            labels.close()
        mask.close()


@pytest.mark.parametrize("inverted", [False, True])
def test_native_empty_mask(imagej_classes, inverted):
    pixels = np.full((32, 32), 0 if inverted else 255, dtype=np.uint8)
    mask = image_from_pixels(imagej_classes, pixels, inverted)
    try:
        records, labels = morphology.measure_final_mask(mask)
        assert records == []
        assert not morphology.processor_pixels(labels).any()
        labels.close()
    finally:
        mask.close()


def test_native_particle_masks_export_consistent_rows_and_qc(imagej_classes, exporter):
    classes = imagej_classes
    source = image_from_pixels(classes, synthetic_objects())
    source.getProcessor().setThreshold(255, 255, classes("ij.process.ImageProcessor").NO_LUT_UPDATE)
    analyzer_type = classes("ij.plugin.filter.ParticleAnalyzer")
    # Reproduce the existing final-mask output; remove only the 240-pixel object.
    analyzer = analyzer_type(analyzer_type.SHOW_MASKS, 0, classes("ij.measure.ResultsTable")(),
                             300.0, float("inf"))
    analyzer.setHideOutputImage(True)
    assert analyzer.analyze(source)
    mask = analyzer.getOutputImage()
    before = morphology.processor_pixels(mask)
    try:
        exporter.particle_size = 300
        exporter.add_image(mask, "WellB12_cell_StarDist_processed.tif", "final.tif")
        assert exporter.save() == "complete"
        np.testing.assert_array_equal(before, morphology.processor_pixels(mask))
        book = load_workbook(exporter.output / "Nuclei_Morphology.xlsx")
        try:
            rows = read_sheet(book, "Nuclei")
            assert [row["Area_px2"] for row in rows] == [800, 709]
            assert all(row["Well"] == "B12" for row in rows)
            assert all(row["Particle_size_px2"] == 300 for row in rows)
            summary = read_sheet(book, "Images")[0]
            assert summary["Nuclei_count"] == 2
            assert summary["Border_nuclei_count"] == 0
            assert summary["Area_px2_Median"] == 754.5
            assert summary["Area_px2_IQR"] == 45.5
            with (exporter.output / "Nuclei_Morphology.csv").open(encoding="utf-8-sig") as handle:
                csv_rows = list(csv.DictReader(handle))
            assert len(csv_rows) == len(rows)
            assert float(csv_rows[0]["Circularity"]) == pytest.approx(rows[0]["Circularity"])
            label_path = exporter.output / rows[0]["Label_map"]
            assert label_path.parent.name == "Morphology_QC"
            saved = classes("ij.IJ").openImage(str(label_path))
            try:
                assert int(saved.getBitDepth()) == 16
                labels = morphology.processor_pixels(saved)
                for row in rows:
                    assert np.count_nonzero(labels == row["Nucleus_ID"]) == row["Area_px2"]
            finally:
                saved.close()
            from PIL import Image
            with Image.open(exporter.output / rows[0]["Numbered_image"]) as png:
                rgb = np.asarray(png)
            assert np.any((rgb[:, :, 0] == 255) & (rgb[:, :, 1] == 0))
            assert not list(exporter.output.glob("*_ids.tif"))
        finally:
            book.close()
    finally:
        source.close()
        mask.close()


def test_native_nonbinary_mask_is_rejected(imagej_classes):
    mask = image_from_pixels(imagej_classes, np.full((32, 32), 128, dtype=np.uint8))
    try:
        with pytest.raises(ValueError, match="not binary"):
            morphology.measure_final_mask(mask)
    finally:
        mask.close()
