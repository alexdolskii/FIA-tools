"""PLOT LAYOUT order/geometry and inference invariance on the current template."""
import json
from copy import copy, deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

import marker_intensity_report as report
import marker_report_data as inputs
import marker_report_plots as plots
import plate_design
from marker_report_statistics import calculate_statistics
from test_collect_marker_intensity_results import sheet_rows
from test_marker_intensity_report import MARKER, annotated, collection, plate


def layout(path, panels):
    book = load_workbook(path)
    sheet = book['Plate Map']
    cells = {c.value: c for row in sheet['B2:M9'] for c in row if c.value}
    for row in sheet['B15:M21']:
        for cell in row:
            cell.value = None
            cell.fill = PatternFill()
            cell.font = Font()
    for row, column, labels in panels:
        for offset, label in enumerate(labels):
            source = cells[label]
            target = sheet.cell(row, column+offset, label)
            target.fill, target.font = copy(source.fill), copy(source.font)
    book.save(path)
    book.close()
    return path


def read_design(path, unit='nucleus', data=None):
    data = data or {'files': {}, 'images': [], 'nuclei': [], 'markers': [MARKER], 'report_kind': 'marker'}
    data.setdefault('files', {})
    return inputs.annotate(data, path, [MARKER], unit)


def test_layout_controls_axes_even_with_control_last(tmp_path):
    path = layout(plate(tmp_path/'FIA.xlsx'), [(15, 2, ['Treatment', 'Control'])])
    data = read_design(path)
    assert data['groups'] == ['Treatment', 'Control']
    assert data['group_order_source'] == 'plot_layout'
    assert [r['Well'] for r in data['design']] == ['A02', 'A03']
    style = plots.prepare_plot_style(data)
    control = next(r for r in style['conditions'] if r['Group'] == 'Control')
    assert control['X_position'] == 2 and control['Condition_color'] == '#B5B1D8'


@pytest.mark.parametrize('change', ['missing', 'unknown', 'duplicate', 'bold', 'color',
                                  'merged', 'formula', 'split', 'vertical', 'legacy_grid'])
def test_invalid_layout_is_rejected_before_aggregation(tmp_path, change):
    path = plate(tmp_path/'FIA.xlsx')
    book = load_workbook(path)
    sheet = book['Plate Map']
    if change == 'missing':
        sheet['C15'] = None
    elif change == 'unknown':
        sheet['C15'] = 'Unknown'
    elif change == 'duplicate':
        sheet['C15'] = 'Control'
        sheet['C15'].font = Font(bold=True)
    elif change == 'bold':
        sheet['B15'].font = Font(bold=False)
    elif change == 'color':
        sheet['C15'].fill = PatternFill('solid', fgColor='FF0000')
    elif change == 'merged':
        sheet.merge_cells('B15:C15')
    elif change == 'formula':
        sheet['C15'] = '="Treatment"'
    elif change in ('split', 'vertical'):
        target = sheet['F15'] if change == 'split' else sheet['B16']
        target.value = sheet['C15'].value
        target._style = copy(sheet['C15']._style)
        sheet['C15'] = None
    else:
        del book['Channels']
    book.save(path)
    book.close()
    with pytest.raises(inputs.ValidationError):
        read_design(path)


@pytest.mark.parametrize('position', [(15, 7), (18, 2), (18, 7)])
def test_blocks_can_be_side_by_side_stacked_or_grid(tmp_path, position):
    groups = {'A02': ('A control', True, 'ABCDEF'), 'A03': ('A treatment', False, 'ABCDEF'),
              'B02': ('B control', True, 'CCBBAA'), 'B03': ('B treatment', False, 'CCBBAA')}
    path = plate(tmp_path/'FIA.xlsx', groups)
    layout(path, [(15, 2, ['A control', 'A treatment']), (*position, ['B treatment', 'B control'])])
    panels = plots.plot_panels(read_design(path))
    assert panels[1]['groups'] == ['B treatment', 'B control']
    assert (panels[1]['layout_row'], panels[1]['layout_column']) == position
    rows, _, places, _ = plots.panel_grid(panels)
    assert rows == (1 if position[0] == 15 else 2)
    assert len(places) == 2


def test_order_survives_filtering_without_changing_statistics(tmp_path):
    base = annotated()
    groups = {r['Well']: (r['Group'], r['Group'] == 'Control') for r in base['design']}
    groups['A06'] = ('No images', False)
    path = plate(tmp_path/'FIA.xlsx', groups)
    original = read_design(path, data=deepcopy(base))
    layout(path, [(15, 2, ['No images', 'Treatment', 'Control'])])
    reordered = read_design(path, data=deepcopy(base))
    for data in (original, reordered):
        data['plot_style'] = plots.prepare_plot_style(data, 3)
        inputs.filter_images(data, 3)
        inputs.aggregate(data)
        calculate_statistics(data)
    key = lambda r: (r['Block'], r['Metric'], r['Control'], r['Treatment'])
    assert sorted(original['statistics'], key=key) == sorted(reordered['statistics'], key=key)
    assert original['images'] == reordered['images'] and original['nuclei'] == reordered['nuclei']
    assert reordered['plot_style']['panels'][0]['groups'] == ['No images', 'Treatment', 'Control']


def test_layout_is_saved_in_report_and_source_snapshot(tmp_path):
    path = collection(tmp_path)
    template = path.parent.parent/'FIA_scheme.xlsx'
    layout(template, [(18, 5, ['Treatment', 'Control'])])
    ok, output = report.create_report(path, path.parent.parent, markers=[MARKER], stats_unit='nucleus')
    assert ok
    assert [r['Group'] for r in sheet_rows(output/report.WORKBOOK, 'Group_Order')] == ['Treatment', 'Control']
    status = json.loads((output/'report_status.json').read_text())
    assert status['Condition_order_source'] == 'plot_layout'
    assert (output/'plate_design.xlsx').read_bytes() == template.read_bytes()
    assert not list((output/'Plots').glob('*Nuclei*'))


def test_blank_template_has_no_assignments_and_golden_is_complete():
    folder = Path(__file__).parents[1]/'templates'
    blank = plate_design.read_design(folder/'FIA_96_well_Template.xlsx', allow_empty=True)
    assert blank.wells == {} and blank.panels == []
    assert len(plate_design.read_design(folder/'FIA_96_well_Golden_Standard.xlsx').wells) == 96
