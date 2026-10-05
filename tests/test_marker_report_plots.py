"""Plot regressions: readable panels, complete distributions and unchanged inference."""

from copy import deepcopy
from io import StringIO
from itertools import pairwise
from zipfile import ZipFile

import collect_marker_intensity_results as collector
import marker_intensity_report as report
import marker_report_data as inputs
import marker_report_plots as plots
import matplotlib.pyplot as plt
import numpy as np
import pytest
from marker_report_statistics import calculate_statistics
from matplotlib.collections import PathCollection, PolyCollection
from matplotlib.figure import Figure
from PIL import Image
from table_workflow import TableProgress

MARKER = 'Foci_1_Channel_2'


def four_blocks(unit, contexts=None):
    data = {'images': [], 'nuclei': [], 'groups': [], 'design': [], 'blocks': {}, 'markers': [MARKER],
            'stats_unit': unit, 'run_id': 'Final_Nuclei_Mask_20260929_100000'}
    fields = [field for _, _, field, _ in inputs.metric_specs([MARKER])[1:]]
    contexts = contexts or ('P4+8i old serum', 'P4+8i new serum', 'P4+13i old serum', 'P4+13i new serum')
    for block, context in enumerate(contexts):
        color = f'color-{block}'
        data['blocks'][color] = {}
        for treatment in ('dmso', 'tgfb inhibitor', 'citric acid', 'tgfb ligand'):
            group = 'BK near (tert-blast) ' + context + ' ' + treatment
            data['groups'].append(group)
            data['blocks'][color][group] = treatment == 'dmso'
            for repeat in range(2):
                well = f'{chr(65 + block)}{len(data["blocks"][color]) * 2 + repeat:02d}'
                identity = {'Group': group, 'Well': well, 'Image_name': f'field_Well{well}.tif', 'Mask_name': well}
                data['design'].append({'Group': group, 'Well': well, 'Color': color})
                values = np.arange(1, 21, dtype=float) + repeat
                image = dict(identity, **{inputs.COUNT: len(values)})
                for field in fields:
                    for stat, value in zip(collector.STATS, (np.mean(values), np.median(values),
                                                           np.percentile(values, 75) - np.percentile(values, 25))):
                        image[field + '_' + stat] = float(value)
                data['images'].append(image)
                for index, value in enumerate(values):
                    data['nuclei'].append(dict(identity, Nucleus_ID=index + 1,
                                               **{field: float(value) for field in fields}))
    if unit is None:
        data['blocks'] = {}
    inputs.aggregate(data)
    return calculate_statistics(data)


@pytest.mark.parametrize('unit', ['nucleus', 'well', None])
def test_four_panels_keep_statistics_observations_and_readable_labels(tmp_path, monkeypatch, unit):
    data = four_blocks(unit)
    before = deepcopy(data)
    snapshots = []
    original_save = Figure.savefig

    def inspect(fig, filename, **kwargs):
        if filename.suffix == '.pdf':
            return original_save(fig, filename, **kwargs)
        fig.canvas.draw()
        renderer = fig.canvas.get_renderer()
        axes = [ax for ax in fig.axes if ax.get_label().startswith('Block_')]
        for ax in axes:
            labels = ax.get_xticklabels()
            boxes = [label.get_window_extent(renderer) for label in labels]
            assert all(left.x1 < right.x0 for left, right in pairwise(boxes))
            samples = [text for text in ax.texts if text.get_gid() == 'sample_size']
            assert len(samples) == len(labels)
            assert all('n=' in label.get_text() and '2 wells' in label.get_text() for label in samples)
            assert all(label.get_fontsize() == 14 for label in labels)
            assert all(label.get_fontsize() == 12 for label in samples)
            sample_boxes = [label.get_window_extent(renderer) for label in samples]
            assert all(left.x1 < right.x0 for left, right in pairwise(sample_boxes))
            assert all(label.y0 > sample.y1 for label, sample in zip(boxes, sample_boxes))
            assert ax.get_yscale() == 'linear'
            significance = [text for text in ax.texts if text.get_gid() == 'significance']
            for left, right in pairwise(significance):
                assert not left.get_window_extent(renderer).overlaps(right.get_window_extent(renderer))
            for text in significance:
                assert not text.get_window_extent(renderer).overlaps(ax.title.get_window_extent(renderer))
            dots = sum(len(c.get_offsets()) for c in ax.collections if isinstance(c, PathCollection))
            assert dots == (8 if filename.name.startswith('Nuclei_count') else 0)
            for label in labels + samples:
                assert not label.get_window_extent(renderer).overlaps(fig.axes[-1].get_window_extent(renderer))
        snapshots.append((filename.name, [ax.get_ylim() for ax in axes],
                          [t.get_text() for ax in axes for t in ax.texts if t.get_gid() == 'significance']))
        return original_save(fig, filename, **kwargs)

    monkeypatch.setattr(Figure, 'savefig', inspect)
    progress = TableProgress(stream=StringIO())
    plots.render_plots(data, tmp_path / 'Plots', progress, plot_format='both')
    assert all(data[key] == value for key, value in before.items())
    assert data['plot_data'] == plots.plot_rows(before)
    assert len(data['plots']) == 2
    assert progress.done == progress.phase_total == 2
    assert {path.name for path in (tmp_path / 'Plots').iterdir()} == {
        name + suffix for name in ('Nuclei_count', MARKER + '_Integrated_density')
        for suffix in ('.png', '.pdf')}
    assert len(data['plot_labels']) == 16
    assert {r['Display_label'] for r in data['plot_labels']} == {'dmso', 'tgfb inhibitor', 'citric acid', 'tgfb ligand'}
    assert {r['Figure_context'] for r in data['plot_labels']} == {'BK near (tert-blast)'}
    assert {r['Group'] for r in data['plot_labels']} == set(data['groups'])
    for metric in ('Nuclei_count', MARKER + '_Integrated_density'):
        selected = [item for item in snapshots if item[0].startswith(metric)]
        assert len(selected) == 1
        assert len({limits for _, axes, _ in selected for limits in axes}) == 1
        assert len(selected[0][1]) == 4
        assert len(selected[0][2]) == (0 if unit is None else 12)
        assert all(text in ('ns', '*', '**', '***', '****', 'Not tested') for _, _, texts in selected for text in texts)
    for plot in data['plots']:
        assert plot['View'] == 'overview' and len(plot['Panels'].split(', ')) == 4
        assert plot['Points'] == sum(row['Plot'] == plot['Plot'] for row in data['plot_data'])
        assert (tmp_path / plot['PDF_file']).read_bytes().startswith(b'%PDF-')
        with Image.open(tmp_path / plot['File']) as image:
            assert image.info['dpi'][0] == pytest.approx(300, abs=0.01)
    tables = {'Plot_Info': report.as_table(data['plots']), 'Plot_Labels': report.as_table(data['plot_labels'])}
    report.write_workbook(tmp_path, tables, data['plots'])
    with ZipFile(tmp_path / report.WORKBOOK) as archive:
        assert len([name for name in archive.namelist() if name.startswith('xl/media/')]) == 2


@pytest.mark.parametrize('values', [[], [7], [7] * 100, [1, 2, 3, 9], list(range(5000)) + [100000]])
def test_nucleus_density_has_no_dots_and_keeps_extremes(values):
    fig, ax = plt.subplots()
    try:
        plots.draw_distribution(ax, values, 1, 'nucleus')
        assert not any(isinstance(c, PathCollection) for c in ax.collections)
        bodies = [c for c in ax.collections if isinstance(c, PolyCollection)]
        assert bool(bodies) == (len(values) >= 5 and max(values) > min(values))
        if bodies:
            vertices = bodies[0].get_paths()[0].vertices
            assert vertices[:, 1].min() == min(values)
            assert vertices[:, 1].max() == max(values)
            median = [line for line in ax.lines if np.all(np.asarray(line.get_ydata()) == np.median(values))]
            assert any(line.get_zorder() > ax.patches[0].get_zorder() for line in median)
        if values:
            assert ax.dataLim.ymin == min(values)
            assert ax.dataLim.ymax == max(values)
    finally:
        plt.close(fig)


def test_more_than_four_blocks_share_one_complete_overview(tmp_path, monkeypatch):
    data = four_blocks('nucleus', contexts=[f'P4+{index}i old serum' for index in range(5)])
    before = deepcopy(data)
    snapshots = []
    original_save = Figure.savefig

    def inspect(fig, filename, **kwargs):
        if filename.suffix == '.pdf':
            return original_save(fig, filename, **kwargs)
        if '__Block_' not in filename.name:
            fig.canvas.draw()
            renderer = fig.canvas.get_renderer()
            axes = [ax for ax in fig.axes if ax.get_label().startswith('Block_')]
            assert len(axes) == 5
            assert len({ax.get_ylim() for ax in axes}) == 1
            labels = [label for ax in axes for label in ax.get_xticklabels()]
            assert len(labels) == 20
            for label in labels:
                box = label.get_window_extent(renderer)
                assert not box.overlaps(fig.axes[-1].get_window_extent(renderer))
                assert box.y0 >= 0 and box.y1 <= fig.bbox.height
            assert sum(sum(t.get_gid() == 'significance' for t in ax.texts) for ax in axes) == 15
            snapshots.append(filename.name)
        return original_save(fig, filename, **kwargs)

    monkeypatch.setattr(Figure, 'savefig', inspect)
    plots.render_plots(data, tmp_path / 'Plots', plot_format='both')
    assert snapshots == ['Nuclei_count.png', MARKER + '_Integrated_density.png']
    assert not list((tmp_path / 'Plots').glob('*__Page_*.png'))
    assert len(data['plots']) == 2
    assert not list((tmp_path / 'Plots').glob('*__Block_*'))
    assert data['statistics'] == before['statistics']
    assert data['plot_data'] == plots.plot_rows(before)
    for spec in plots.plot_specs(data):
        views = [row for row in data['plots'] if row['Plot'] == spec[0]]
        assert len(views) == 1
        overview = next(row for row in views if row['View'] == 'overview')
        assert len(overview['Panels'].split(', ')) == 5
        assert overview['Points'] == sum(row['Plot'] == spec[0] for row in data['plot_data'])


def test_display_shortening_never_merges_conditions():
    groups = ['control', 'control treated', 'control  treated']
    _, labels = plots.short_labels(groups)
    assert len(set(labels.values())) == len(groups)
    assert labels == {group: group for group in groups}
    data = four_blocks(None)
    data['design'].append(dict(data['design'][0], Color='conflicting color'))
    panels = plots.plot_panels(data)
    assert len(panels) == 1 and panels[0]['groups'] == data['groups']


@pytest.mark.parametrize('p,expected', [(None, 'Not tested'), (1, 'ns'), (0.05, 'ns'),
                                       (0.049, '*'), (0.01, '*'), (0.009, '**'),
                                       (0.001, '**'), (0.0009, '***'), (0.0001, '***'), (0.00009, '****')])
def test_plot_symbols_use_adjusted_p_and_preserve_exact_values(p, expected):
    row = {'P_Holm': p, 'P_raw': 0.0000001}
    original = dict(row)
    assert plots.p_label(row) == expected
    assert row == original


def test_long_condition_token_wraps_without_losing_characters():
    name = 'VeryLongTreatmentNameWithoutAnySpaces_10uM'
    wrapped = plots.wrap_label(name, 90, 14)
    assert '\n' in wrapped and wrapped.replace('\n', '') == name
    font = plots.FontProperties(family=plots.plot_font(), size=14)
    assert all(plots.TextPath((0, 0), line, prop=font).get_extents().width <= 90 for line in wrapped.split('\n'))


def test_eight_panels_retain_all_blocks_with_short_caption(tmp_path, monkeypatch):
    data = four_blocks('nucleus')
    blocks = {}
    for color, groups in data['blocks'].items():
        members = list(groups)
        blocks[color + '-a'] = {members[0]: True, members[1]: False}
        blocks[color + '-b'] = {members[2]: True, members[3]: False}
    data['blocks'] = blocks
    calculate_statistics(data)
    data['min_nuclei'] = 20
    data['image_exclusion_rule'] = 'Images with Non_border_nuclei_count < 20 excluded before analysis.'
    before = deepcopy(data)
    captured = []
    original_save = Figure.savefig

    def inspect(fig, filename, **kwargs):
        if filename.suffix == '.png' and '__Block_' not in filename.name:
            fig.canvas.draw()
            renderer = fig.canvas.get_renderer()
            axes = [ax for ax in fig.axes if ax.get_label().startswith('Block_')]
            assert len(axes) == 8
            assert [ax.get_title().split()[0] for ax in axes] == list('ABCDEFGH')
            assert len({ax.get_ylim() for ax in axes}) == 1
            assert len(fig.axes[0].texts) == 2
            assert not any('Final_Nuclei_Mask' in t.get_text() for t in fig.axes[0].texts)
            for ax in axes:
                texts = list(ax.get_xticklabels()) + [t for t in ax.texts if t.get_gid() == 'sample_size']
                for text in texts:
                    box = text.get_window_extent(renderer)
                    assert box.x0 >= 0 and box.x1 <= fig.bbox.width and box.y0 >= 0
                    assert not box.overlaps(fig.axes[-1].get_window_extent(renderer))
                    for other in axes:
                        if other is not ax:
                            assert not box.overlaps(other.get_window_extent(renderer))
            captured.append(filename.name)
        return original_save(fig, filename, **kwargs)

    monkeypatch.setattr(Figure, 'savefig', inspect)
    plots.render_plots(data, tmp_path / 'Plots', plot_format='both')
    assert captured == ['Nuclei_count_all.png', 'Nuclei_count.png', MARKER + '_Integrated_density.png']
    assert len(data['plots']) == 3
    assert data['statistics'] == before['statistics']
    assert data['plot_data'] == plots.plot_rows(before)
    assert len(list((tmp_path / 'Plots').glob('*.pdf'))) == 3
    assert not list((tmp_path / 'Plots').glob('*__Block_*'))
    for plot in data['plots']:
        if plot['Plot'] == 'Nuclei_count_all':
            assert '<20' in plot['Display_caption'] and 'Diagnostic only' in plot['Display_caption']
            assert 'Welch' not in plot['Display_caption']
            continue
        assert '≥20' in plot['Display_caption'] and 'Welch + Holm' in plot['Display_caption']
        assert 'within-well dependence' in plot['Caption']
        assert len(plot['Display_caption']) < len(plot['Caption'])


def test_descriptive_template_retains_direct_colors_without_requiring_controls(tmp_path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from test_marker_intensity_report import plate

    path = plate(tmp_path / 'plate.xlsx')
    book = load_workbook(path)
    book.active['C2'].font = Font()
    book.active['D2'].fill = PatternFill('solid', fgColor='ABC123')
    book.save(path)
    _, design = inputs.read_template(path, {}, statistics=False)
    assert len({row['Color'] for row in design}) == 2
    with pytest.raises(inputs.ValidationError, match='one control'):
        # Color extraction itself remains independent of the statistical control validation.
        inputs.annotate({'images': [], 'nuclei': [], 'markers': [], 'files': {}}, path, [], 'nucleus')
