"""Build real current-format schemes for focused regression fixtures."""
from pathlib import Path
from copy import copy
import json
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
import plate_design as plate

TEMPLATE = Path(__file__).parents[1] / 'templates/FIA_96_well_Template.xlsx'
M1 = plate.marker_id('Nuclear intensity', 'Marker 1')
M2 = plate.marker_id('Nuclear intensity', 'Marker 2')
M3 = plate.marker_id('Nuclear intensity', 'Marker 3')

def workbook(path, groups=None, channels=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    groups = groups or {'A02': ('Control', True), 'A03': ('Treatment', False)}
    channels = channels or ['Nuclei: DAPI', 'Nuclear intensity: Marker 1', 'Nuclear intensity: Marker 2']
    book = load_workbook(TEMPLATE)
    sheet, roles = book['Plate Map'], book['Channels']
    labels = {}
    for well, item in groups.items():
        group, control = item[:2]
        color = item[2] if len(item) > 2 else 'ABCDEF'
        cell = sheet.cell(ord(well[0])-65+2, int(well[1:])+1, group)
        cell.fill = PatternFill('solid', fgColor=color)
        cell.font = Font(bold=control)
        labels.setdefault(color, {})[group] = cell
        values = channels[well] if isinstance(channels, dict) else channels
        for i, value in enumerate(values):
            roles.cell(6+4*(ord(well[0])-65)+i, int(well[1:])+2, value)
    for r, block in enumerate(labels.values(), 15):
        for c, cell in enumerate(block.values(), 2):
            target = sheet.cell(r,c,cell.value)
            target.fill, target.font = copy(cell.fill), copy(cell.font)
    book.save(path)
    book.close()
    return path

def design(root, groups=None, channels=None):
    return plate.read_design(workbook(Path(root)/'FIA_scheme.xlsx', groups, channels))

def prepared(folder, design):
    folder = Path(folder)
    plate.save_snapshot(folder, design)
    names = [p for p in folder.iterdir() if p.is_file() and not p.name.startswith('.') and p.suffix.lower() in ('.tif', '.tiff')]
    (folder/plate.PREPARATION).write_text(json.dumps({'Status':'COMPLETE', 'Files':{p.name:plate.file_digest(p) for p in names}, 'Sources':{}}))
