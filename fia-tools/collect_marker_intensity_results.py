"""Collect nuclear morphology and marker-intensity spreadsheets without ImageJ."""

import csv
import hashlib
import io
import json
import math
import re
import shutil
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

OUTPUT_PREFIX = 'FIA_Marker_Intensity_Combined_Results_'
COMBINED_NAME = 'FIA_Marker_Intensity_Combined.xlsx'
NUCLEI_PATTERN = re.compile(r'^Final_Nuclei_Mask_(\d{8}_\d{6})$')
INTENSITY_PATTERN = re.compile(r'^Nuclear_Intensity_(?:(Foci_[1-9][0-9]*_Channel_[1-9][0-9]*)_)?(\d{8}_\d{6})$')
MARKER_PATTERN = re.compile(r'^Foci_[1-9][0-9]*_Channel_([1-9][0-9]*)$')
MASK_SUFFIX = '_nuclei_projection_StarDist_processed_processed'
MORPH_METRICS = (
    'Area_px2', 'Perimeter_px', 'Circularity', 'Aspect_ratio', 'Solidity',
    'Major_axis_px', 'Minor_axis_px', 'Feret_max_px', 'Feret_min_px',
    'Equivalent_diameter_px', 'Roundness', 'Eccentricity',
)
INTENSITY_METRICS = ('Marker_Mean', 'Marker_Median', 'Marker_StdDev',
                     'Marker_Min', 'Marker_Max', 'Marker_RawIntDen')
COUNTS = ('Nuclei_count_total', 'Border_nuclei_count', 'Non_border_nuclei_count')
STATS = ('Mean', 'Median', 'IQR')
INFO_COLUMNS = ['Category', 'Item', 'Status', 'Source', 'Details', 'SHA256']


class ValidationError(ValueError):
    """A result cannot be selected or safely combined."""


class Cancelled(Exception):
    """Interactive selection was canceled before collection."""


def parts(path):
    """Compare recorded paths across slash conventions without resolving old drives."""
    return tuple(part for part in str(path).replace('\\', '/').split('/') if part)


def basename(path):
    values = parts(path)
    return values[-1] if values else ''


def number(value, field, integer=False, blank=False):
    if value in ('', None) and blank:
        return None
    try:
        result = float(value)
    except (TypeError, ValueError) as error:
        raise ValidationError(f'Invalid numeric {field}: {value!r}') from error
    if not math.isfinite(result) or (integer and not result.is_integer()):
        raise ValidationError(f'Invalid numeric {field}: {value!r}')
    return int(result) if integer else result


def snapshot(path, files):
    path = Path(path)
    if path.is_symlink() or path.parent.is_symlink() or not path.is_file():
        raise ValidationError(f'Missing, nonregular or linked result file: {path}')
    data = path.read_bytes()
    if path in files and files[path] != data:
        raise ValidationError(f'Source changed while reading: {path}')
    files[path] = data
    return data


def json_snapshot(path, files):
    result = json.loads(snapshot(path, files).decode('utf-8-sig'))
    if not isinstance(result, dict):
        raise ValidationError(f'Expected a JSON object: {path}')
    return result


def csv_snapshot(path, files, required):
    text = snapshot(path, files).decode('utf-8-sig')
    # Run_Info can contain fingerprints for hundreds of original images.
    csv.field_size_limit(max(csv.field_size_limit(), len(text)))
    reader = csv.DictReader(io.StringIO(text, newline=''), strict=True)
    header = reader.fieldnames
    if not header or any(not col for col in header) or len(set(header)) != len(header):
        raise ValidationError(f'Empty or duplicate CSV headers: {path}')
    missing = set(required) - set(header)
    if missing:
        raise ValidationError(f'{path.name}: missing columns {sorted(missing)}')
    rows = []
    for row in reader:
        if None in row or any(value is None for value in row.values()):
            raise ValidationError(f'{path.name}: malformed CSV row {reader.line_num}')
        rows.append(row)
    return header, rows


def matching_cell(value, text, allow_truncated=False):
    if value is None:
        return text == ''
    if isinstance(value, bool):
        return str(value) == text
    if isinstance(value, (float, int)):
        try:
            return math.isclose(float(value), float(text), rel_tol=1e-12, abs_tol=1e-12)
        except ValueError:
            return False
    return str(value) == text or (allow_truncated and len(text) > 32767 and str(value) == text[:32767])


def verify_workbook(path, files, tables):
    """Check CSV pairs, allowing Excel precision and its Run_Info text-cell limit."""
    book = None
    try:
        book = load_workbook(io.BytesIO(snapshot(path, files)), read_only=True, data_only=False)
        for name, (columns, rows) in tables.items():
            if name not in book.sheetnames:
                raise ValidationError(f'{path.name}: missing sheet {name}')
            iterator = iter(book[name].values)
            if list(next(iterator, ())) != columns:
                raise ValidationError(f'{path.name}: {name} headers differ from CSV')
            for index, row in enumerate(rows, 2):
                values = next(iterator, None)
                if values is not None and len(values) <= len(columns):
                    values = tuple(values) + (None,) * (len(columns) - len(values))
                if values is None or len(values) != len(columns) or not all(
                        matching_cell(value, row[column], allow_truncated=name == 'Run_Info' and column == 'Value')
                        for value, column in zip(values, columns)):
                    raise ValidationError(f'{path.name}: {name} row {index} differs from CSV')
            if next(iterator, None) is not None:
                raise ValidationError(f'{path.name}: extra rows in {name}')
    except Exception as error:
        raise ValidationError(f'Invalid workbook {path.name}: {error}') from error
    finally:
        if book is not None:
            book.close()


def read_tables(run, kind, files):
    if kind == 'morphology':
        names = ('Nuclei_Morphology.csv', 'Nuclei_Images.csv', 'Nuclei_Run_Info.csv', 'Nuclei_Morphology.xlsx')
        identity = ['Dataset_path', 'Run_ID', 'Mask_name', 'Image_name', 'Particle_size_px2']
        nucleus = identity + ['Nucleus_ID', 'Touches_border'] + list(MORPH_METRICS)
        images = identity + ['Status', 'Error'] + list(COUNTS)
        images += [f'{metric}_{stat}' for metric in MORPH_METRICS for stat in STATS]
    else:
        names = ('Nuclear_Intensity.csv', 'Nuclear_Intensity_Images.csv',
                 'Nuclear_Intensity_Run_Info.csv', 'Nuclear_Intensity.xlsx')
        identity = ['Dataset_path', 'Nuclei_run_ID', 'Intensity_run_ID', 'Mask_file',
                    'Image_name', 'Source_file', 'Marker_channel', 'Particle_size_px2']
        nucleus = identity + ['Nucleus_ID', 'Area_px2'] + list(INTENSITY_METRICS)
        images = identity + ['Status', 'Error'] + list(COUNTS) + ['Nuclei_measured_count', 'Failed_nuclei_count']
        images += [f'{metric}_{stat}' for metric in ('Area_px2', *INTENSITY_METRICS) for stat in STATS]
    tables = {
        'Nuclei': csv_snapshot(run / names[0], files, nucleus),
        'Images': csv_snapshot(run / names[1], files, images),
        'Run_Info': csv_snapshot(run / names[2], files, ['Parameter', 'Value']),
    }
    verify_workbook(run / names[3], files, tables)
    parameters = {}
    for row in tables['Run_Info'][1]:
        if row['Parameter'] in parameters:
            raise ValidationError('Duplicate Run_Info parameter')
        parameters[row['Parameter']] = row['Value']
    if parameters.get('Status') != 'complete':
        raise ValidationError('Run_Info status is not complete')
    return tables, parameters, [run / name for name in names]


def validate_rows(bundle, kind):
    """Validate unique mask/ID keys, counts and the source dataset namespace."""
    run, meta, tables = bundle['run'], bundle['metadata'], bundle['tables']
    nucleus_rows, image_rows = tables['Nuclei'][1], tables['Images'][1]
    if not image_rows:
        raise ValidationError('No image rows')
    namespace = parts(image_rows[0]['Dataset_path'])
    if not namespace:
        raise ValidationError('Empty dataset path')
    run_id = run.name if kind == 'morphology' else basename(meta['Nuclei_run'])
    particle_size = number(meta['particle_size_pixels_squared'] if kind == 'morphology'
                           else meta['Particle_size_px2'], 'particle size')
    if particle_size < 0:
        raise ValidationError('Negative particle size')
    images, nuclei = {}, {}
    for table, rows in (('Images', image_rows), ('Nuclei', nucleus_rows)):
        for row in rows:
            if parts(row['Dataset_path']) != namespace:
                raise ValidationError('Mixed datasets inside one result')
            if row['Run_ID' if kind == 'morphology' else 'Nuclei_run_ID'] != run_id:
                raise ValidationError('Nuclei run ID differs from metadata')
            if number(row['Particle_size_px2'], 'Particle_size_px2') != particle_size:
                raise ValidationError('Particle size differs from metadata')
            mask = row['Mask_name'] if kind == 'morphology' else basename(row['Mask_file'])
            if not mask or mask.startswith('.') or basename(mask) != mask or '\x00' in mask:
                raise ValidationError('Invalid or hidden mask filename')
            if kind == 'intensity':
                if row['Intensity_run_ID'] != run.name:
                    raise ValidationError('Intensity run ID differs from its folder')
                if number(row['Marker_channel'], 'Marker_channel', integer=True) != bundle['channel']:
                    raise ValidationError('Marker channel differs from metadata')
                if meta.get('Schema_version') == 2 and row.get('Marker_folder') != bundle['marker']:
                    raise ValidationError('Marker folder differs from metadata')
                if parts(row['Mask_file'])[:-1] != parts(meta['Nuclei_run']):
                    raise ValidationError('Mask path differs from the recorded nucleus run')
                if basename(row['Source_file']) != row['Image_name']:
                    raise ValidationError('Source filename differs from Image_name')
                if (Path(row['Image_name']).suffix.lower() not in ('.nd2', '.tif', '.tiff')
                        or Path(row['Image_name']).stem + MASK_SUFFIX != Path(mask).stem):
                    raise ValidationError('Source filename does not identify the corresponding mask')
            if table == 'Images':
                if mask in images:
                    raise ValidationError(f'Duplicate image mask: {mask}')
                if row['Status'] != 'complete' or row['Error'].strip():
                    raise ValidationError(f'Unsuccessful image row: {mask}')
                counts = [number(row[key], key, integer=True) for key in COUNTS]
                if min(counts) < 0 or counts[0] - counts[1] != counts[2]:
                    raise ValidationError(f'Inconsistent nucleus counts: {mask}')
                if kind == 'intensity' and (
                        number(row['Nuclei_measured_count'], 'Nuclei_measured_count', integer=True) != counts[2]
                        or number(row['Failed_nuclei_count'], 'Failed_nuclei_count', integer=True) != 0):
                    raise ValidationError(f'Incomplete nucleus measurements: {mask}')
                metrics = MORPH_METRICS if kind == 'morphology' else ('Area_px2', *INTENSITY_METRICS)
                for metric in metrics:
                    for stat in STATS:
                        value = number(row[f'{metric}_{stat}'], f'{metric}_{stat}', blank=True)
                        if counts[2] == 0 and value is not None:
                            raise ValidationError(f'Nonempty summary for zero eligible nuclei: {mask}')
                        if counts[2] > 0 and value is None and (kind == 'intensity' or metric == 'Area_px2'):
                            raise ValidationError(f'Missing summary for measured nuclei: {mask}')
                images[mask] = row
            else:
                nucleus_id = number(row['Nucleus_ID'], 'Nucleus_ID', integer=True)
                area = number(row['Area_px2'], 'Area_px2', integer=True)
                if nucleus_id < 1 or area < 1:
                    raise ValidationError('Nucleus IDs and areas must be positive')
                key = (mask, nucleus_id)
                if key in nuclei:
                    raise ValidationError(f'Duplicate nucleus: {key}')
                if kind == 'morphology' and row['Touches_border'].lower() not in ('false', '0'):
                    raise ValidationError(f'Border nucleus in the measurement table: {key}')
                metrics = MORPH_METRICS if kind == 'morphology' else INTENSITY_METRICS
                for metric in metrics:
                    number(row[metric], metric, blank=kind == 'morphology' and metric != 'Area_px2')
                nuclei[key] = row
    for mask, row in images.items():
        if sum(key[0] == mask for key in nuclei) != number(row['Non_border_nuclei_count'], 'count', integer=True):
            raise ValidationError(f'Nucleus-row count differs from image summary: {mask}')
    if any(mask not in images for mask, _ in nuclei):
        raise ValidationError('Nucleus row has no corresponding image row')
    bundle.update(namespace=namespace, images=images, nuclei=nuclei, particle_size=particle_size)


def load_morphology(run):
    files = {}
    meta = json_snapshot(run / 'nuclei_run.json', files)
    if (meta.get('schema_version') != 1 or meta.get('status') != 'complete'
            or meta.get('morphology_status') != 'complete' or meta.get('skipped_files')):
        raise ValidationError('Nuclei and morphology runs must both be complete')
    tables, info, copies = read_tables(run, 'morphology', files)
    bundle = {'run': run, 'metadata': meta, 'tables': tables, 'files': files, 'copies': copies}
    validate_rows(bundle, 'morphology')
    if info.get('Run_ID') != run.name or number(info.get('Particle_size_px2'), 'Run_Info particle size') != bundle['particle_size']:
        raise ValidationError('Morphology Run_Info identifies a different run or particle size')
    processed = meta.get('processed_files')
    if not isinstance(processed, list) or not processed or not all(
            isinstance(name, str) and basename(name) == name and not name.startswith('.') for name in processed):
        raise ValidationError('Invalid processed_files manifest')
    expected = {Path(name).stem + '_processed.tif' for name in processed}
    if len(expected) != len(processed) or expected != set(bundle['images']):
        raise ValidationError('Morphology images differ from the completed nuclei manifest')
    return bundle


def intensity_candidate(run):
    files = {}
    match = INTENSITY_PATTERN.fullmatch(run.name)
    if match is None:
        raise ValidationError('Unrecognized intensity folder name')
    datetime.strptime(match[2], '%Y%m%d_%H%M%S').replace(tzinfo=timezone.utc)
    meta = json_snapshot(run / 'intensity_run.json', files)
    if meta.get('Schema_version') not in (1, 2):
        raise ValidationError('Unsupported intensity metadata schema')
    channel = number(meta.get('Marker_channel'), 'Marker_channel', integer=True)
    if channel < 1:
        raise ValidationError('Marker channel must be positive')
    marker = meta.get('Marker_folder')
    if marker is None and meta['Schema_version'] == 1 and match[1] is None:
        marker = f'Channel_{channel}'
    elif not isinstance(marker, str) or not MARKER_PATTERN.fullmatch(marker):
        raise ValidationError('Invalid marker folder in intensity metadata')
    if marker.startswith('Foci_') and int(MARKER_PATTERN.fullmatch(marker)[1]) != channel:
        raise ValidationError('Marker folder/channel conflict')
    if match[1] is not None and match[1] != marker:
        raise ValidationError('Marker metadata differs from the intensity folder name')
    if not NUCLEI_PATTERN.fullmatch(basename(meta.get('Nuclei_run', ''))):
        raise ValidationError('Missing or invalid nucleus-run reference')
    return {'run': run, 'metadata': meta, 'files': files, 'marker': marker,
            'channel': channel, 'timestamp': match[2]}


def load_intensity(candidate):
    bundle = {**candidate, 'files': dict(candidate['files'])}
    run, meta = bundle['run'], bundle['metadata']
    if meta.get('Status') != 'complete':
        raise ValidationError('Intensity run is not complete')
    if basename(meta.get('Intensity_run', '')) != run.name:
        raise ValidationError('Intensity metadata identifies another run')
    tables, info, copies = read_tables(run, 'intensity', bundle['files'])
    bundle.update(tables=tables, copies=copies)
    validate_rows(bundle, 'intensity')
    if (parts(meta['Nuclei_run']) != bundle['namespace'] + ('foci_assay', basename(meta['Nuclei_run']))
            or parts(info.get('Nuclei_run', '')) != parts(meta['Nuclei_run'])
            or basename(info.get('Intensity_run', '')) != run.name
            or number(info.get('Marker_channel'), 'Run_Info channel', integer=True) != bundle['channel']
            or number(info.get('Particle_size_px2'), 'Run_Info particle size') != bundle['particle_size']):
        raise ValidationError('Intensity Run_Info/dataset namespace differs from metadata')
    if meta.get('Schema_version') == 2 and info.get('Marker_folder') != bundle['marker']:
        raise ValidationError('Intensity Run_Info marker differs from metadata')
    return bundle


def validate_compatibility(morphology, intensity):
    if (intensity['namespace'] != morphology['namespace']
            or basename(intensity['metadata']['Nuclei_run']) != morphology['run'].name
            or intensity['particle_size'] != morphology['particle_size']):
        raise ValidationError('Intensity refers to another dataset, nucleus run or particle size')
    if set(morphology['images']) != set(intensity['images']):
        raise ValidationError('Image sets differ between morphology and intensity')
    if set(morphology['nuclei']) != set(intensity['nuclei']):
        raise ValidationError('Nucleus ID sets differ between morphology and intensity')
    for key, row in morphology['nuclei'].items():
        if number(row['Area_px2'], 'area', integer=True) != number(intensity['nuclei'][key]['Area_px2'], 'area', integer=True):
            raise ValidationError(f'Nucleus area differs between morphology and intensity: {key}')
    for mask, row in morphology['images'].items():
        for field in COUNTS:
            if number(row[field], field, integer=True) != number(intensity['images'][mask][field], field, integer=True):
                raise ValidationError(f'{field} differs between morphology and intensity: {mask}')


def info_row(category, item, status, source='', details='', digest=''):
    return dict(zip(INFO_COLUMNS, (category, item, status, str(source), str(details), digest)))


def discover_sources(input_path):
    path = Path(input_path).expanduser().resolve()
    if path.name.startswith('.'):
        raise ValidationError('Hidden input manifests are not accepted')
    payload = json.loads(path.read_text(encoding='utf-8-sig'))
    folders = payload.get('paths_to_files') if isinstance(payload, dict) else None
    if not isinstance(folders, list) or not folders or not all(isinstance(p, str) and p.strip() for p in folders):
        raise ValidationError('Input JSON must contain a nonempty paths_to_files list')
    result = []
    for folder in folders:
        root = Path(folder).expanduser().resolve()
        if root in result:
            continue
        if root.name.startswith('.') or not root.is_dir():
            print(f'Skipping missing or hidden experiment: {root}')
            continue
        result.append(root)
    return result


def choose_indices(prompt, count, allow_none=False):
    while True:
        answer = input(prompt).strip().lower()
        if answer == 'q':
            raise Cancelled()
        if allow_none and answer in ('none', 'n'):
            return []
        if answer in ('all', 'a'):
            return list(range(count))
        try:
            indices = sorted({int(value.strip()) - 1 for value in answer.split(',')})
            if indices and min(indices) >= 0 and max(indices) < count:
                return indices
        except ValueError:
            pass
        print(f'Enter numbers from 1 to {count}, all, ' + ('none, ' if allow_none else '') + 'or q.')


def scan_source(root):
    context = {'root': root, 'morphology': [], 'intensity': [], 'markers': set(), 'checks': []}
    assay = root / 'foci_assay'
    if not assay.is_dir() or assay.is_symlink():
        print(f'No regular foci_assay directory: {root}')
        return context
    for run in sorted(assay.iterdir(), reverse=True):
        if run.name.startswith('.') or not run.is_dir() or run.is_symlink():
            continue
        try:
            match = NUCLEI_PATTERN.fullmatch(run.name)
            if match:
                datetime.strptime(match[1], '%Y%m%d_%H%M%S').replace(tzinfo=timezone.utc)
                bundle = load_morphology(run)
                bundle['root'] = root
                context['morphology'].append(bundle)
                print(f"Valid morphology: {run.name} | -p {bundle['particle_size']:g} "
                      f"| images {len(bundle['images'])} | non-border nuclei {len(bundle['nuclei'])}")
            elif run.name.startswith('Nuclear_Intensity_'):
                candidate = intensity_candidate(run)
                context['intensity'].append(candidate)
                context['markers'].add(candidate['marker'])
        except (OSError, ValueError, csv.Error, KeyError) as error:
            context['checks'].append(info_row('Discovery', run.name, 'REJECTED', run, error))
            print(f'Unavailable: {run}: {error}')
    foci = assay / 'Foci'
    if foci.is_dir() and not foci.is_symlink():
        for folder in foci.iterdir():
            if folder.is_dir() and not folder.is_symlink() and MARKER_PATTERN.fullmatch(folder.name):
                context['markers'].add(folder.name)
    return context


def select_morphologies(contexts):
    bundles = [bundle for context in contexts for bundle in context['morphology']]
    if not bundles:
        raise ValidationError('No completed valid morphology spreadsheets were found')
    print('\nChoose nucleus runs: 1 = latest valid per experiment; 2 = all valid; 3 = manual; q = cancel.')
    while True:
        answer = input('Nucleus-run selection: ').strip().lower()
        if answer == 'q':
            raise Cancelled()
        if answer == '1':
            return [max(context['morphology'], key=lambda b: b['run'].name)
                    for context in contexts if context['morphology']]
        if answer == '2':
            return bundles
        if answer == '3':
            for index, bundle in enumerate(bundles, 1):
                print(f"{index}. {bundle['root']} / {bundle['run'].name} "
                      f"| -p {bundle['particle_size']:g}")
            return [bundles[index] for index in choose_indices('Select run numbers, all, or q: ', len(bundles))]
        print('Enter 1, 2, 3, or q.')


def select_markers(bundles, contexts):
    roots = {bundle['root'] for bundle in bundles}
    names = sorted(set().union(*(context['markers'] for context in contexts if context['root'] in roots)))
    if not names:
        print('No marker results or marker folders found; collecting morphology only.')
        return []
    print('\nMarkers available for collection:')
    for index, marker in enumerate(names, 1):
        print(f'{index}. {marker}')
        for bundle in bundles:
            candidates = contexts_by_root(contexts)[bundle['root']]['intensity']
            count = sum(c['marker'] == marker and c['metadata'].get('Status') == 'complete'
                        and basename(c['metadata']['Nuclei_run']) == bundle['run'].name for c in candidates)
            print(f"   {bundle['root']} / {bundle['run'].name}: {count} completed intensity runs")
    return [names[index] for index in choose_indices(
        'Select marker numbers, all, none (morphology only), or q: ', len(names), allow_none=True)]


def contexts_by_root(contexts):
    return {context['root']: context for context in contexts}


def latest_intensity(context, morphology, marker, checks):
    candidates = sorted((c for c in context['intensity'] if c['marker'] == marker
                         and basename(c['metadata']['Nuclei_run']) == morphology['run'].name),
                        key=lambda c: c['timestamp'], reverse=True)
    for candidate in candidates:
        try:
            bundle = load_intensity(candidate)
        except (OSError, ValueError, csv.Error, KeyError) as error:
            checks.append(info_row('Intensity selection', marker, 'REJECTED', candidate['run'], error))
            print(f"Skipping intensity result {candidate['run']}: {error}")
            continue
        # Never search for older matching data after a cross-analysis conflict.
        validate_compatibility(morphology, bundle)
        checks.append(info_row('Intensity selection', marker, 'SELECTED', bundle['run'],
                               'Latest valid intensity table for the selected nucleus run'))
        return bundle
    checks.append(info_row('Intensity selection', marker, 'MISSING', details=
                           'No completed valid intensity table for the selected nucleus run'))
    print(f"Missing marker intensity: {morphology['run']} / {marker}")
    return None


def typed_morphology(row):
    numeric = set(MORPH_METRICS) | set(COUNTS) | {
        'Nucleus_ID', 'Particle_size_px2', 'Orientation_deg', 'Centroid_X_px', 'Centroid_Y_px',
    } | {f'{metric}_{stat}' for metric in MORPH_METRICS for stat in STATS}
    integers = set(COUNTS) | {'Nucleus_ID', 'Area_px2'}
    return {('Nuclei_run_ID' if key == 'Run_ID' else key):
            (number(value, key, integer=key in integers, blank=True) if key in numeric
             else (value.lower() in ('true', '1') if key == 'Touches_border' else value))
            for key, value in row.items()}


def combine(morphology, selected, markers, checks):
    """Join by mask/ID within one dataset and nucleus run, never by display image name."""
    source_map = {}
    for marker, bundle in selected.items():
        for mask, row in bundle['images'].items():
            identity = (row['Image_name'], parts(row['Source_file']))
            if mask in source_map and source_map[mask][0] != identity:
                raise ValidationError(f'Source image differs between selected markers: {mask} ({marker})')
            source_map[mask] = (identity, row['Source_file'])
    originals = [path for path in morphology['root'].iterdir()
                 if not path.name.startswith('.') and path.is_file()
                 and path.suffix.lower() in ('.nd2', '.tif', '.tiff')]
    for mask in morphology['images']:
        if mask in source_map:
            continue
        stem = Path(mask).stem
        matches = [path for path in originals if stem == path.stem + MASK_SUFFIX]
        if len(matches) == 1:
            source_map[mask] = ((matches[0].name, parts(matches[0])), str(matches[0]))
        else:
            checks.append(info_row('Image identity', mask, 'UNRESOLVED', morphology['run'],
                                   'Original filename unavailable or ambiguous; morphology name is retained separately'))
    output = {}
    for sheet in ('Nuclei', 'Images'):
        original_columns = morphology['tables'][sheet][0]
        columns = [('Nuclei_run_ID' if col == 'Run_ID' else col) for col in original_columns]
        columns += ['Morphology_Image_name', 'Source_file']
        records = []
        metric_columns = list(INTENSITY_METRICS) if sheet == 'Nuclei' else [
            f'{metric}_{stat}' for metric in INTENSITY_METRICS for stat in STATS]
        if sheet == 'Images':
            metric_columns += ['Nuclei_measured_count', 'Failed_nuclei_count']
        for marker in markers:
            columns += [f'{marker}_{field}' for field in ('Availability', 'Intensity_run_ID', *metric_columns)]
        for key, row in (morphology['nuclei'] if sheet == 'Nuclei' else morphology['images']).items():
            mask = key[0] if sheet == 'Nuclei' else key
            identity = source_map.get(mask)
            record = typed_morphology(row)
            record.update(Morphology_Image_name=row['Image_name'],
                          Image_name=identity[0][0] if identity else '', Source_file=identity[1] if identity else '')
            for marker in markers:
                bundle = selected.get(marker)
                record[f'{marker}_Availability'] = 'AVAILABLE' if bundle else 'MISSING'
                record[f'{marker}_Intensity_run_ID'] = bundle['run'].name if bundle else ''
                if bundle:
                    values = bundle['nuclei' if sheet == 'Nuclei' else 'images'][key]
                    for field in metric_columns:
                        record[f'{marker}_{field}'] = number(values[field], field, blank=True,
                            integer=field in ('Nuclei_measured_count', 'Failed_nuclei_count'))
            records.append(record)
        output[sheet] = (columns, records)
    return output


def write_workbook(path, tables):
    workbook = Workbook(write_only=True)
    for name, (columns, rows) in tables.items():
        if len(columns) > 16384 or len(rows) + 1 > 1048576:
            raise ValidationError(f'Table exceeds Excel limits: {name}')
        sheet = workbook.create_sheet(name)
        sheet.freeze_panes = 'A2'
        sheet.append(columns)
        for row in rows:
            cells = []
            for column in columns:
                value = row.get(column)
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = 's'
                cells.append(cell)
            sheet.append(cells)
    workbook.save(path)


def write_csv(path, columns, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def create_output(root):
    stamp = datetime.now().astimezone()
    while True:
        output = root / (OUTPUT_PREFIX + stamp.strftime('%Y%m%d_%H%M%S'))
        try:
            output.mkdir()
            return output
        except FileExistsError:
            stamp += timedelta(seconds=1)


def ignore_missing_sidecar(function, path, exc_info):
    """macOS may remove an AppleDouble companion when its main file is deleted."""
    if isinstance(exc_info[1], FileNotFoundError) and Path(path).name.startswith('._'):
        return
    raise exc_info[1]


def publish_staging(staging, root, final_name, filenames):
    """Move only the expected spreadsheets, with the completion workbook last."""
    names = list(filenames)
    if (final_name not in names or len(set(names)) != len(names)
            or any(name.startswith('.') or Path(name).name != name
                   or Path(name).suffix not in ('.csv', '.xlsx') for name in names)):
        raise ValidationError('Invalid publication file list')
    paths = sorted((staging / name for name in names), key=lambda path: path.name == final_name)
    output = create_output(root)
    try:
        for path in paths:
            if path.is_symlink() or not path.is_file():
                raise FileNotFoundError(f'Missing or nonregular expected spreadsheet: {path}')
            path.replace(output / path.name)
        return output
    except BaseException:
        # Only remove this newly allocated collector output, never a previous run.
        try:
            shutil.rmtree(output, onerror=ignore_missing_sidecar)
        except OSError as cleanup_error:
            # Preserve the publication error even if rollback also fails.
            print(f'Could not fully remove incomplete output {output}: {cleanup_error}')
        raise


def copy_name(path, marker=None):
    if marker is None:
        return path.name
    return f'{path.stem}_{marker}{path.suffix}'


def collect_one(morphology, context, markers, manifest):
    root = morphology['root']
    checks = list(context['checks'])
    checks += [info_row('Collection', 'Source folder', 'SELECTED', root),
               info_row('Collection', 'Input manifest', 'SELECTED', manifest),
               info_row('Morphology selection', morphology['run'].name, 'SELECTED', morphology['run'],
                        f"Particle_size_px2={morphology['particle_size']:g}"),
               info_row('Policy', 'Population', 'INFO', details='Non-border nuclei only; no new filtering or statistics'),
               info_row('Policy', 'Join', 'INFO', details='Dataset, nucleus run, mask basename, Nucleus_ID; Area_px2 verified'),
               info_row('Policy', 'Missing data', 'INFO', details='Missing marker measurements remain blank, never zero'),
               info_row('Policy', 'Long metadata', 'INFO', details='Run_Info text may be truncated by Excel at 32767 characters; original CSV retains the full value'),
               info_row('Policy', 'Source copies', 'INFO', details='Original spreadsheet bytes are preserved; paths refer to their original runs')]
    selected, snapshots = {}, dict(morphology['files'])
    try:
        for marker in markers:
            bundle = latest_intensity(context, morphology, marker, checks)
            if bundle:
                selected[marker] = bundle
                snapshots.update(bundle['files'])
        tables = combine(morphology, selected, markers, checks)
        status = 'SUCCESS_WITH_MISSING_INTENSITY' if len(selected) != len(markers) else 'SUCCESS'
        checks.append(info_row('Collection', 'Status', status, details=datetime.now(timezone.utc).isoformat()))
        checks.append(info_row('Collection', 'Rows', 'INFO', details=
                              f"Images={len(morphology['images'])}; non-border nuclei={len(morphology['nuclei'])}"))
        copies = [(path, copy_name(path)) for path in morphology['copies']]
        copies += [(path, copy_name(path, marker)) for marker, bundle in selected.items() for path in bundle['copies']]
        for path, data in snapshots.items():
            checks.append(info_row('Source snapshot', path.name, 'VALIDATED', path,
                                   f'{len(data)} bytes', hashlib.sha256(data).hexdigest()))
        for path, target in copies:
            checks.append(info_row('Copied spreadsheet', target, 'COPIED', path,
                                   digest=hashlib.sha256(snapshots[path]).hexdigest()))
        tables['Collection_Info'] = (INFO_COLUMNS, checks)
        with tempfile.TemporaryDirectory(prefix='.fia_marker_collection_', dir=root) as temporary:
            staging = Path(temporary)
            for path, target in copies:
                (staging / target).write_bytes(snapshots[path])
            for sheet in ('Nuclei', 'Images'):
                write_csv(staging / f'FIA_Marker_Intensity_{sheet}.csv', *tables[sheet])
            write_workbook(staging / COMBINED_NAME, tables)
            for path, data in snapshots.items():
                if path.is_symlink() or path.read_bytes() != data:
                    raise ValidationError(f'Source changed during collection: {path}')
            filenames = [target for _, target in copies] + [
                'FIA_Marker_Intensity_Nuclei.csv', 'FIA_Marker_Intensity_Images.csv', COMBINED_NAME]
            output = publish_staging(staging, root, COMBINED_NAME, filenames)
        print(f'{status}: {output}')
        return True, output
    except (OSError, ValueError, csv.Error, KeyError) as error:
        checks = [row for row in checks if not (row['Category'] == 'Collection' and row['Item'] == 'Status')
                  and row['Category'] != 'Copied spreadsheet']
        status = 'IO_FAILED' if isinstance(error, OSError) else 'VALIDATION_FAILED'
        checks.append(info_row('Collection', 'Status', status, morphology['run'], error))
        # A failed collection contains a diagnostic spreadsheet only.
        with tempfile.TemporaryDirectory(prefix='.fia_marker_collection_', dir=root) as temporary:
            staging = Path(temporary)
            write_workbook(staging / 'Collection_Report.xlsx', {'Collection_Info': (INFO_COLUMNS, checks)})
            output = publish_staging(staging, root, 'Collection_Report.xlsx', ['Collection_Report.xlsx'])
        print(f'{status}: {error}. Report: {output}')
        return False, output


def main(input_path):
    try:
        roots = discover_sources(input_path)
        if not roots:
            raise ValidationError('No existing visible experiment folders')
        for index, root in enumerate(roots, 1):
            print(f'{index}. {root}')
        roots = [roots[index] for index in choose_indices('Select experiments, all, or q: ', len(roots))]
        contexts = [scan_source(root) for root in roots]
        bundles = select_morphologies(contexts)
        markers = select_markers(bundles, contexts)
        print(f'\nCollecting {len(bundles)} separate nucleus-run collections; markers: {markers or "none"}.')
        failures = 0
        for bundle in bundles:
            try:
                success, _ = collect_one(bundle, contexts_by_root(contexts)[bundle['root']], markers, input_path)
                failures += not success
            except (OSError, ValueError) as error:
                print(f"Cannot write collection for {bundle['root']}: {error}")
                failures += 1
        return 1 if failures else 0
    except (Cancelled, KeyboardInterrupt, EOFError):
        print('Collection canceled. A completed collection is identified by its combined workbook.')
        return 130
    except (OSError, ValueError, csv.Error) as error:
        print(f'Marker-intensity collection could not complete: {error}')
        return 1
