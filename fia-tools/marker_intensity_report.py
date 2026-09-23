"""Create a spreadsheet-based marker-intensity report without starting ImageJ."""

import argparse
import hashlib
import json
import logging
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import collect_marker_intensity_results as collect
import marker_report_data as inputs
import matplotlib
import numpy as np
import openpyxl
import scipy
from marker_report_plots import PLOT_COLUMNS, render_plots
from marker_report_statistics import STAT_COLUMNS, calculate_statistics, observations
from openpyxl.cell import WriteOnlyCell
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_PREFIX = 'FIA_Marker_Intensity_Report_'
WORKBOOK = 'FIA_Marker_Intensity_Report.xlsx'
MORPHOLOGY_WORKBOOK = 'Nuclei_Morphology_Summary.xlsx'
MORPHOLOGY_SHEETS = ('Morphology_By_Condition', 'Morphology_Comparisons')


def create_output(root):
    stamp = datetime.now().astimezone()
    while True:
        path = Path(root) / (OUTPUT_PREFIX + stamp.strftime('%Y%m%d_%H%M%S'))
        try:
            path.mkdir()
            return path
        except FileExistsError:
            stamp += timedelta(seconds=1)


def write_status(output, status, **details):
    payload = dict(Schema_version=1, Status=status, Updated_UTC=datetime.now(timezone.utc).isoformat(), **details)
    temporary = output / '.report_status.partial.json'
    temporary.write_text(json.dumps(payload, indent=2, allow_nan=False), encoding='utf-8')
    temporary.replace(output / 'report_status.json')


def discover_collections(root):
    candidates = []
    for path in sorted(Path(root).iterdir()):
        if path.is_symlink() or not path.is_dir() or not inputs.COLLECTION_PATTERN.fullmatch(path.name):
            continue
        try:
            info = inputs.collection_info(path, {})
            run = next((row['Item'] for row in info if row['Category'] == 'Morphology selection'), 'unknown run')
            candidates.append((path, run))
        except (OSError, ValueError) as error:
            print(f'Skipping incomplete collection {path.name}: {error}')
    return candidates


def choose_collections(roots, mode):
    contexts = {root: discover_collections(root) for root in roots}
    missing = [str(root) for root, candidates in contexts.items() if not candidates]
    for root in missing:
        print(f'No completed collections: {root}')
    candidates = [item for values in contexts.values() for item in values]
    if not candidates:
        raise inputs.ValidationError('Run fia_collect_marker_intensity_results first: no completed collections found')
    for index, (path, run) in enumerate(candidates, 1):
        print(f'{index}. {path} | {run}')
    if mode == 'ask':
        while True:
            answer = input('Collections: 1 = latest per experiment; 2 = all; 3 = manual; q = cancel: ').strip().lower()
            if answer == 'q':
                raise collect.Cancelled()
            if answer in ('1', '2', '3'):
                mode = {'1': 'latest', '2': 'all', '3': 'manual'}[answer]
                break
            print('Enter 1, 2, 3, or q.')
    if mode == 'latest':
        selected = [values[-1][0] for values in contexts.values() if values]
    elif mode == 'all':
        selected = [path for path, _ in candidates]
    else:
        selected = [candidates[index][0] for index in collect.choose_indices(
            'Select collection numbers, all, or q: ', len(candidates))]
    return selected, missing


def choose_markers(data, selection):
    catalog, notes = inputs.marker_catalog(data)
    data['notes'].extend(notes)
    for note in notes:
        print(note)
    if isinstance(selection, (list, tuple)):
        return list(selection)
    if selection is not None:
        if selection == 'all':
            return catalog
        if selection == 'none':
            return []
        return [name.strip() for name in selection.split(',') if name.strip()]
    if not catalog:
        print('No marker columns: morphology-only report.')
        return []
    for index, marker in enumerate(catalog, 1):
        print(f'{index}. {marker}: {data["availability"][marker]}')
    return [catalog[index] for index in collect.choose_indices(
        'Select markers, all, none (morphology only), or q: ', len(catalog), allow_none=True)]


def as_table(rows, columns=()):
    return list(columns) or list(dict.fromkeys(key for row in rows for key in row)), rows


def morphology_tables(data):
    """Present morphology separately, using the same observations and tests as the report."""
    statistics = ('N', 'Mean', 'SD', 'Median', 'IQR')
    columns = ['Metric', 'Unit', 'Observation_unit']
    columns += [f'{group}: {stat}' for group in data['groups'] for stat in statistics]
    rows = []
    for category, _, field, unit in inputs.metric_specs([]):
        if category != 'Morphology':
            continue
        row = {'Metric': field, 'Unit': unit,
               'Observation_unit': 'well' if data['stats_unit'] == 'well' else 'nucleus'}
        for group in data['groups']:
            values, _ = observations(data, field, group)
            summaries = {
                'N': len(values),
                'Mean': float(np.mean(values)) if values else None,
                'SD': float(np.std(values, ddof=1)) if len(values) >= 2 else None,
                'Median': float(np.median(values)) if values else None,
                'IQR': float(np.percentile(values, 75) - np.percentile(values, 25)) if values else None,
            }
            row.update({f'{group}: {stat}': summaries[stat] for stat in statistics})
        rows.append(row)
    # Reuse the full report's Holm results without creating a new test family.
    comparison_columns = [column for column in STAT_COLUMNS if column not in ('Category', 'Marker')]
    comparisons = [{column: row[column] for column in comparison_columns}
                   for row in data['statistics'] if row['Category'] == 'Morphology']
    return {MORPHOLOGY_SHEETS[0]: (columns, rows),
            MORPHOLOGY_SHEETS[1]: (comparison_columns, comparisons)}


def report_tables(data, output, manifest):
    unit = data['stats_unit'] or 'disabled'
    notes = [
        ('Population', 'Non-border nuclei only. No extra morphology or intensity filtering.'),
        ('Count plot', 'One point = non-border nuclei in one image, including zero-count images.'),
        ('Intensity plots', 'One point = one nucleus; original Marker_RawIntDen; no averaging or normalization.'),
        ('Morphology plots', 'Add a boxplot for each morphology metric with at least one tested control comparison having P_Holm < 0.05. Show all conditions and comparisons for that metric; one point = one non-border nucleus.'),
        ('Statistics unit', unit),
        ('Count test exception', 'Requested nucleus mode uses images for count tests; well mode uses wells.'),
        ('Well aggregation', 'Mean per image across usable nuclei, then mean of usable image means per well. Equal image weight.'),
        ('Missing data', 'Missing marker measurements are blank, not zero. Empty images contribute zero to counts only.'),
        ('Tests', 'Two-sided Welch comparisons of each treatment to the bold control in its plate-map color block.'),
        ('Multiplicity', 'Holm family includes count, 12 morphology metrics and six metrics per selected marker, across all treatments in a color block, including unavailable planned tests.'),
        ('Confidence intervals', '95% Welch intervals for treatment minus control; nominal, not multiplicity-adjusted.'),
        ('Minimum observations', 'At least two usable units per condition. Both groups constant: not tested.'),
        ('Dependence', 'Nucleus/image tests are exploratory and do not account for within-well clustering. Holm does not correct this dependence.'),
        ('Replicates', 'Wells are within-plate observations, not inferred biological replicates; experiments and runs are never pooled.'),
        ('Summary fields', 'Median and IQR columns describe distributions; no separate tests of these summary columns.'),
        ('Morphology summary', 'Nuclei_Morphology_Summary.xlsx contains all 12 morphology metrics side by side by condition, existing control comparisons and Run_Info. Significant morphology plots are in the main report and Plots folder; no additional tests.'),
        ('Size units', 'Area_px2 and pixel-based morphology; no rescaling or conversion to micrometers.'),
        ('Marker identity', 'Selected folder names are used verbatim. No inferred biological marker names or legacy alias merging.'),
        ('Completion', 'Use this report only when report_status.json says SUCCESS.'),
    ]
    notes += [('Selection note', note) for note in data['notes']]
    provenance = []
    for path, content in data['files'].items():
        if path == data['template']:
            relative = Path('Inputs') / 'PlateMap' / path.name
        elif manifest is not None and path == manifest:
            relative = Path('Inputs') / 'input_paths.json'
        else:
            relative = Path('Inputs') / 'Collection' / path.name
        destination = output / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(content)
        provenance.append({'Source': str(path), 'Archived_copy': str(relative), 'Bytes': len(content),
                           'SHA256': hashlib.sha256(content).hexdigest()})
    metadata = {
        'Report_schema_version': 1, 'Created_UTC': datetime.now(timezone.utc).isoformat(),
        'Collection': str(data['path']), 'Nuclei_run_ID': data['run_id'],
        'Particle_size_px2': data['particle_size'], 'Markers': ', '.join(data['markers']),
        'Plate_map': str(data['template']), 'Plate_map_sheet': data['template_sheet'],
        'Requested_statistics_unit': unit, 'Non_border_nuclei': len(data['nuclei']),
        'Morphology_summary_observation_unit': 'well' if data['stats_unit'] == 'well' else 'nucleus',
        'Morphology_summary_population': 'Non-border nuclei only; no additional filtering.',
        'Morphology_summary_N': 'Usable observations for each metric, in Morphology_summary_observation_unit.',
        'Morphology_summary_SD': 'Sample standard deviation (ddof=1); blank for fewer than two observations.',
        'Morphology_summary_well_values': 'Mean of usable per-image nucleus means within each well; equal image weight.',
        'Morphology_summary_tests': 'Existing Welch/Holm results from Statistics, including its full planned family across count, morphology and selected markers. No tests when Requested_statistics_unit is disabled.',
        'Morphology_plot_rule': 'At least one TESTED control comparison with P_Holm < 0.05 in the selected statistics unit; no morphology plots when statistics are disabled or no comparisons qualify.',
        'Images': len(data['images']), 'Python': sys.version.split()[0],
        'NumPy': np.__version__, 'SciPy': scipy.__version__, 'Matplotlib': matplotlib.__version__,
        'openpyxl': openpyxl.__version__,
    }
    return {
        'Overview': as_table([{'Item': key, 'Description': value} for key, value in notes]),
        **morphology_tables(data),
        'Statistics': as_table(data['statistics'], STAT_COLUMNS),
        'Summary': as_table(data['summary']),
        'Nuclei': as_table(data['nuclei'], data['nuclei_columns'] + ['Group']),
        'Images': as_table(data['images']),
        'Image_Values': as_table(data['image_values']),
        'Well_Values': as_table(data['well_values']),
        'Plate_Map': as_table(data['design']),
        'Plot_Data': as_table(data['plot_data'], PLOT_COLUMNS),
        'Plot_Info': as_table(data['plots']),
        'Run_Info': as_table([{'Parameter': key, 'Value': value} for key, value in metadata.items()]),
        'Source_Files': as_table(provenance),
    }


def write_workbook(output, tables, plots, filename=WORKBOOK):
    book = openpyxl.Workbook(write_only=True)
    for title, (columns, rows) in tables.items():
        if len(columns) > 16384 or len(rows) + 1 > 1048576:
            raise inputs.ValidationError(f'Table exceeds Excel limits: {title}')
        sheet = book.create_sheet(title)
        sheet.freeze_panes = 'D2' if title == MORPHOLOGY_SHEETS[0] else 'A2'
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{len(rows) + 1}'
        for index, column in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(45, max(14, len(column) + 2))
        cells = []
        for value in columns:
            cell = WriteOnlyCell(sheet, value)
            cell.data_type = 's'
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='234F68')
            if title in MORPHOLOGY_SHEETS:
                cell.alignment = Alignment(wrap_text=True, vertical='center')
            cells.append(cell)
        if title in MORPHOLOGY_SHEETS:
            sheet.row_dimensions[1].height = 42
        sheet.append(cells)
        for row in rows:
            cells = []
            for column in columns:
                value = row.get(column)
                if isinstance(value, str) and len(value) > 32767:
                    raise inputs.ValidationError(f'Text exceeds Excel cell limit: {title}/{column}')
                cell = WriteOnlyCell(sheet, value)
                if isinstance(value, str):
                    cell.data_type = 's'
                elif title == MORPHOLOGY_SHEETS[0] and value is not None:
                    cell.number_format = '0' if column.endswith(': N') else '0.000'
                cells.append(cell)
            sheet.append(cells)
    if plots:
        sheet = book.create_sheet('Plots')
        anchor = 1
        for plot in plots:
            image = Image(output / plot['File'])
            height = image.height * min(1, 1050 / image.width)
            image.width = min(1050, image.width)
            image.height = height
            sheet.add_image(image, f'A{anchor}')
            anchor += int(height / 20) + 3
    temporary = output / ('.' + filename + '.partial.xlsx')
    book.save(temporary)
    temporary.replace(output / filename)


def create_report(collection, root, template=None, markers=None, stats_unit=None, sheet=None, manifest=None):
    output = create_output(root)
    logger = logging.getLogger(f'fia_marker_report.{output.name}')
    logger.setLevel(logging.INFO)
    logger.propagate = False
    handler = logging.FileHandler(output / 'report.log', encoding='utf-8')
    handler.setFormatter(logging.Formatter('%(asctime)s %(levelname)s %(message)s'))
    logger.addHandler(handler)
    stage = 'input validation'
    try:
        write_status(output, 'RUNNING', Collection=str(collection))
        logger.info('Starting report for %s; statistics=%s', collection, stats_unit or 'disabled')
        data = inputs.load_collection(collection)
        selected = choose_markers(data, markers)
        plate_map = Path(template).expanduser().absolute() if template else inputs.find_template(Path(collection), sheet)
        if plate_map.name.startswith(('.', '~$')):
            raise inputs.ValidationError('Hidden or temporary plate maps are not accepted')
        # Keep headers even when every image has zero non-border nuclei.
        excluded = set(data['markers']) - set(selected)
        source_columns = collect.csv_snapshot(Path(collection) / 'FIA_Marker_Intensity_Nuclei.csv', data['files'], [])[0]
        data['nuclei_columns'] = [c for c in source_columns if not any(c.startswith(m + '_') for m in excluded)]
        inputs.annotate(data, plate_map, selected, stats_unit, sheet)
        manifest_path = Path(manifest).expanduser().absolute() if manifest else None
        if manifest_path:
            collect.snapshot(manifest_path, data['files'])
        stage = 'aggregation and statistics'
        inputs.aggregate(data)
        calculate_statistics(data)
        stage = 'plots'
        render_plots(data, output / 'Plots')
        stage = 'spreadsheet export'
        tables = report_tables(data, output, manifest_path)
        for title, table in tables.items():
            collect.write_csv(output / (title + '.csv'), *table)
        write_workbook(output, tables, data['plots'])
        morphology = {title: tables[title] for title in (*MORPHOLOGY_SHEETS, 'Run_Info')}
        write_workbook(output, morphology, [], MORPHOLOGY_WORKBOOK)
        # Verify exact published tables and unchanged inputs before marking success.
        stage = 'final verification'
        exported = {title: collect.csv_snapshot(output / (title + '.csv'), {}, columns)
                    for title, (columns, _) in tables.items()}
        collect.verify_workbook(output / WORKBOOK, {}, exported)
        collect.verify_workbook(output / MORPHOLOGY_WORKBOOK, {},
                                {title: exported[title] for title in morphology})
        for path, content in data['files'].items():
            if path.is_symlink() or path.read_bytes() != content:
                raise inputs.ValidationError(f'Input changed during report generation: {path}')
        logger.info('Verified %s non-border nuclei, %s images, %s markers, %s planned comparisons',
                    len(data['nuclei']), len(data['images']), len(selected), len(data['statistics']))
        write_status(output, 'SUCCESS', Collection=str(collection), Nuclei_run_ID=data['run_id'],
                     Statistics_unit=stats_unit or 'disabled', Markers=selected,
                     Non_border_nuclei=len(data['nuclei']), Images=len(data['images']),
                     Planned_comparisons=len(data['statistics']),
                     Tested_comparisons=sum(r['Status'] == 'TESTED' for r in data['statistics']))
        print(f'SUCCESS: {output}')
        return True, output
    except (collect.Cancelled, KeyboardInterrupt, EOFError):
        write_status(output, 'CANCELED', Stage=stage, Collection=str(collection))
        raise
    except Exception as error:
        logger.exception('Report failed during %s', stage)
        write_status(output, 'FAILED', Stage=stage, Collection=str(collection), Error=str(error))
        collect.write_workbook(output / 'Report_Diagnostics.xlsx', {
            'Diagnostics': (['Status', 'Stage', 'Error'], [{'Status': 'FAILED', 'Stage': stage, 'Error': str(error)}])})
        print(f'FAILED during {stage}: {error}. Diagnostics: {output}')
        return False, output
    finally:
        logger.removeHandler(handler)
        handler.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description='Report collected nuclear morphology and marker intensity; no image processing.')
    parser.add_argument('-i', '--input', required=True, help='JSON with paths_to_files experiment folders')
    parser.add_argument('--stats-unit', choices=('nucleus', 'well'), help='Omit for descriptive results without tests')
    parser.add_argument('--template', help='96-well plate-map XLSX; otherwise discover it inside each collection folder')
    parser.add_argument('--sheet', help='Plate-map worksheet name; default: first sheet')
    parser.add_argument('--all-experiments', action='store_true', help='Use all valid manifest folders without the experiment prompt')
    parser.add_argument('--collections', choices=('ask', 'latest', 'all'), default='ask', help='Default: interactive selection')
    parser.add_argument('--markers', help='Comma-separated marker folders, all, or none; default: interactive selection per collection')
    args = parser.parse_args(argv)
    try:
        roots = collect.discover_sources(args.input)
        if not roots:
            raise inputs.ValidationError('No existing experiment folders')
        if not args.all_experiments:
            for index, root in enumerate(roots, 1):
                print(f'{index}. {root}')
            roots = [roots[index] for index in collect.choose_indices('Select experiments, all, or q: ', len(roots))]
        selected, missing = choose_collections(roots, args.collections)
        print(f'Statistics: {args.stats_unit or "disabled (descriptive only)"}. Reports remain separate per collection.')
        failures = len(missing)
        for collection in selected:
            try:
                success, _ = create_report(collection, collection.parent, args.template, args.markers,
                                           args.stats_unit, args.sheet, args.input)
                failures += not success
            except OSError as error:
                print(f'Cannot write report for {collection}: {error}')
                failures += 1
        return 1 if failures else 0
    except (collect.Cancelled, KeyboardInterrupt, EOFError):
        print('Report canceled. Completed reports have report_status.json with Status SUCCESS.')
        return 130
    except (OSError, ValueError) as error:
        print(f'Marker-intensity report could not complete: {error}')
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
