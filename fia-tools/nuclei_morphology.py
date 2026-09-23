"""Export pixel-based nuclear morphology without modifying final masks."""

import csv
import math
import re
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from scyjava import jimport


METRICS = [
    "Area_px2", "Perimeter_px", "Circularity", "Aspect_ratio", "Solidity",
    "Major_axis_px", "Minor_axis_px", "Feret_max_px", "Feret_min_px",
    "Equivalent_diameter_px", "Roundness", "Eccentricity",
]
IDENTIFIERS = [
    "Dataset", "Dataset_path", "Image_name", "Mask_name", "Well",
    "Run_ID", "Particle_size_px2", "StarDist_source",
]
NUCLEI_COLUMNS = IDENTIFIERS + ["Nucleus_ID"] + METRICS + [
    "Orientation_deg", "Centroid_X_px", "Centroid_Y_px", "Touches_border",
    "Label_map", "Numbered_image",
]
IMAGE_COLUMNS = IDENTIFIERS + [
    "Status", "Error", "Nuclei_count_total", "Border_nuclei_count",
    "Non_border_nuclei_count",
] + [f"{metric}_{stat}" for metric in METRICS for stat in ("Mean", "Median", "IQR")]


def processor_pixels(imp):
    """Copy raw pixels, independent of the ImageJ display LUT."""
    dtype = np.uint8 if int(imp.getBitDepth()) == 8 else np.uint16
    return np.asarray(imp.getProcessor().getPixels()).astype(dtype).reshape(
        int(imp.getHeight()), int(imp.getWidth()))


def measure_final_mask(mask):
    """Measure a duplicate with ImageJ and return rows plus an ID label image.

    The caller supplies the unchanged `show=Masks` output of ParticleAnalyzer.
    These masks display particles in black, including when the LUT is inverted.
    No watershed, size filter or exclusion of border objects is applied here.
    The caller owns the returned label image and must close it.
    """
    ImagePlus = jimport("ij.ImagePlus")
    Calibration = jimport("ij.measure.Calibration")
    Measurements = jimport("ij.measure.Measurements")
    ResultsTable = jimport("ij.measure.ResultsTable")
    ParticleAnalyzer = jimport("ij.plugin.filter.ParticleAnalyzer")
    ImageProcessor = jimport("ij.process.ImageProcessor")
    Color = jimport("java.awt.Color")

    if int(mask.getBitDepth()) != 8 or int(mask.getStackSize()) != 1:
        raise ValueError("Morphology requires a single 8-bit final mask.")
    original = processor_pixels(mask)
    if not np.isin(original, (0, 255)).all():
        raise ValueError("The final nuclei mask is not binary.")
    foreground = int(mask.getProcessor().getBestIndex(Color.black))
    if foreground not in (0, 255):
        raise ValueError("Unrecognized final-mask LUT.")

    duplicate = ImagePlus("Nuclei morphology", mask.getProcessor().duplicate())
    duplicate.setIgnoreGlobalCalibration(True)
    duplicate.setCalibration(Calibration())
    duplicate.getProcessor().setThreshold(
        foreground, foreground, ImageProcessor.NO_LUT_UPDATE)
    table = ResultsTable()
    flags = (Measurements.AREA | Measurements.PERIMETER | Measurements.ELLIPSE
             | Measurements.SHAPE_DESCRIPTORS | Measurements.FERET
             | Measurements.CENTROID)
    analyzer = ParticleAnalyzer(
        ParticleAnalyzer.SHOW_ROI_MASKS, flags, table, 0.0, float("inf"))
    analyzer.setHideOutputImage(True)
    labels_image = None
    try:
        if not analyzer.analyze(duplicate):
            raise RuntimeError("ImageJ could not measure the final nuclei mask.")
        count = int(table.size())
        if count > 65534:
            raise ValueError("Too many nuclei for unique ImageJ count-mask IDs.")
        labels_image = analyzer.getOutputImage()
        if labels_image is None:
            raise RuntimeError("ImageJ did not return a nucleus ID map.")
        labels = processor_pixels(labels_image)
        if not np.array_equal(labels > 0, original == foreground):
            raise ValueError("Measured objects do not match the saved final mask.")
        sizes = np.bincount(labels.ravel(), minlength=count + 1)
        if len(sizes) != count + 1 or np.any(sizes[1:] == 0):
            raise ValueError("Nucleus IDs do not match the measurement rows.")
        border_ids = set(np.concatenate(
            (labels[0], labels[-1], labels[:, 0], labels[:, -1])))

        def value(heading, row):
            number = float(table.getValue(heading, row))
            return number if math.isfinite(number) else None

        records = []
        for index in range(count):
            area = value("Area", index)
            if area is None or not math.isclose(area, int(sizes[index + 1])):
                raise ValueError("ImageJ area does not match the ID-map pixel count.")
            major, minor = value("Major", index), value("Minor", index)
            eccentricity = None
            if major is not None and major > 0 and minor is not None:
                eccentricity = math.sqrt(max(0.0, 1.0 - (minor / major) ** 2))
            records.append({
                "Nucleus_ID": index + 1,
                "Area_px2": int(sizes[index + 1]),
                "Perimeter_px": value("Perim.", index),
                "Circularity": value("Circ.", index),
                "Aspect_ratio": value("AR", index),
                "Solidity": value("Solidity", index),
                "Major_axis_px": major,
                "Minor_axis_px": minor,
                "Feret_max_px": value("Feret", index),
                "Feret_min_px": value("MinFeret", index),
                "Equivalent_diameter_px": math.sqrt(4.0 * area / math.pi),
                "Roundness": value("Round", index),
                "Eccentricity": eccentricity,
                "Orientation_deg": value("Angle", index),
                "Centroid_X_px": value("X", index),
                "Centroid_Y_px": value("Y", index),
                "Touches_border": index + 1 in border_ids,
            })
        return records, labels_image
    except Exception:
        if labels_image is not None:
            labels_image.close()
        raise
    finally:
        duplicate.close()


def save_numbered_image(mask, records, path):
    """Draw IDs on a separate RGB image, leaving the binary mask unchanged."""
    ImagePlus = jimport("ij.ImagePlus")
    FileSaver = jimport("ij.io.FileSaver")
    Color = jimport("java.awt.Color")
    JavaFont = jimport("java.awt.Font")
    processor = mask.getProcessor().convertToRGB()
    processor.setFont(JavaFont("SansSerif", JavaFont.BOLD, 12))
    processor.setColor(Color.red)
    for row in records:
        text = str(row["Nucleus_ID"])
        width = int(processor.getStringWidth(text))
        x = round(row["Centroid_X_px"]) - width // 2
        y = round(row["Centroid_Y_px"])
        x = max(0, min(x, int(mask.getWidth()) - width))
        y = min(int(mask.getHeight()) - 1, max(12, y))
        processor.drawString(text, x, y)
    preview = ImagePlus("Nucleus IDs", processor)
    try:
        if not FileSaver(preview).saveAsPng(str(path)):
            raise OSError(f"Could not save numbered image: {path}")
    finally:
        preview.close()


class NucleiMorphologyExport:
    """Accumulate per-nucleus measurements and per-image quality information."""

    def __init__(self, output_folder, stardist_folder, particle_size, ij_version):
        self.output = Path(output_folder)
        self.source = Path(stardist_folder).resolve()
        self.particle_size = particle_size
        self.ij_version = str(ij_version)
        self.nuclei = []
        self.images = []

    def identifiers(self, filename, mask_name=""):
        source = Path(filename)
        stem = source.stem.removesuffix("_StarDist_processed")
        image_name = stem + source.suffix
        well = re.search(r"(?:^|[_\s-])Well([A-P])(0?[1-9]|1[0-9]|2[0-4])"
                         r"(?=[_\s.\-]|$)", image_name, re.IGNORECASE)
        dataset = self.output.parent.parent.resolve()
        return {
            "Dataset": dataset.name, "Dataset_path": str(dataset),
            "Image_name": image_name, "Mask_name": mask_name,
            "Well": f"{well[1].upper()}{int(well[2])}" if well else "",
            "Run_ID": self.output.name, "Particle_size_px2": self.particle_size,
            "StarDist_source": str(self.source),
        }

    def add_image(self, mask, filename, mask_name):
        records, label_image = measure_final_mask(mask)
        qc = self.output / "Morphology_QC"
        stem = Path(mask_name).stem
        label_path = qc / f"{stem}_ids.tif"
        preview_path = qc / f"{stem}_ids.png"
        try:
            qc.mkdir(exist_ok=True)
            FileSaver = jimport("ij.io.FileSaver")
            if not FileSaver(label_image).saveAsTiff(str(label_path)):
                raise OSError(f"Could not save nucleus ID map: {label_path}")
            save_numbered_image(mask, records, preview_path)
        finally:
            label_image.close()
        # Exclude edge-touching nuclei from every measurement table and summary.
        # Keep their IDs in QC images so the unchanged final mask remains traceable.
        included_records = [row for row in records if not row["Touches_border"]]
        identifiers = self.identifiers(filename, mask_name)
        self.nuclei.extend({
            **identifiers, **row,
            "Label_map": str(label_path.relative_to(self.output)),
            "Numbered_image": str(preview_path.relative_to(self.output)),
        } for row in included_records)
        summary = {
            **identifiers, "Status": "complete", "Error": "",
            "Nuclei_count_total": len(records),
            "Border_nuclei_count": len(records) - len(included_records),
            "Non_border_nuclei_count": len(included_records),
        }
        # Orientation is axial and is deliberately not summarized by linear statistics.
        for metric in METRICS:
            values = [row[metric] for row in included_records if row[metric] is not None]
            summary[f"{metric}_Mean"] = float(np.mean(values)) if values else None
            summary[f"{metric}_Median"] = float(np.median(values)) if values else None
            summary[f"{metric}_IQR"] = (
                float(np.percentile(values, 75) - np.percentile(values, 25))
                if values else None)
        self.images.append(summary)

    def record_failure(self, filename, error, mask_name=""):
        self.images.append({**self.identifiers(filename, mask_name),
                            "Status": "failed", "Error": str(error),
                            "Nuclei_count_total": None, "Border_nuclei_count": None,
                            "Non_border_nuclei_count": None})

    def run_info(self):
        return [
            ("Run_ID", self.output.name),
            ("StarDist_source", str(self.source)),
            ("Particle_size_px2", self.particle_size),
            ("Status", "complete" if self.images and all(
                row["Status"] == "complete" for row in self.images) else "incomplete"),
            ("ImageJ_version", self.ij_version),
            ("Measurement_method", "ImageJ ParticleAnalyzer on a duplicate final mask"),
            ("Area_units", "pixels squared; no physical calibration"),
            ("Length_units", "pixels of the processed image"),
            ("Shape_space", "2D processed-image pixel coordinates"),
            ("Circularity", "ImageJ: min(1, 4*pi*Area/Perimeter^2)"),
            ("Aspect_ratio", "Major_axis_px / Minor_axis_px"),
            ("Roundness", "4*Area_px2/(pi*Major_axis_px^2)"),
            ("Solidity", "Area / convex hull area (ImageJ definition)"),
            ("Feret_diameters", "Maximum and minimum caliper diameters (ImageJ)"),
            ("Equivalent_diameter_px", "sqrt(4*Area_px2/pi)"),
            ("Eccentricity", "sqrt(1-(Minor_axis_px/Major_axis_px)^2)"),
            ("Orientation_deg", "ImageJ fitted-ellipse angle to the x axis (0-180)"),
            ("Centroid", "ImageJ pixel coordinates; origin at the top left"),
            ("Border_policy", "Nuclei touching any image edge are excluded from "
             "per-nucleus tables and summary metrics; counted separately"),
            ("Summary_population", "Final-mask nuclei that do not touch the image border"),
            ("Non_border_nuclei_count", "Nuclei_count_total - Border_nuclei_count"),
            ("Summary_statistics", "Per-image arithmetic mean, median and IQR; "
             "blank if no non-border nuclei; no hypothesis tests"),
            ("Nucleus_ID", "Local to Image_name and Run_ID; matches Morphology_QC ID map"),
            ("QC_population", "All final-mask nuclei, including border nuclei; "
             "exported Nucleus_ID values may have gaps"),
            ("ID_compatibility", "Not guaranteed to match StarDist or quantify_foci IDs"),
            ("Intensity_and_texture", "Not measured"),
            ("Missing_values", "Blank means unavailable; zero nuclei is recorded explicitly"),
        ]

    def save(self):
        """Write Excel and UTF-8 CSV tables without rounding numeric values."""
        tables = [("Nuclei", NUCLEI_COLUMNS, self.nuclei),
                  ("Images", IMAGE_COLUMNS, self.images),
                  ("Run_Info", ["Parameter", "Value"],
                   [dict(zip(("Parameter", "Value"), pair)) for pair in self.run_info()])]
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, columns, records in tables:
            csv_name = "Nuclei_Morphology.csv" if name == "Nuclei" else f"Nuclei_{name}.csv"
            csv_path = self.output / csv_name
            temporary = csv_path.with_name("." + csv_name + ".tmp")
            with temporary.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(handle, fieldnames=columns)
                writer.writeheader()
                writer.writerows(records)
            temporary.replace(csv_path)

            sheet = workbook.create_sheet(name)
            sheet.append(columns)
            for record in records:
                sheet.append([record.get(column) for column in columns])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="264653")
                sheet.column_dimensions[cell.column_letter].width = min(
                    32, max(16, len(str(cell.value)) + 2))
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    elif isinstance(cell.value, float):
                        cell.number_format = "0.0000"
        destination = self.output / "Nuclei_Morphology.xlsx"
        temporary = destination.with_name("." + destination.name)
        try:
            workbook.save(temporary)
            temporary.replace(destination)
        finally:
            workbook.close()
        return dict(self.run_info())["Status"]
