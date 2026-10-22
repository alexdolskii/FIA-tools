"""Authoritative per-well routing and plot layout for the FIA 96-well workbook.

This module has no imaging dependencies. A channel number always refers to the
original image; neither Ignore nor a missing channel changes that numbering.
"""

import hashlib
import io
import json
import re
import unicodedata
from zipfile import BadZipFile
from pathlib import Path

from openpyxl import load_workbook

from assay_layout import ASSAY_DIR

METHODS = ('Nuclei', 'Foci', 'Nuclear intensity', 'Field signal', 'Fiber')
PREFIXES = dict(zip(METHODS, ('Nuclei', 'Foci', 'NuclearIntensity', 'FieldSignal', 'Fiber')))
MARKER_PATTERN = re.compile(r'^(?:Foci|NuclearIntensity|FieldSignal|Fiber)_[A-Za-z0-9_]+_[0-9a-f]{12}$')
WELL_PATTERN = re.compile(r'Well([A-H])(0?[1-9]|1[0-2])(?!\d)', re.IGNORECASE)
SNAPSHOT = 'plate_design.json'
PREPARATION = 'preparation_run.json'
_OVERRIDES = {}
_ANNOUNCED = set()

JOURNALS = dict(zip((
    'select_channels', 'generate_nuclei_mask', 'generate_foci_mask', 'quantify_foci',
    'quantify_nuclear_intensity', 'fia_collect_marker_intensity_results',
    'fia_marker_intensity_report', 'fia_nuclei_report', 'quantify_field_signal', 'fia_field_signal_report'), (
    '1_log.log', '2_log.log', '3_foci_mask.log', '4_foci_quantification.log',
    '3_nuclei_intensity.log', '4_collect_marker_intensity.log', '5_marker_intensity_report.log',
    '3_nuclei_report.log', '3_field_signal.log', '4_field_signal_report.log')))


def record_validation_failure(command, input_path, roots, errors):
    """Save actionable input errors even before analysis output is created."""
    import logging
    from table_workflow import TableJournal
    audit = TableJournal(JOURNALS[command], input_path)
    try:
        for root in roots:
            if not root.is_dir():
                continue
            audit.register(root)
            message = errors.get(root, 'Batch stopped: another experiment has an invalid plate scheme.')
            audit.event(root, 'INPUT_VALIDATION_FAILED | %s | Action: place one completed FIA*.xlsx beside raw images and rerun.',
                        message, level=logging.ERROR)
        audit.status = 'FAILED'
    finally:
        audit.finish()


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     separators=(',', ':')).encode('utf-8')).hexdigest()


def file_digest(path):
    value = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            value.update(chunk)
    return value.hexdigest()


def marker_id(method, name):
    """Portable, collision-resistant identity, independent of source channel."""
    label = unicodedata.normalize('NFC', name.strip())
    slug = re.sub(r'[^A-Za-z0-9]+', '_', label).strip('_')[:40] or 'marker'
    return f'{PREFIXES[method]}_{slug}_{digest([method, label])[:12]}'


def well_from_name(value):
    name = Path(value).name
    matches = WELL_PATTERN.findall(name)
    if len(matches) != 1:
        raise ValueError(f'Expected one unambiguous WellA01–WellH12 identifier: {name}')
    well = f'{matches[0][0].upper()}{int(matches[0][1]):02d}'
    points = re.findall(r'Point([A-H])(0?[1-9]|1[0-2])(?!\d)', name, re.IGNORECASE)
    if any(f'{row.upper()}{int(column):02d}' != well for row, column in points):
        raise ValueError(f'Well and Point identifiers disagree: {name}')
    return well


def _literal(cell):
    if cell.data_type in ('f', 'e'):
        raise ValueError(f'{cell.parent.title}!{cell.coordinate}: use literal values, not formulas/errors')
    if cell.value is None:
        return ''
    if not isinstance(cell.value, str):
        raise ValueError(f'{cell.parent.title}!{cell.coordinate}: expected text')
    return unicodedata.normalize('NFC', cell.value.strip())


def _color(cell):
    color = cell.fill.fgColor
    if cell.fill.patternType != 'solid' or color.type != 'rgb':
        raise ValueError(f'{cell.parent.title}!{cell.coordinate}: use a direct solid RGB block fill')
    return str(color.rgb)[-6:].upper()


class PlateDesign:
    def __init__(self, path, content, wells, panels):
        self.path = Path(path)
        self.content = content
        self.wells = wells
        self.panels = panels
        self.workbook_sha256 = hashlib.sha256(content).hexdigest()

    def channel(self, well, method, marker=None):
        values = [row['channel'] for row in self.wells.get(well, {}).get('channels', [])
                  if row and row['method'] == method and (marker is None or row['id'] == marker)]
        if len(values) != 1:
            raise ValueError(f'{well}: expected one {method} channel for {marker or method}; found {len(values)}')
        return values[0]

    def assigned(self, well, method=None, marker=None):
        return well in self.wells and (method is None or any(
            row and row['method'] == method and (marker is None or row['id'] == marker)
            for row in self.wells[well]['channels']))

    def validate_channels(self, well, count):
        expected = sum(row is not None for row in self.wells[well]['channels'])
        if int(count) != expected:
            raise ValueError(f'{well}: Excel declares {expected} existing channels, original has {count}. '
                             'Use Ignore for an existing unused channel; blank only for absent channels.')

    def markers(self, method):
        result = {}
        for well, item in self.wells.items():
            for row in item['channels']:
                if row and row['method'] == method:
                    marker = result.setdefault(row['id'], {k: row[k] for k in ('id', 'name', 'method')})
                    marker.setdefault('channels', {})[well] = row['channel']
        return list(result.values())

    def scope_hash(self, method=None, marker=None):
        rows = {}
        for well, item in self.wells.items():
            if not self.assigned(well, method, marker):
                continue
            if method == 'Nuclei':
                assignments = [self.channel(well, 'Nuclei')]
            elif method:
                assignments = [row for row in item['channels'] if row and (
                    row['method'] == 'Nuclei' or (row['method'] == method and
                                                (marker is None or row['id'] == marker)))]
            else:
                assignments = item['channels']
            rows[well] = {'channels': assignments,
                          'count': sum(row is not None for row in item['channels'])}
        return digest(rows)

    def report_rows(self):
        order = {group: index for index, group in enumerate(
            (g for panel in self.panels for g in panel['groups']), 1)}
        positions = {group: panel for panel in self.panels for group in panel['groups']}
        return [{'Well': well, 'Group': item['condition'], 'Color': item['block'],
                 'Is_control': item['control'], 'Excel_cell': item['cell'],
                 'Group_order': order[item['condition']], 'Order_source': 'plot_layout',
                 'Order_cell': positions[item['condition']]['cells'][item['condition']],
                 'Order_group_cell': positions[item['condition']]['cells'][item['condition']]}
                for well, item in self.wells.items()]


def read_design(path, allow_empty=False):
    path = Path(path)
    content = path.read_bytes()
    book = load_workbook(io.BytesIO(content), data_only=False)
    try:
        if not {'Plate Map', 'Channels'} <= set(book.sheetnames):
            raise ValueError(f'{path.name}: required sheets are Plate Map and Channels; use the current FIA template')
        plate, channels = book['Plate Map'], book['Channels']
        if ([plate.cell(1, c).value for c in range(2, 14)] != list(range(1, 13)) or
                [plate.cell(r, 1).value for r in range(2, 10)] != list('ABCDEFGH')):
            raise ValueError('Plate Map: expected the 96-well grid at A1:M9')
        if ([channels.cell(5, c).value for c in range(3, 15)] != list(range(1, 13)) or
                [channels.cell(43, c).value for c in range(3, 15)] != list(range(1, 13))):
            raise ValueError('Channels: expected channel grid at C6:N37 and intersections at C44:N51')
        for sheet, regions in ((plate, [(2, 9, 2, 13), (15, 21, 2, 13)]),
                               (channels, [(6, 37, 3, 14), (44, 51, 3, 14)])):
            for area in sheet.merged_cells.ranges:
                if any(area.min_row <= r2 and area.max_row >= r1 and area.min_col <= c2 and area.max_col >= c1
                       for r1, r2, c1, c2 in regions):
                    raise ValueError(f'{sheet.title}: merged input cells are not supported ({area})')
            for rule in sheet.conditional_formatting:
                if any(a.min_row <= r2 and a.max_row >= r1 and a.min_col <= c2 and a.max_col >= c1
                       for a in rule.sqref.ranges for r1, r2, c1, c2 in regions):
                    raise ValueError(f'{sheet.title}: use direct formatting in input cells')
        wells, conditions = {}, {}
        for row_index, letter in enumerate('ABCDEFGH'):
            if channels.cell(6 + 4 * row_index, 1).value != letter:
                raise ValueError(f'Channels: missing row label {letter}')
            if [channels.cell(6 + 4 * row_index + k, 2).value for k in range(4)] != ['C1', 'C2', 'C3', 'C4']:
                raise ValueError(f'Channels: {letter} must have C1, C2, C3, C4 in that order')
            for column in range(1, 13):
                well = f'{letter}{column:02d}'
                cell = plate.cell(row_index + 2, column + 1)
                condition = _literal(cell)
                cells = [channels.cell(6 + 4 * row_index + k, column + 2) for k in range(4)]
                entries = [_literal(c) for c in cells]
                intersection_cell = channels.cell(44 + row_index, column + 2)
                requested = _literal(intersection_cell)
                if not condition:
                    if any(entries) or requested:
                        raise ValueError(f'{well}: channels/intersections specified for an unannotated well')
                    continue
                block, control = _color(cell), bool(cell.font.bold)
                if condition in conditions and conditions[condition] != (block, control):
                    raise ValueError(f'{cell.coordinate}: condition {condition!r} has inconsistent block/control formatting')
                conditions[condition] = block, control
                assignments, seen, absent = [], set(), False
                for index, (entry, source_cell) in enumerate(zip(entries, cells), 1):
                    if not entry:
                        absent = True
                        assignments.append(None)
                        continue
                    if absent:
                        raise ValueError(f'{well}: missing channel before C{index}; use Ignore for existing unused channels')
                    if entry == 'Ignore':
                        assignments.append({'channel': index, 'method': 'Ignore', 'name': '', 'id': ''})
                        continue
                    method, sep, label = entry.partition(':')
                    method, label = method.strip(), label.strip()
                    if not sep or method not in METHODS or not label:
                        raise ValueError(f'Channels!{source_cell.coordinate}: expected Method: marker name, or Ignore')
                    identity = marker_id(method, label)
                    if identity in seen:
                        raise ValueError(f'{well}: duplicate {method}: {label}; use distinct marker names')
                    seen.add(identity)
                    assignments.append({'channel': index, 'method': method, 'name': label, 'id': identity})
                if sum(bool(a and a['method'] == 'Nuclei') for a in assignments) != 1:
                    raise ValueError(f'{well}: exactly one Nuclei channel is required')
                intersections = []
                for value in requested.split(';') if requested else []:
                    value = value.strip()
                    if not re.fullmatch(r'C[1-4](\s*\+\s*C[1-4]){1,2}', value):
                        raise ValueError(f'{well}: invalid intersection {value!r}; use C2+C3; C2+C3+C4')
                    numbers = sorted(int(v.strip()[1:]) for v in value.split('+'))
                    if len(set(numbers)) != len(numbers):
                        raise ValueError(f'{well}: repeated channel in {value}')
                    if any(not assignments[n-1] or assignments[n-1]['method'] != 'Foci' for n in numbers):
                        raise ValueError(f'{well}: all channels in {value} must be assigned Foci')
                    if numbers not in intersections:
                        intersections.append(numbers)
                wells[well] = {'condition': condition, 'block': block, 'control': control,
                               'cell': cell.coordinate, 'channels': assignments, 'intersections': intersections}
        if not wells and not allow_empty:
            raise ValueError('Plate Map has no annotated wells; fill the template before analysis')
        panels, located, blocks = [], set(), set()
        for row in range(15, 22):
            column = 2
            while column <= 13:
                if not _literal(plate.cell(row, column)):
                    column += 1
                    continue
                start, groups, cells = column, [], {}
                block = _color(plate.cell(row, column))
                if block in blocks:
                    raise ValueError('PLOT LAYOUT: each comparison block must occupy one horizontal strip')
                while column <= 13 and _literal(plate.cell(row, column)):
                    cell = plate.cell(row, column)
                    group = _literal(cell)
                    if group not in conditions or conditions[group] != (_color(cell), bool(cell.font.bold)):
                        raise ValueError(f'PLOT LAYOUT!{cell.coordinate}: label/fill/bold must match Plate Map')
                    if _color(cell) != block:
                        raise ValueError(f'PLOT LAYOUT!{cell.coordinate}: leave a blank column between blocks')
                    if group in located:
                        raise ValueError(f'PLOT LAYOUT: repeated condition {group!r}')
                    located.add(group)
                    groups.append(group)
                    cells[group] = cell.coordinate
                    column += 1
                expected = {g for g, (color, _) in conditions.items() if color == block}
                if set(groups) != expected:
                    raise ValueError(f'PLOT LAYOUT: strip at row {row} must contain every condition of its color block')
                blocks.add(block)
                panels.append({'row': row, 'column': start, 'end_column': column - 1,
                               'block': block, 'groups': groups, 'cells': cells})
        if located != set(conditions):
            raise ValueError('PLOT LAYOUT: missing conditions: ' + ', '.join(sorted(set(conditions) - located)))
        return PlateDesign(path, content, wells, panels)
    finally:
        book.close()


def find_design(root, template=None):
    root = Path(root).resolve()
    override = template or _OVERRIDES.get(root)
    if override:
        return Path(override).expanduser().resolve()
    paths = sorted(p for p in root.iterdir() if p.is_file() and not p.is_symlink()
                   and p.name.casefold().startswith('fia') and p.suffix.casefold() == '.xlsx')
    if len(paths) != 1:
        raise ValueError(f'{root}: expected exactly one FIA*.xlsx beside raw images; found {len(paths)}. '
                         + ', '.join(p.name for p in paths))
    return paths[0]


def load_design(root, template=None):
    path = find_design(root, template)
    try:
        return read_design(path)
    except BadZipFile as error:
        raise ValueError(f'{path}: unreadable Excel workbook; save a valid .xlsx file') from error


def preflight(input_path, template=None, command=None):
    """Validate every experiment before starting an imaging engine or analysis."""
    from validate_folders import validate_input_file
    results, errors = {}, {}
    try:
        values = validate_input_file(input_path)
    except (OSError, ValueError) as error:
        if command:
            print(f'INPUT_VALIDATION_FAILED | {error}')
        raise
    roots = [Path(value).resolve() for value in values]
    for root in roots:
        try:
            design = load_design(root, template)
            results[root] = design
            if template:
                _OVERRIDES[root] = design.path
            notice = (str(root), design.workbook_sha256)
            if notice not in _ANNOUNCED:
                print(f'Plate scheme: {root.name} | {design.path.name} | annotated wells: {len(design.wells)}')
                _ANNOUNCED.add(notice)
        except (OSError, ValueError, KeyError) as error:
            errors[root] = str(error)
    if errors:
        if command:
            record_validation_failure(command, input_path, roots, errors)
        raise ValueError('Plate scheme validation failed before analysis:\n' + '\n'.join(errors.values()))
    return results


def sources(root, design, method=None, marker=None):
    selected, skipped, stems = [], [], set()
    for path in sorted(Path(root).iterdir()):
        if not path.is_file() or path.name.startswith('.') or path.suffix.lower() not in ('.nd2', '.tif', '.tiff'):
            continue
        well = well_from_name(path.name)
        if not design.assigned(well, method, marker):
            skipped.append(path)
            continue
        if path.stem in stems:
            raise ValueError(f'Ambiguous original image stem: {path.stem}')
        stems.add(path.stem)
        selected.append(path)
    return selected, skipped


def selected_markers(designs, method, requested=None):
    markers = {m['id']: m for design in designs for m in design.markers(method)}
    if requested:
        values = [requested] if isinstance(requested, str) else requested
        chosen = set()
        for value in values:
            matches = [m['id'] for m in markers.values() if value in (m['id'], m['name'])]
            if len(matches) != 1:
                raise ValueError(f'{method}: marker {value!r} is absent or ambiguous in the selected schemes')
            chosen.add(matches[0])
        markers = {key: value for key, value in markers.items() if key in chosen}
    return list(markers.values())


def save_snapshot(folder, design):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    value = {'Schema_version': 1, 'Workbook': str(design.path), 'Workbook_SHA256': design.workbook_sha256,
             'Wells': design.wells, 'Panels': design.panels, 'Routing_SHA256': design.scope_hash(),
             'Nuclei_SHA256': design.scope_hash('Nuclei')}
    temporary = folder / ('.' + SNAPSHOT + '.tmp')
    temporary.write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    temporary.replace(folder / SNAPSHOT)
    (folder / 'plate_design.xlsx').write_bytes(design.content)


def snapshot_design(folder):
    path = Path(folder) / 'plate_design.xlsx'
    meta = json.loads((Path(folder) / SNAPSHOT).read_text(encoding='utf-8'))
    design = read_design(path)
    if meta.get('Schema_version') != 1 or design.workbook_sha256 != meta.get('Workbook_SHA256'):
        raise ValueError(f'Plate scheme snapshot changed or is invalid: {folder}')
    return design


def validate_snapshot(folder, design, method='Nuclei', marker=None):
    try:
        original = snapshot_design(folder)
    except (OSError, ValueError, KeyError, BadZipFile) as error:
        raise ValueError(f'{folder}: missing/invalid scheme provenance; rerun this analysis with the current template') from error
    if original.scope_hash(method, marker) != design.scope_hash(method, marker):
        raise ValueError(f'{folder}: {method} assignments changed; regenerate the affected analysis')
    return original


def preparation_files(folder):
    """Only successfully prepared files from a complete current preparation."""
    folder = Path(folder)
    root = folder.parent.parent if folder.name == 'Nuclei' else folder.parents[2]
    design = load_design(root)
    method = 'Nuclei' if folder.name == 'Nuclei' else next(
        (method for method in METHODS if any(m['id'] == folder.name for m in design.markers(method))), None)
    if method is None:
        raise ValueError(f'{folder}: marker is not declared by the current scheme')
    validate_snapshot(folder, design, method, None if method == 'Nuclei' else folder.name)
    meta = json.loads((folder / PREPARATION).read_text(encoding='utf-8'))
    if meta.get('Status') != 'COMPLETE':
        raise ValueError(f'{folder}: incomplete channel preparation; rerun select_channels')
    for name, checksum in meta.get('Sources', {}).items():
        well = well_from_name(name)
        if not design.assigned(well, method, None if method == 'Nuclei' else folder.name):
            continue
        source = root / name
        if Path(name).name != name or not source.is_file() or file_digest(source) != checksum:
            raise ValueError(f'{source}: original changed after preparation; rerun select_channels')
    files = meta.get('Files', {})
    for name, checksum in files.items():
        path = folder / name
        if path.name != name or not path.is_file() or file_digest(path) != checksum:
            raise ValueError(f'{path}: prepared file changed or is missing; rerun select_channels')
    return sorted(files)
