"""Validate collected FIA spreadsheets and annotate nuclei from a plate map."""

import hashlib
import io
import math
import re
from collections import defaultdict
from pathlib import Path
from zipfile import BadZipFile

import collect_marker_intensity_results as collect
import numpy as np
from openpyxl import load_workbook

ValidationError = collect.ValidationError
COLLECTION_PATTERN = re.compile(r'^FIA_Marker_Intensity_Combined_Results_(\d{8}_\d{6})$')
MARKER_PATTERN = re.compile(r'^(?:Foci_[1-9][0-9]*_)?Channel_([1-9][0-9]*)$')
WELL_PATTERN = re.compile(r'Well([A-H])(0?[1-9]|1[0-2])(?!\d)', re.IGNORECASE)
COUNT = 'Non_border_nuclei_count'


def workbook_rows(path, sheet, files):
    try:
        book = load_workbook(io.BytesIO(collect.snapshot(path, files)), read_only=True, data_only=False)
        try:
            values = iter(book[sheet].values)
            columns = list(next(values))
            if not columns or len(set(columns)) != len(columns) or any(not x for x in columns):
                raise ValidationError(f'Invalid {sheet} headers')
            return [dict(zip(columns, tuple(row) + (None,) * (len(columns) - len(row)))) for row in values]
        finally:
            book.close()
    except Exception as error:
        raise ValidationError(f'Cannot read {path.name}/{sheet}: {error}') from error


def collection_info(path, files):
    rows = workbook_rows(path / collect.COMBINED_NAME, 'Collection_Info', files)
    status = [row['Status'] for row in rows if row.get('Category') == 'Collection' and row.get('Item') == 'Status']
    if len(status) != 1 or status[0] not in ('SUCCESS', 'SUCCESS_WITH_MISSING_INTENSITY'):
        raise ValidationError('Collection must have one successful completion status')
    return rows


def _unique(rows, key, label):
    result = {}
    for row in rows:
        identity = key(row)
        if identity in result:
            raise ValidationError(f'Duplicate {label}: {identity}')
        result[identity] = row
    return result


def _nucleus_key(row):
    return row['Mask_name'], collect.number(row['Nucleus_ID'], 'Nucleus_ID', integer=True)


def _check_values(expected, actual, mapping=None):
    mapping = mapping or {}
    for field, value in expected.items():
        target = mapping.get(field, field)
        if target not in actual or not collect.matching_cell(value, actual[target]):
            raise ValidationError(f'Combined/source table mismatch in {target}')


def _read_intensity(path, marker, files, copied):
    names = [f'Nuclear_Intensity_{marker}.csv', f'Nuclear_Intensity_Images_{marker}.csv',
             f'Nuclear_Intensity_Run_Info_{marker}.csv', f'Nuclear_Intensity_{marker}.xlsx']
    if any(name not in copied for name in names):
        raise ValidationError(f'Missing collector source record for {marker}')
    identity = ['Dataset_path', 'Nuclei_run_ID', 'Intensity_run_ID', 'Mask_file', 'Image_name',
                'Source_file', 'Marker_channel', 'Particle_size_px2']
    required = [identity + ['Nucleus_ID', 'Area_px2'] + list(collect.INTENSITY_METRICS),
                identity + ['Status', 'Error', 'Nuclei_measured_count', 'Failed_nuclei_count'] + list(collect.COUNTS)
                + [f'{metric}_{stat}' for metric in ('Area_px2', *collect.INTENSITY_METRICS) for stat in collect.STATS],
                ['Parameter', 'Value']]
    tables = {sheet: collect.csv_snapshot(path / name, files, fields)
              for sheet, name, fields in zip(('Nuclei', 'Images', 'Run_Info'), names, required)}
    collect.verify_workbook(path / names[3], files, tables)
    info_rows = _unique(tables['Run_Info'][1], lambda row: row['Parameter'], 'run parameter')
    meta = {key: row['Value'] for key, row in info_rows.items()}
    if meta.get('Status') != 'complete':
        raise ValidationError(f'Incomplete intensity: {marker}')
    meta['Schema_version'] = collect.number(meta.get('Schema_version'), 'Schema_version', integer=True)
    if meta['Schema_version'] not in (1, 2):
        raise ValidationError('Unsupported intensity schema')
    channel = int(MARKER_PATTERN.fullmatch(marker)[1])
    if (collect.number(meta.get('Marker_channel'), 'Marker_channel', integer=True) != channel
            or (meta['Schema_version'] == 2 and meta.get('Marker_folder') != marker)):
        raise ValidationError('Marker identity differs from its metadata')
    run = Path(collect.basename(meta.get('Intensity_run', '')))
    if not collect.INTENSITY_PATTERN.fullmatch(run.name):
        raise ValidationError('Invalid intensity run reference')
    bundle = {'run': run, 'metadata': meta, 'marker': marker, 'channel': channel, 'tables': tables}
    collect.validate_rows(bundle, 'intensity')
    return bundle


def load_collection(path):
    """Use archived tables only; original images and original drives are not required."""
    path = Path(path)
    files = {}
    info = collection_info(path, files)
    copied_rows = [r for r in info if r.get('Category') == 'Copied spreadsheet' and r.get('Status') == 'COPIED']
    copied = _unique(copied_rows, lambda row: row['Item'], 'copied spreadsheet')
    for name, row in copied.items():
        if Path(name).name != name or name.startswith('.') or Path(name).suffix not in ('.csv', '.xlsx'):
            raise ValidationError('Invalid archived spreadsheet name')
        content = collect.snapshot(path / name, files)
        if hashlib.sha256(content).hexdigest() != row.get('SHA256'):
            raise ValidationError(f'Archived spreadsheet changed since collection: {name}')
    morph_names = {'Nuclei_Morphology.csv', 'Nuclei_Images.csv', 'Nuclei_Run_Info.csv', 'Nuclei_Morphology.xlsx'}
    if not morph_names <= set(copied):
        raise ValidationError('Collection lacks the four morphology source records')
    tables, meta, _ = collect.read_tables(path, 'morphology', files)
    run_id = meta.get('Run_ID', '')
    if not collect.NUCLEI_PATTERN.fullmatch(run_id):
        raise ValidationError('Invalid nuclei run reference')
    morphology = {'run': Path(run_id), 'metadata': {'particle_size_pixels_squared': meta.get('Particle_size_px2')},
                  'tables': tables}
    collect.validate_rows(morphology, 'morphology')
    combined = {
        'Nuclei': collect.csv_snapshot(path / 'FIA_Marker_Intensity_Nuclei.csv', files,
                                      ['Mask_name', 'Nucleus_ID', 'Nuclei_run_ID', 'Image_name', 'Source_file']),
        'Images': collect.csv_snapshot(path / 'FIA_Marker_Intensity_Images.csv', files,
                                      ['Mask_name', 'Nuclei_run_ID', 'Image_name', 'Source_file']),
    }
    collect.verify_workbook(path / collect.COMBINED_NAME, files, combined)
    merged_nuclei = _unique(combined['Nuclei'][1], _nucleus_key, 'combined nucleus')
    merged_images = _unique(combined['Images'][1], lambda row: row['Mask_name'], 'combined image')
    if set(merged_nuclei) != set(morphology['nuclei']) or set(merged_images) != set(morphology['images']):
        raise ValidationError('Combined and morphology populations differ')
    for sheet, merged, source in (('Nuclei', merged_nuclei, morphology['nuclei']),
                                   ('Images', merged_images, morphology['images'])):
        for key, row in source.items():
            _check_values(collect.typed_morphology(row), merged[key], {'Image_name': 'Morphology_Image_name'})
    marker_lists = [{col.removesuffix('_Availability') for col in combined[sheet][0] if col.endswith('_Availability')}
                    for sheet in ('Nuclei', 'Images')]
    if marker_lists[0] != marker_lists[1] or any(not MARKER_PATTERN.fullmatch(m) for m in marker_lists[0]):
        raise ValidationError('Inconsistent marker columns')
    markers = sorted(marker_lists[0])
    bundles, availability = {}, {}
    for marker in markers:
        states = {row.get(marker + '_Availability') for sheet in combined.values() for row in sheet[1]}
        if len(states) != 1 or not states <= {'AVAILABLE', 'MISSING'}:
            raise ValidationError(f'Inconsistent availability: {marker}')
        availability[marker] = states.pop()
        if availability[marker] == 'MISSING':
            for _, rows in combined.values():
                if any(value for row in rows for col, value in row.items()
                       if col.startswith(marker + '_') and col != marker + '_Availability'):
                    raise ValidationError('Missing marker contains measurements')
            continue
        bundle = _read_intensity(path, marker, files, copied)
        collect.validate_compatibility(morphology, bundle)
        bundles[marker] = bundle
        for sheet, merged, source in (('Nuclei', merged_nuclei, bundle['nuclei']),
                                      ('Images', merged_images, bundle['images'])):
            metrics = list(collect.INTENSITY_METRICS) if sheet == 'Nuclei' else [
                f'{metric}_{stat}' for metric in collect.INTENSITY_METRICS for stat in collect.STATS]
            if sheet == 'Images':
                metrics += ['Nuclei_measured_count', 'Failed_nuclei_count']
            for key, row in source.items():
                expected = {'Image_name': row['Image_name'], 'Source_file': row['Source_file'],
                            marker + '_Intensity_run_ID': bundle['run'].name}
                expected.update({marker + '_' + metric: collect.number(row[metric], metric, blank=True) for metric in metrics})
                _check_values(expected, merged[key])
    images, nuclei = [], []
    for row in merged_images.values():
        record = collect.typed_morphology(row)
        record['Well'] = image_well(record)
        images.append(record)
    image_map = {row['Mask_name']: row for row in images}
    for key, row in merged_nuclei.items():
        record = collect.typed_morphology(row)
        image = image_map[key[0]]
        if record['Image_name'] != image['Image_name'] or image_well(record) != image['Well']:
            raise ValidationError('Nucleus/image identity mismatch')
        record['Well'] = image['Well']
        for marker in markers:
            for metric in collect.INTENSITY_METRICS:
                field = marker + '_' + metric
                record[field] = collect.number(row.get(field), field, blank=availability[marker] == 'MISSING')
        nuclei.append(record)
    return {'path': path, 'files': files, 'info': info, 'images': images, 'nuclei': nuclei,
            'markers': markers, 'availability': availability, 'bundles': bundles,
            'run_id': run_id, 'particle_size': morphology['particle_size'], 'notes': []}


def image_well(row):
    value = str(row.get('Well', ''))
    match = re.fullmatch(r'([A-H])(0?[1-9]|1[0-2])', value, re.IGNORECASE)
    if not match:
        raise ValidationError(f'Invalid or missing well: {value!r}')
    well = f'{match[1].upper()}{int(match[2]):02d}'
    name = row.get('Image_name') or row['Mask_name']
    matches = WELL_PATTERN.findall(name)
    if len(matches) != 1 or f'{matches[0][0].upper()}{int(matches[0][1]):02d}' != well:
        raise ValidationError(f'Well differs from filename: {name}')
    points = re.findall(r'Point([A-H])(0?[1-9]|1[0-2])(?!\d)', name, re.IGNORECASE)
    if any(f'{letter.upper()}{int(column):02d}' != well for letter, column in points):
        raise ValidationError(f'Well/Point tokens disagree: {name}')
    return well


def marker_catalog(data):
    """Prefer named marker runs over legacy channel-only exports; never pool aliases."""
    markers, notes = list(data['markers']), []
    for marker in data['markers']:
        if not marker.startswith('Channel_'):
            continue
        channel = MARKER_PATTERN.fullmatch(marker)[1]
        named = [m for m in markers if m.startswith('Foci_') and MARKER_PATTERN.fullmatch(m)[1] == channel]
        if named:
            markers.remove(marker)
            notes.append(f'{marker} excluded: use named marker result(s) {", ".join(named)}; no legacy fallback.')
    return markers, notes


def _color(cell):
    color = cell.fill.fgColor
    if cell.fill.patternType != 'solid' or color.type not in ('rgb', 'theme', 'indexed'):
        raise ValidationError(f'{cell.coordinate}: use a direct solid fill')
    return f'{color.type}:{color.value}:tint={color.tint}'


def read_template(path, files, sheet_name=None, statistics=False):
    book = load_workbook(io.BytesIO(collect.snapshot(path, files)), read_only=False, data_only=False)
    try:
        sheet = book[sheet_name] if sheet_name else book.worksheets[0]
        if any(area.min_row <= 9 and area.min_col <= 13 for area in sheet.merged_cells.ranges):
            raise ValidationError('Merged cells in the plate grid')
        if ([str(sheet.cell(1, col).value) for col in range(2, 14)] != [str(n) for n in range(1, 13)]
                or [str(sheet.cell(row, 1).value).upper() for row in range(2, 10)] != list('ABCDEFGH')):
            raise ValidationError('Expected a 96-well grid at A1:M9')
        if statistics:
            for conditional in sheet.conditional_formatting:
                if any(area.min_row <= 9 and area.max_row >= 2 and area.min_col <= 13 and area.max_col >= 2
                       for area in conditional.sqref.ranges):
                    raise ValidationError('Use direct fills/bold, not conditional formatting, for the plate grid')
        design = []
        for row, letter in enumerate('ABCDEFGH', 2):
            for column in range(1, 13):
                cell = sheet.cell(row, column + 1)
                if cell.value is None or not str(cell.value).strip():
                    continue
                if cell.data_type in ('f', 'e') or not isinstance(cell.value, str):
                    raise ValidationError(f'{cell.coordinate}: group names must be literal text')
                design.append({'Well': f'{letter}{column:02d}', 'Group': cell.value, 'Excel_cell': cell.coordinate,
                               'Is_control': bool(cell.font.bold), 'Color': _color(cell) if statistics else ''})
        if not design:
            raise ValidationError('No annotated wells')
        return sheet.title, design
    finally:
        book.close()


def find_template(collection, sheet_name=None):
    candidates = []
    for path in sorted(collection.iterdir()):
        if (path.name.startswith(('.', '~$')) or path.suffix.lower() != '.xlsx' or path.is_symlink()
                or path.name == collect.COMBINED_NAME or path.name == 'Nuclei_Morphology.xlsx'
                or path.name.startswith('Nuclear_Intensity_')):
            continue
        try:
            read_template(path, {}, sheet_name)
            candidates.append(path)
        except (OSError, ValueError, KeyError, BadZipFile):
            continue
    if len(candidates) != 1:
        raise ValidationError(f'Expected one plate-map workbook, found {len(candidates)}. Supply --template explicitly.')
    return candidates[0]


def annotate(data, template, markers, stats_unit, sheet_name=None):
    if stats_unit not in (None, 'nucleus', 'well'):
        raise ValidationError('Statistics unit must be nucleus, well, or omitted')
    if len(set(markers)) != len(markers) or not set(markers) <= set(marker_catalog(data)[0]):
        raise ValidationError('Select unique marker folders from the available catalog')
    if len({MARKER_PATTERN.fullmatch(m)[1] for m in markers}) != len(markers):
        raise ValidationError('Multiple marker folders refer to the same channel; select one result per channel')
    sheet, design = read_template(template, data['files'], sheet_name, statistics=stats_unit is not None)
    mapping = {row['Well']: row['Group'] for row in design}
    for row in data['nuclei'] + data['images']:
        if row['Well'] not in mapping:
            raise ValidationError(f'Unannotated imaged well: {row["Well"]}')
        row['Group'] = mapping[row['Well']]
    groups = list(dict.fromkeys(row['Group'] for row in design))
    blocks, roles = {}, {}
    if stats_unit is not None:
        for row in design:
            identity = row['Color'], row['Is_control']
            if row['Group'] in roles and roles[row['Group']] != identity:
                raise ValidationError(f'Inconsistent fill or control role: {row["Group"]}')
            roles[row['Group']] = identity
            blocks.setdefault(row['Color'], {})[row['Group']] = row['Is_control']
        for color, members in blocks.items():
            if sum(members.values()) != 1 or len(members) < 2:
                raise ValidationError(f'{color}: each comparison block needs one control condition and a treatment')
    # Remove excluded marker aliases from the report tables, while retaining archived sources.
    excluded = set(data['markers']) - set(markers)
    for sheet_rows in (data['images'], data['nuclei']):
        for row in sheet_rows:
            for field in list(row):
                if any(field.startswith(marker + '_') for marker in excluded):
                    del row[field]
    data.update(markers=markers, template=Path(template), template_sheet=sheet,
                design=design, groups=groups, blocks=blocks, stats_unit=stats_unit)
    return data


def metric_specs(markers):
    metrics = [('Nuclei', '', COUNT, 'Nuclei per image')]
    for field in collect.MORPH_METRICS:
        units = 'px2' if field == 'Area_px2' else ('px' if field.endswith('_px') else 'dimensionless')
        metrics.append(('Morphology', '', field, units))
    for marker in markers:
        metrics.extend(('Intensity', marker, marker + '_' + field,
                        'summed raw pixel values' if field == 'Marker_RawIntDen' else 'raw intensity')
                       for field in collect.INTENSITY_METRICS)
    return metrics


def aggregate(data):
    """Keep nuclei as plot observations; compute equal-image-weight well means."""
    by_mask = defaultdict(list)
    for nucleus in data['nuclei']:
        by_mask[nucleus['Mask_name']].append(nucleus)
    image_values, well_values, summary = [], [], []
    for category, marker, field, unit in metric_specs(data['markers']):
        per_well = defaultdict(list)
        for image in data['images']:
            values = ([image[COUNT]] if field == COUNT else
                      [row[field] for row in by_mask[image['Mask_name']] if row[field] is not None])
            record = {'Category': category, 'Marker': marker, 'Metric': field, 'Unit': unit,
                      'Group': image['Group'], 'Well': image['Well'], 'Image_name': image['Image_name'],
                      'Mask_name': image['Mask_name'], 'N_nuclei': image[COUNT],
                      'N_values': len(values), 'Value': float(np.mean(values)) if values else None}
            image_values.append(record)
            per_well[image['Well']].append(record)
            if field != COUNT and values:
                # Existing image summaries are independently reconciled with individual nuclei.
                for stat, value in zip(collect.STATS, (np.mean(values), np.median(values),
                                      np.percentile(values, 75) - np.percentile(values, 25))):
                    recorded = collect.number(image.get(field + '_' + stat), field + '_' + stat, blank=True)
                    if recorded is None or not math.isclose(recorded, float(value), rel_tol=1e-9, abs_tol=1e-9):
                        raise ValidationError(f'Image summary disagrees with nuclei: {field}/{stat}')
        for annotation in data['design']:
            records = per_well[annotation['Well']]
            valid = [row['Value'] for row in records if row['Value'] is not None]
            well_values.append({'Category': category, 'Marker': marker, 'Metric': field, 'Unit': unit,
                                'Group': annotation['Group'], 'Well': annotation['Well'],
                                'N_images': len(records), 'N_images_used': len(valid),
                                'N_nuclei': sum(row['N_nuclei'] for row in records),
                                'N_values': sum(row['N_values'] for row in records),
                                'Value': float(np.mean(valid)) if valid else None})
        for group in data['groups']:
            population = data['images'] if field == COUNT else data['nuclei']
            values = [row[field] for row in population if row['Group'] == group and row[field] is not None]
            summary.append({'Category': category, 'Marker': marker, 'Metric': field, 'Group': group,
                            'Unit': unit, 'Observation': 'image' if field == COUNT else 'nucleus',
                            'N': len(values), 'Mean': float(np.mean(values)) if values else None,
                            'Median': float(np.median(values)) if values else None,
                            'IQR': float(np.percentile(values, 75) - np.percentile(values, 25)) if values else None})
    data.update(image_values=image_values, well_values=well_values, summary=summary)
    data['counts_by_group'] = {group: {'Nuclei': sum(r['Group'] == group for r in data['nuclei']),
                                      'Images': sum(r['Group'] == group for r in data['images']),
                                      'Wells': len({r['Well'] for r in data['images'] if r['Group'] == group})}
                               for group in data['groups']}
    return data
